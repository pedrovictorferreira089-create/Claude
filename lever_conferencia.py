"""Conferência do painel comercial (dados do Lever) com as vendas do Belle, card a card.

Uso: python3 lever_conferencia.py [dd/mm/aaaa início] [dd/mm/aaaa fim]
     período = "Data de Fechamento" do painel (última movimentação do card nas fases de venda do funil SDRs).
     Gera Conferencia_Lever_Belle.xlsx (ou BELLE_OUT).

Casamento: telefone do contato do card no Lever com o celular/telefone do cliente no Belle (DDD + 8 últimos dígitos,
o que ignora o nono dígito); sem telefone em comum, pelo nome completo. Cada card fica com o plano aprovado do cliente
mais próximo da data de fechamento do card (até JANELA dias). A classificação da venda (Avaliação, Cabine SDR, Cabine)
vem de belle_avaliacao.coletar().
"""
import collections
import datetime as dt
import os
import re
import sys

from openpyxl import Workbook

import lever_api
from belle_auditoria import DATA, NUM, base, data_br, nome, secao, titulo
from belle_avaliacao import AV, CAB, ESTABS, SDR, coletar, reais
from belle_sabados import BRL, HOJE, UNIDADES, historico

ANO = dt.date(2026, 1, 1)
JANELA = 45  # dias entre a movimentação do card e a venda no Belle para considerar a mesma venda
GRUPOS = (AV, SDR, CAB)


def fone(s):
    """DDD + 8 últimos dígitos (sem o 55 e sem o nono dígito), ou None."""
    d = re.sub(r"\D", "", str(s or ""))
    if len(d) >= 12 and d.startswith("55"):
        d = d[2:]
    return d[:2] + d[-8:] if len(d) >= 10 else None


def data_local(s):
    return (dt.datetime.fromisoformat(s.replace("Z", "+00:00")) - dt.timedelta(hours=3)).date() if s else None


def data_campo(v):
    m = re.match(r"(\d{4})[/-](\d{2})[/-](\d{2})", v or "")
    return dt.date(*map(int, m.groups())) if m else None


def clientes_belle(fim):
    """Telefones e nomes dos clientes do Belle que compraram de 01/01 até fim (Vendas Detalhado)."""
    por_fone, por_nome = collections.defaultdict(set), collections.defaultdict(set)
    for e in UNIDADES:
        for v in historico("vendas_detalhado", "estab", e, fim, desde=ANO):
            dc = v.get("dados_cliente") or {}
            c = v.get("cod_cliente")
            for k in ("celular", "celular2", "telefone"):
                f = fone(dc.get(k))
                if f:
                    por_fone[f].add(c)
            if dc.get("nome"):
                por_nome[nome(dc["nome"]).casefold()].add(c)
    return por_fone, por_nome


def conferir(ini, fim):
    painel = lever_api.painel()
    cards = lever_api.cards(painel)
    b = coletar(ANO, max(fim, HOJE), ESTABS)
    por_fone, por_nome = clientes_belle(max(fim, HOJE))
    planos = collections.defaultdict(list)  # cliente -> planos aprovados de 2026 nas unidades da análise
    for x in b["vendas"]:
        if x["tipo"] == "Plano" and x["conta"]:
            planos[x["cliente_cod"]].append(x)

    linhas, usados = [], {}
    for c in cards:
        fech = data_local(c["updatedAt"])
        fones = [fone(k.get("phoneNumber")) for k in c["contatos"]]
        nomes = [nome(k.get("name") or c.get("title") or "").casefold() for k in c["contatos"]] or [nome(c.get("title") or "").casefold()]
        cli, via = set(), ""
        for f in filter(None, fones):
            cli |= por_fone.get(f, set())
        if cli:
            via = "Telefone"
        else:
            for n in filter(None, nomes):
                cli |= por_nome.get(n, set())
            via = "Nome" if cli else ""
        cand = sorted((p for k in cli for p in planos.get(k, [])), key=lambda p: (abs((p["data"] - fech).days), p["data"]))
        plano = cand[0] if cand and abs((cand[0]["data"] - fech).days) <= JANELA else None
        if plano:
            usados.setdefault(plano["orc"], []).append(c["key"])
        linhas.append(dict(card=c, fech=fech, aval=data_campo(lever_api.campo(c, "data-de-avalia") or lever_api.campo(c, "data-avalia")),
                           unidade=lever_api.campo(c, "unidade"), valor=float(c.get("monetaryAmount") or 0),
                           cliente=", ".join(sorted(str(k) for k in cli)), via=via, plano=plano, cand=cand,
                           tem_fone=any(fones)))
    for x in linhas:
        p = x["plano"]
        if not x["cliente"]:
            x["motivo"] = "Cliente não encontrado no Belle (telefone e nome sem correspondência)"
        elif not x["cand"]:
            x["motivo"] = "Cliente sem plano aprovado em 2026 nas 4 unidades"
        elif not p:
            x["motivo"] = f"Plano mais próximo a mais de {JANELA} dias do fechamento do card"
        elif len(usados[p["orc"]]) > 1:
            x["motivo"] = "Mesmo plano casado com mais de um card: " + ", ".join(usados[p["orc"]])
        else:
            x["motivo"] = ""
    return dict(linhas=linhas, belle=b, painel=painel)


def montar(d, ini, fim, destino):
    linhas = d["linhas"]
    per = [x for x in linhas if ini <= x["fech"] <= fim]
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    titulo(ws, "Conferência painel (Lever) x Belle", f"Cards das fases Convertidos e Convertidos avulsos do funil SDRs, "
           f"Data de Fechamento de {ini:%d/%m/%Y} a {fim:%d/%m/%Y}. Gerado em {dt.date.today():%d/%m/%Y}.")
    r = 4
    secao(ws, r, "Cards do período")
    r += 1
    resumo = [("Cards", len(per), sum(x["valor"] for x in per)),
              ("Casados com plano do Belle", sum(1 for x in per if x["plano"]), sum(x["plano"]["valor"] for x in per if x["plano"]))]
    for motivo, n in collections.Counter(x["motivo"].split(":")[0] for x in per if x["motivo"]).most_common():
        resumo.append((motivo, n, sum(x["valor"] for x in per if x["motivo"].split(":")[0] == motivo)))
    for g in GRUPOS:
        sel = [x for x in per if x["plano"] and x["plano"]["grupo"] == g]
        resumo.append((f"  {g}", len(sel), sum(x["plano"]["valor"] for x in sel)))
    for rot, n, v in resumo:
        ws.cell(r, 1, rot)
        ws.cell(r, 2, n)
        ws.cell(r, 3, v).number_format = BRL
        r += 1
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["C"].width = 18

    cols = ["Card", "Fase", "Título", "Unidade (Lever)", "Data de fechamento (Lever)", "Data de avaliação (Lever)",
            "Valor no Lever", "Casado por", "Cód. cliente Belle", "Data da venda (Belle)", "Unidade (Belle)", "Plano",
            "Valor do plano (Belle)", "Diferença Lever − Belle", "Grupo (regras Belle)", "Perfil da venda", "Dias card x venda",
            "No período do painel", "Observação"]
    dados = []
    for x in sorted(linhas, key=lambda x: (x["fech"], x["card"]["key"]), reverse=True):
        p, c = x["plano"], x["card"]
        dados.append([c["key"], c["fase"], nome(c.get("title") or ""), x["unidade"], x["fech"], x["aval"], x["valor"],
                      x["via"] or None, x["cliente"] or None, p and p["data"], p and p["unidade"], p and p["desc"],
                      p and p["valor"], p and round(x["valor"] - p["valor"], 2), p and p["grupo"], p and p["perfil"],
                      p and (p["data"] - x["fech"]).days, "Sim" if ini <= x["fech"] <= fim else "Não", x["motivo"]])
    fm = [None, None, None, None, DATA, DATA, BRL, None, None, DATA, None, None, BRL, BRL, None, None, NUM, None, None]
    base(wb.create_sheet("Cards"), cols, dados, fm, "Cards", [11, 18, 30, 15, 13, 13, 13, 10, 12, 13, 15, 30, 13, 13, 12, 45, 9, 9, 60])
    wb.save(destino)


if __name__ == "__main__":
    a = sys.argv[1:]
    ini = data_br(a[0]) if a else dt.date(2026, 9, 28)
    fim = data_br(a[1]) if len(a) > 1 else dt.date(2026, 10, 4)
    d = conferir(ini, fim)
    out = os.environ.get("BELLE_OUT", "Conferencia_Lever_Belle.xlsx")
    montar(d, ini, fim, out)
    per = [x for x in d["linhas"] if ini <= x["fech"] <= fim]
    print(f"{len(per)} cards ({reais(sum(x['valor'] for x in per))}) – {sum(1 for x in per if x['plano'])} casados – {out}")
