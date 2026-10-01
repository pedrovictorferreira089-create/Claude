"""Recompra de planos: quantos clientes voltam a comprar, depois de quanto tempo e em que momento do tratamento
(API Belle Software).

Uso: python3 belle_recompra.py
Lê o histórico completo de vendas de planos, atendimentos e vendas (cadastro dos clientes) de todas as unidades
e gera Recompra_Clientes_Drenesse.xlsx (ou BELLE_OUT).

Regras:
- Compra = dia em que o cliente fechou pelo menos um plano pago (preço final > 0) com status Aprovado, ou
  Suspenso depois de pago (estorno não conta). Planos do mesmo cliente no mesmo dia são uma compra só.
- Troca de tratamento = plano vendido até 3 dias antes ou depois da suspensão de outro plano do mesmo cliente.
  Se não custa mais que o plano trocado, continua a compra original (não é recompra); se custa mais, é nova compra.
- Sessões usadas = saldo do Belle (vendidas − restantes − agendamentos futuros) nos serviços pagos (sem cortesia),
  distribuídas pelas datas dos atendimentos realizados de cada plano. Plano suspenso conta só o que chegou a usar.
- Dias sem vir = dias desde a última sessão realizada ou a última compra, o que for mais recente.
"""
import collections
import datetime as dt
import os
import statistics

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

from belle_auditoria import DATA, NUM, PCT, base, compactar_xlsx, data_br, linha, nome, secao, titulo
from belle_sabados import FB, FI, FILL_IN, HOJE, SUB, UNIDADES, api, br, cabecalho, historico, larguras

TROCA_DIAS = 3
ATIVO_DIAS = 45    # sem vir há até 45 dias, ou com agendamento marcado = em tratamento
RETA_FINAL = 0.75  # parte do plano usada ou agendada a partir da qual a renovação deve ser oferecida
INATIVOS = ("Desmarcado", "Cancelado", "Falhou", "Atendido")
FAMILIAS = [  # (família, trechos do nome do serviço principal) – vale a primeira que bater
    ("Drenagem Método Drenesse", ("DRENAGEM",)),
    ("Crioturbo e pós turbo", ("CRIO", "PÓS TURBO", "POS TURBO")),
    ("Laser (depilação)", ("DEPILA", "LASER", "PASSE LIVRE")),
    ("Diástase", ("DIÁSTASE", "DIASTASE")),
    ("Limpeza de pele", ("LIMPEZA",)),
    ("Injetáveis (toxina, preenchimento)", ("BOTOX", "DYSPORT", "NABOTA", "TOXINA", "PREENCH", "BIOESTIMULADOR RENOVA",
                                           "SCULPTRA", "RADIESSE", "PEIM", "INJET")),
    ("Facial (peeling, microagulhamento, manchas)", ("PEELING", "MICROAGULH", "MICRO AGULH", "MELLAN", "MELAN", "LINE SKIN",
                                                     "ACNE", "FACIAL", "HAIR TECH", "MICROROBOTICO")),
    ("Stimulusculp, glúteos e flacidez", ("STIMULUS", "GLUTE", "GLÚTE", "BUMBUM", "FLACI", "STRIORT", "ESTRIA", "RÁDIO",
                                          "RADIO", "PLASMA", "BLEFARO", "FIRM")),
    ("Modelagem (lipo escultura, enzimas, detox)", ("LIPO", "ENZIMA", "DETOX", "SLIM", "SHAPE", "ORTOSHOCK", "TERMOTERAPIA",
                                                    "BBB")),
    ("Massagem", ("MASSAGEM",)),
]
FAMILIAS_TODAS = [f for f, _ in FAMILIAS] + ["Outros"]
FAIXAS_INT = [("Até 7 dias", 0, 7), ("8 a 30 dias", 8, 30), ("31 a 90 dias", 31, 90), ("91 a 180 dias", 91, 180),
              ("181 a 365 dias", 181, 365), ("1 a 2 anos", 366, 730), ("Mais de 2 anos", 731, None)]
MOMENTOS = ["Complemento: anterior com menos de 50% usado", "Renovação: anterior com 50% a 99% usado",
            "Renovação: anterior concluído (até 30 dias)", "Retorno após 1 a 3 meses sem vir",
            "Retorno após 3 a 12 meses sem vir", "Retorno após mais de 1 ano sem vir"]
SITUACOES = [  # (situação, ação sugerida, quando agir) – na ordem de prioridade da lista de campanhas
    ("Reta final do plano", "Oferecer o próximo plano nas últimas sessões (renovação antecipada com condição especial)",
     "Ao passar de 75% das sessões usadas ou agendadas"),
    ("Parado com sessões a usar", "Ligar ou chamar no WhatsApp para reagendar as sessões que faltam; na volta, apresentar o próximo passo",
     "A partir de 45 dias sem vir"),
    ("Concluiu há até 3 meses", "Contato pós-tratamento: avaliar resultado e oferecer manutenção ou tratamento complementar",
     "Até 30 dias após a última sessão"),
    ("Concluiu há 3 a 12 meses", "Campanha de retorno: reavaliação gratuita e condição de volta por tempo limitado",
     "Uma vez por mês, para quem completou 3 meses sem vir"),
    ("Em tratamento", "Apresentar um tratamento complementar na reavaliação (a maior parte das recompras é complemento)",
     "Nas primeiras semanas do plano"),
    ("Sem vir há mais de 1 ano", "Reativação em datas promocionais e pesquisa do motivo de não ter voltado",
     "Dia das Mães, aniversário do cliente, Black Friday"),
]
LISTA = ["Reta final do plano", "Parado com sessões a usar", "Concluiu há até 3 meses", "Concluiu há 3 a 12 meses"]
FAIXAS_VALOR = [("Até R$ 500", 0, 500), ("R$ 500 a R$ 1.000", 500, 1000), ("R$ 1.000 a R$ 2.000", 1000, 2000),
                ("R$ 2.000 a R$ 4.000", 2000, 4000), ("Acima de R$ 4.000", 4000, None)]
TIPOS = ["Plano personalizado", "Pacote ou promoção"]


def familia(servico):
    s = (servico or "").upper()
    return next((f for f, chaves in FAMILIAS if any(c in s for c in chaves)), "Outros")


def cliente_de(p):
    cod, _, nm = p["cliente"].partition(" - ")
    return int(cod), nome(nm)


def carregar():
    """Planos (todos os status), validade, atendimentos realizados e agendamentos futuros por plano, visitas e
    próximo agendamento por cliente, e o cadastro mais recente de cada cliente (das vendas)."""
    planos = {}
    for e in UNIDADES:
        for p in historico("venda_planos", "codEstab", e, HOJE):
            if p.get("cliente") and p["codOrcamento"] not in planos:
                planos[p["codOrcamento"]] = dict(p, estab=e)
    validade = {}
    for e in UNIDADES:
        for x in historico("relatorios/sessoes_planos", "codEstab", e, HOJE, p_ini="dtVndIni", p_fim="dtVndFim"):
            v = data_br(x.get("dtValidade"))
            if v:
                validade[x["idPlano"]] = max(v, validade.get(x["idPlano"], v))
    sessoes = collections.defaultdict(list)   # plano -> [(data, código do serviço)] atendidos
    futuros = collections.Counter()            # (plano, código do serviço) -> agendamentos em aberto de hoje em diante
    visitas = collections.defaultdict(set)    # cliente -> dias com atendimento realizado
    agendado = {}                              # cliente -> próximo agendamento em aberto
    vistos = set()
    amanha = HOJE + dt.timedelta(days=1)
    for e in UNIDADES:
        ags = historico("relatorios/relatorio_atendimentos", "codEstab", e, HOJE)
        ags = ags + api("relatorios/relatorio_atendimentos", dtInicio=amanha.strftime("%d/%m/%Y"),
                        dtFim=(amanha + dt.timedelta(days=88)).strftime("%d/%m/%Y"), codEstab=str(e))
        for a in ags:
            if a["idAgendamento"] in vistos or not a.get("codigoCliente"):
                continue
            vistos.add(a["idAgendamento"])
            d, c, st = data_br(a["dataAgendamento"]), a["codigoCliente"], a["statusAgendamento"]
            cod = str(a["codigoServico"] or "")
            if st == "Atendido" and d <= HOJE:
                visitas[c].add(d)
                if a["idOrcamento"]:
                    sessoes[a["idOrcamento"]].append((d, cod))
            elif d >= HOJE and st not in INATIVOS:
                agendado[c] = min(d, agendado.get(c, d))
                if a["idOrcamento"]:
                    futuros[(a["idOrcamento"], cod)] += 1
    cadastro = {}
    for e in UNIDADES:
        for v in historico("vendas_detalhado", "estab", e, HOJE):
            dc = v.get("dados_cliente") or {}
            c = v.get("cod_cliente") or dc.get("codigo")
            if c and (c not in cadastro or (v.get("data_venda") or "") >= cadastro[c][0]):
                cadastro[c] = (v.get("data_venda") or "", dc)
    return planos, validade, sessoes, futuros, visitas, agendado, {c: dc for c, (_, dc) in cadastro.items()}


def vendido(p):
    """Plano que conta como venda: preço final > 0, Aprovado ou Suspenso depois de pago (sem estorno)."""
    if br(p["precoFinal"]) <= 0 or not data_br(p.get("dataVenda")):
        return False
    if p["statusPlano"] == "Aprovado":
        return True
    return (p["statusPlano"] == "Suspenso" and (p.get("motivoRescisao") or "").strip().lower() != "estorno de venda"
            and any(x["confirmado"] == "Sim" for x in p["parcelas"]))


def pagos(p):
    """Sessões pagas (sem cortesia) por serviço do plano."""
    out = collections.Counter()
    for s in p["servicos"]:
        if s["cortesia"] != "Sim" and s["qtdSessoes"]:
            out[str(s["codigoServico"])] += s["qtdSessoes"]
    return out


def uso_plano(p, sessoes, futuros):
    """Sessões pagas do plano: (vendidas, a agendar hoje, [(data, sessões consumidas)] dos atendimentos realizados).
    O total consumido vem do saldo do Belle (um atendimento pode consumir mais de uma sessão, ex.: áreas de crio)."""
    vend, rest = collections.Counter(), collections.Counter()
    for s in p.get("saldoPlano") or []:
        vend[str(s["codigoServico"])] += int(br(s["qtdSessaoVendida"]))
        rest[str(s["codigoServico"])] += int(br(s["qtdSessaoRestante"]))
    feitos = collections.defaultdict(list)
    for d, cod in sessoes.get(p["codOrcamento"], []):
        feitos[cod].append(d)
    vendidas = a_agendar = 0
    eventos = []
    for cod, qtd in pagos(p).items():
        datas = sorted(feitos[cod])
        if cod in vend:
            usadas = min(qtd, max(vend[cod] - rest[cod] - futuros[(p["codOrcamento"], cod)], 0))
            restam = min(qtd, rest[cod])
        else:
            usadas = min(qtd, len(datas))
            restam = qtd - usadas
        if p["statusPlano"] == "Suspenso":
            vendidas += usadas
        else:
            vendidas += qtd
            a_agendar += restam
        eventos += [(d, usadas / len(datas)) for d in datas] if datas else []
    return vendidas, a_agendar, eventos


def principal(planos_compra):
    """Serviço pago de maior valor (preço cheio menos desconto) entre os planos da compra."""
    peso = collections.Counter()
    for p in planos_compra:
        for s in p["servicos"]:
            if s["cortesia"] == "Sim":
                continue
            v = br(s["valorTotalServico"])
            v = v * (1 - br(s["desconto"]) / 100) if s["tipoDesconto"] == "%" else v - br(s["desconto"])
            peso[" ".join((s["nomeServico"] or "").split())] += max(v, 0.0) + (s["qtdSessoes"] or 0) * 1e-6
    return peso.most_common(1)[0][0] if peso else ""


def montar_compras(planos, sessoes, futuros):
    """Agrupa os planos de cada cliente em compras (ocasiões), tratando trocas de tratamento."""
    por_cliente = collections.defaultdict(list)
    for p in planos.values():
        if p["statusPlano"] in ("Aprovado", "Suspenso") and data_br(p.get("dataVenda")):
            por_cliente[cliente_de(p)[0]].append(p)
    compras, trocas = [], collections.Counter()
    for c, ps in por_cliente.items():
        ps.sort(key=lambda p: (data_br(p["dataVenda"]), p["codOrcamento"]))
        papel, origem, chave = {}, {}, {}  # plano -> troca/upgrade; troca -> plano substituído; plano -> data da compra
        for s in ps:
            if s["statusPlano"] != "Suspenso" or not data_br(s.get("dataRescisao")):
                continue
            r = data_br(s["dataRescisao"])
            cand = [p for p in ps if p is not s and data_br(p["dataVenda"]) >= data_br(s["dataVenda"])
                    and abs((data_br(p["dataVenda"]) - r).days) <= TROCA_DIAS and p["codOrcamento"] not in papel]
            if not cand:
                continue
            dia = min(data_br(p["dataVenda"]) for p in cand)
            cand = [p for p in cand if data_br(p["dataVenda"]) == dia]
            upgrade = sum(br(p["precoFinal"]) for p in cand) > br(s["precoFinal"]) + 0.005
            for p in cand:
                papel[p["codOrcamento"]] = "upgrade" if upgrade else "troca"
                origem[p["codOrcamento"]] = s["codOrcamento"]
        for p in ps:  # ordem cronológica: a compra original já tem data quando a troca aparece
            orc = p["codOrcamento"]
            if papel.get(orc) == "troca":
                chave[orc] = chave.get(origem[orc], data_br(p["dataVenda"]))
            else:
                chave[orc] = data_br(p["dataVenda"])
            if vendido(p) and orc in papel:
                trocas[papel[orc]] += 1
        grupos = collections.defaultdict(list)
        for p in ps:
            grupos[chave[p["codOrcamento"]]].append(p)
        for d in sorted(grupos):
            g = grupos[d]
            pagos_g = [p for p in g if vendido(p) and papel.get(p["codOrcamento"]) != "troca"]
            if not pagos_g:
                continue
            vendidas = a_agendar = 0
            eventos = []
            for p in g:  # inclui trocas e planos de valor zero do mesmo dia (bônus), que fazem parte do tratamento
                v, r, ev = uso_plano(p, sessoes, futuros)
                vendidas, a_agendar, eventos = vendidas + v, a_agendar + r, eventos + ev
            nomes = collections.Counter(" ".join(p["nomePlano"].split()) for p in pagos_g)
            maior = max(pagos_g, key=lambda p: br(p["precoFinal"]))
            servico = principal(pagos_g)
            compras.append(dict(
                cliente=c, nome=cliente_de(maior)[1], data=d, estab=maior["estab"],
                unidade=UNIDADES.get(maior["estab"], str(maior["estab"])),
                planos=", ".join(f"{k} (x{v})" if v > 1 else k for k, v in nomes.items()),
                qtd_planos=len(pagos_g), valor=round(sum(br(p["precoFinal"]) for p in pagos_g), 2),
                servico=servico, familia=familia(servico),
                tipo=TIPOS[any(p["nomePlano"].strip().upper() != "PLANO PERSONALIZADO" for p in pagos_g)],
                vendedor=nome(maior.get("vendedor")), upgrade=any(papel.get(p["codOrcamento"]) == "upgrade" for p in pagos_g),
                orcs=[p["codOrcamento"] for p in g], sessoes=vendidas, a_agendar=a_agendar, eventos=sorted(eventos)))
    compras.sort(key=lambda x: (x["cliente"], x["data"]))
    return compras, trocas


def momento(dias_sem_vir, pct_ant):
    """Em que ponto da relação com a clínica a recompra aconteceu (mesma regra da fórmula da aba Compras)."""
    if dias_sem_vir <= 30:
        if pct_ant is not None and pct_ant < 0.5:
            return MOMENTOS[0]
        return MOMENTOS[1] if pct_ant is not None and pct_ant < 1 else MOMENTOS[2]
    return MOMENTOS[3] if dias_sem_vir <= 90 else MOMENTOS[4] if dias_sem_vir <= 365 else MOMENTOS[5]


def situacao(c):
    """Situação atual do cliente pela última compra (mesma regra da fórmula da aba Clientes)."""
    if c["a_agendar"] > 0 and (c["dias_sem_vir"] <= ATIVO_DIAS or c["agendado"]):
        return SITUACOES[0][0] if c["pct"] >= RETA_FINAL else "Em tratamento"
    if c["a_agendar"] > 0 and c["dias_sem_vir"] <= 365:
        return "Parado com sessões a usar"
    if c["dias_sem_vir"] <= 90:
        return "Concluiu há até 3 meses"
    return "Concluiu há 3 a 12 meses" if c["dias_sem_vir"] <= 365 else "Sem vir há mais de 1 ano"


def analisar(compras, visitas, agendado, validade):
    """Completa cada compra com intervalo, uso da compra anterior, dias sem vir e ciclo até a próxima compra;
    devolve a visão por cliente."""
    por_cliente = collections.defaultdict(list)
    for x in compras:
        por_cliente[x["cliente"]].append(x)
    clientes = []
    for c, cs in por_cliente.items():
        vis = sorted(visitas.get(c, ()))
        for k, x in enumerate(cs):
            x["n"] = k + 1
            x["chave"] = f"{c}-{k + 1}"
            x["proxima"] = cs[k + 1]["data"] if k + 1 < len(cs) else None
            fim = x["proxima"] or HOJE + dt.timedelta(days=1)
            x["ultima_sessao"] = max((d for d in vis if x["data"] <= d < fim), default=None) or x["data"]
            x.update(anterior=None, dias_ant=None, pct_ant=None, fam_ant="", dias_sem_vir=None, momento="")
            if k == 0:
                continue
            ant = cs[k - 1]
            usado = sum(p for d, p in ant["eventos"] if d < x["data"])
            contato = max([d for d in vis if d < x["data"]] + [ant["data"]])
            x.update(anterior=ant["data"], dias_ant=(x["data"] - ant["data"]).days, fam_ant=ant["familia"],
                     pct_ant=round(min(usado / ant["sessoes"], 1.0), 4) if ant["sessoes"] else None,
                     dias_sem_vir=(x["data"] - contato).days)
            x["momento"] = momento(x["dias_sem_vir"], x["pct_ant"])
        p, u = cs[0], cs[-1]
        ult = max((d for d in vis if d <= HOJE), default=None)
        cli = dict(cliente=c, nome=u["nome"], primeira=p, ultima=u, n=len(cs), segunda=cs[1]["data"] if len(cs) > 1 else None,
                   total=round(sum(x["valor"] for x in cs), 2), a_agendar=u["a_agendar"],
                   pct=(u["sessoes"] - u["a_agendar"]) / u["sessoes"] if u["sessoes"] else 0.0,
                   ult_visita=ult, dias_sem_vir=(HOJE - max(d for d in (ult, u["data"]) if d)).days,
                   agendado=agendado.get(c), validade=max((validade[o] for o in u["orcs"] if o in validade), default=None))
        cli["situacao"] = situacao(cli)
        clientes.append(cli)
    clientes.sort(key=lambda c: (c["primeira"]["data"], c["cliente"]))
    return clientes


def chance_por_tempo(compras, so_primeira=False, dias=(30, 60, 90, 180)):
    """Ciclos (da compra até a próxima) com a última sessão há 1 ano ou mais: % que recomprou até d dias depois da
    última sessão e, para quem não recomprou até d, a chance de recomprar até 365 dias depois dela."""
    obs = [x for x in compras if (HOJE - x["ultima_sessao"]).days >= 365 and (x["n"] == 1 or not so_primeira)]
    t = [(x["proxima"] - x["ultima_sessao"]).days if x["proxima"] else None for x in obs]
    ate = lambda d: sum(1 for v in t if v is not None and v <= d)
    return len(obs), {d: (ate(d) / len(obs), (ate(365) - ate(d)) / (len(obs) - ate(d))) for d in dias + (365,)}


def numeros(compras, clientes, trocas):
    """Principais números (para o relatório e para conferir a planilha)."""
    rec = [c for c in clientes if c["n"] > 1]
    d12 = [(c["segunda"] - c["primeira"]["data"]).days for c in rec]
    out = dict(clientes=len(clientes), recompraram=len(rec), pct=len(rec) / len(clientes), tres_mais=sum(c["n"] >= 3 for c in clientes),
               compras=len(compras), mediana_1_2=statistics.median(d12), media_1_2=statistics.mean(d12),
               mediana_todas=statistics.median(x["dias_ant"] for x in compras if x["n"] > 1), trocas=dict(trocas))
    for w in (91, 183, 365):
        eleg = [c for c in clientes if (HOJE - c["primeira"]["data"]).days >= w]
        out[f"ate_{w}"] = sum(1 for c in eleg if c["segunda"] and (c["segunda"] - c["primeira"]["data"]).days <= w) / len(eleg)
    v1 = sum(x["valor"] for x in compras if x["n"] == 1)
    vr = sum(x["valor"] for x in compras if x["n"] > 1)
    out.update(valor=v1 + vr, valor_rec=vr, pct_valor_rec=vr / (v1 + vr),
               ticket_1=v1 / len(clientes), ticket_rec=vr / (len(compras) - len(clientes)),
               gasto_1x=statistics.mean(c["total"] for c in clientes if c["n"] == 1), gasto_rec=statistics.mean(c["total"] for c in rec))
    out["momentos"] = collections.Counter(x["momento"] for x in compras if x["n"] > 1)
    out["momentos_2a"] = collections.Counter(x["momento"] for x in compras if x["n"] == 2)
    out["situacoes"] = collections.Counter(c["situacao"] for c in clientes)
    out["chance"] = chance_por_tempo(compras)
    out["chance_1a"] = chance_por_tempo(compras, so_primeira=True)
    return out


def escolher_exemplos(clientes):
    """Um cliente real para cada padrão de comportamento (o de maior gasto entre os que se encaixam)."""
    def melhor(cond, chave=lambda c: c["total"]):
        cands = [c for c in clientes if cond(c)]
        return max(cands, key=chave) if cands else None

    regras = [
        ("Cliente fiel: compra um plano novo a cada 1 a 2 meses",
         lambda c: 8 <= c["n"] <= 14 and c["ultima"]["data"] >= HOJE - dt.timedelta(days=120)
         and 25 <= statistics.median(x["dias_ant"] for x in c["_todas"][1:]) <= 60),
        ("Complemento rápido: 2ª compra de outro tratamento poucas semanas depois da 1ª",
         lambda c: 3 <= c["n"] <= 6 and c["_todas"][1]["momento"] == MOMENTOS[0] and c["_todas"][1]["dias_ant"] <= 30
         and c["_todas"][1]["familia"] != c["primeira"]["familia"]),
        ("Renovação: compra de novo o mesmo tratamento perto do fim ou logo depois de concluir",
         lambda c: 3 <= c["n"] <= 6 and sum(x["momento"] in MOMENTOS[1:3] and x["familia"] == x["fam_ant"] for x in c["_todas"]) >= 2),
        ("Retorno depois de uma pausa: concluiu, ficou meses sem vir e voltou",
         lambda c: 2 <= c["n"] <= 5 and any(x["momento"] == MOMENTOS[4] for x in c["_todas"])),
        ("Comprou uma vez, concluiu e não voltou",
         lambda c: c["n"] == 1 and c["situacao"] == "Concluiu há 3 a 12 meses"),
        ("Comprou uma vez e parou com sessões a usar",
         lambda c: c["n"] == 1 and c["situacao"] == "Parado com sessões a usar" and c["primeira"]["sessoes"] >= 8),
    ]
    return [(t, melhor(r)) for t, r in regras if melhor(r)]


# ---------------------------------------------------------------- planilha
BRL0 = 'R$ #,##0;-R$ #,##0;"-"'


def colunas(defs):
    """{chave: letra} a partir da lista de colunas (chave, título, formato, largura)."""
    return {k: get_column_letter(i) for i, (k, *_) in enumerate(defs, 1)}


CLIENTES = [
    ("cod", "Código", "0", 10), ("nome", "Cliente", None, 32), ("cel", "Celular", None, 18),
    ("unid1", "Unidade da 1ª compra", None, 14), ("data1", "Data da 1ª compra", DATA, 11), ("ano1", "Ano da 1ª compra", "0", 8),
    ("planos1", "1ª compra – planos", None, 30), ("valor1", "Valor da 1ª compra (R$)", BRL0, 13),
    ("serv1", "Serviço principal da 1ª compra", None, 30), ("fam1", "Família da 1ª compra", None, 26),
    ("tipo1", "Tipo da 1ª compra", None, 18), ("n", "Nº de compras", NUM, 9), ("voltou", "Voltou a comprar?", None, 9),
    ("data2", "Data da 2ª compra", DATA, 11), ("dias2", "Dias até a 2ª compra", NUM, 10),
    ("r3", "Recompra em até 3 meses", None, 12), ("r6", "Recompra em até 6 meses", None, 12),
    ("r12", "Recompra em até 12 meses", None, 12), ("dataU", "Data da última compra", DATA, 11),
    ("total", "Total em planos (R$)", BRL0, 13), ("ticket", "Ticket médio por compra (R$)", BRL0, 12),
    ("intervalo", "Intervalo médio entre compras (dias)", NUM, 11), ("planosU", "Última compra – planos", None, 30),
    ("famU", "Família da última compra", None, 26), ("sessU", "Sessões pagas da última compra", NUM, 9),
    ("restU", "Sessões a agendar (saldo Belle)", NUM, 9), ("pctU", "% usado ou agendado", PCT, 9),
    ("ultS", "Última sessão realizada", DATA, 11), ("semVir", "Dias sem vir", NUM, 9),
    ("agenda", "Próximo agendamento", DATA, 11), ("validade", "Validade do último plano", DATA, 11),
    ("situacao", "Situação atual", None, 24), ("acao", "Ação sugerida", None, 60),
]
COMPRAS = [
    ("chave", "Chave", None, 12), ("cod", "Código do cliente", "0", 10), ("nome", "Cliente", None, 32), ("n", "Nº da compra", NUM, 8),
    ("data", "Data da compra", DATA, 11), ("ano", "Ano", "0", 7), ("unid", "Unidade", None, 14), ("planos", "Planos", None, 32),
    ("qtd", "Nº de planos", NUM, 8), ("valor", "Valor (R$)", BRL0, 12), ("serv", "Serviço principal", None, 30),
    ("fam", "Família", None, 26), ("tipo", "Tipo", None, 18), ("vend", "Vendedor", None, 26), ("sess", "Sessões pagas", NUM, 8),
    ("dataAnt", "Data da compra anterior", DATA, 11), ("dias", "Dias desde a compra anterior", NUM, 10),
    ("faixa", "Faixa do intervalo", None, 14), ("famAnt", "Família da compra anterior", None, 26),
    ("mesma", "Mesma família da anterior?", None, 10), ("pctAnt", "% da compra anterior usado na data", PCT, 10),
    ("semVir", "Dias sem vir antes desta compra", NUM, 10), ("momento", "Momento da recompra", None, 40),
    ("ultCiclo", "Última sessão antes da próxima compra", DATA, 11), ("prox", "Data da próxima compra", DATA, 11),
    ("diasProx", "Dias da última sessão até a próxima compra", NUM, 11),
    ("obs", "Última sessão há 1 ano ou mais?", None, 11), ("upgrade", "Troca com upgrade?", None, 9),
    ("orcs", "Nº dos planos no Belle", None, 24),
]
CL, CP = colunas(CLIENTES), colunas(COMPRAS)


def montar(compras, clientes, cadastro, trocas, exemplos, destino, compactar=True):
    wb = Workbook()
    nC, nP = len(clientes) + 1, len(compras) + 1
    inicio = min(x["data"] for x in compras)

    def C(k):
        return f"Clientes!${CL[k]}$2:${CL[k]}${nC}"

    def P(k):
        return f"Compras!${CP[k]}$2:${CP[k]}${nP}"

    fonte = (f"Período: {inicio:%d/%m/%Y} a {HOJE:%d/%m/%Y} • Todas as unidades ({', '.join(UNIDADES.values())}) • "
             f"Fonte: API Belle Software, extraído em {HOJE:%d/%m/%Y}")

    # ---------------- Resumo
    rs = wb.active
    rs.title = "Resumo"
    titulo(rs, "Drenesse – Recompra de planos: quantos clientes voltam a comprar e em quanto tempo", fonte)
    rs["A4"], rs["B4"] = "Data de referência", HOJE
    rs["A4"].font, rs["B4"].font, rs["B4"].number_format, rs["B4"].fill = FB, FB, DATA, FILL_IN
    rs["C4"] = "Usada nos cálculos de dias sem vir e de quem já teve tempo de recomprar (data da extração)."
    rs["C4"].font = FI
    wb.defined_names["DataRef"] = DefinedName("DataRef", attr_text="Resumo!$B$4")
    secao(rs, 6, "1. Quantos clientes voltam a comprar")
    cabecalho(rs, 7, ["Indicador", "Valor", "Como ler"])
    elig = lambda k: f'(COUNTIFS({C(k)},"SIM")+COUNTIFS({C(k)},"NÃO"))'
    inds = [
        ("Clientes que compraram pelo menos um plano", f"=COUNT({C('cod')})", NUM, "Clientes com ao menos uma compra de plano pago no período."),
        ("Voltaram a comprar (2 compras ou mais)", f'=COUNTIFS({C("n")},">=2")', NUM, "Compraram outro plano em outro dia, depois da 1ª compra."),
        ("% que voltou a comprar (até hoje)", "=IFERROR(B9/B8,0)", PCT, "Inclui clientes recentes, que ainda podem voltar: veja as taxas por prazo abaixo."),
        ("Compraram 3 vezes ou mais", f'=COUNTIFS({C("n")},">=3")', NUM, ""),
        ("% de recompra em até 3 meses", f'=IFERROR(COUNTIFS({C("r3")},"SIM")/{elig("r3")},0)', PCT,
         "Só clientes cuja 1ª compra foi há 3 meses ou mais (todos tiveram o mesmo prazo para voltar)."),
        ("% de recompra em até 6 meses", f'=IFERROR(COUNTIFS({C("r6")},"SIM")/{elig("r6")},0)', PCT, "Só clientes com a 1ª compra há 6 meses ou mais."),
        ("% de recompra em até 12 meses", f'=IFERROR(COUNTIFS({C("r12")},"SIM")/{elig("r12")},0)', PCT, "Só clientes com a 1ª compra há 1 ano ou mais."),
        ("Mediana de dias entre a 1ª e a 2ª compra", f"=MEDIAN({C('dias2')})", NUM, "Metade de quem voltou fez a 2ª compra em até este número de dias."),
        ("Média de dias entre a 1ª e a 2ª compra", f"=AVERAGE({C('dias2')})", NUM, "A média é puxada para cima por quem volta depois de anos."),
        ("Mediana de dias entre compras seguidas (todas as recompras)", f"=MEDIAN({P('dias')})", NUM, "Considera também 2ª→3ª, 3ª→4ª etc."),
        ("Compras de planos no período", f"=COUNTA({P('chave')})", NUM, "Compra = dia em que o cliente fechou um ou mais planos pagos."),
        ("Valor vendido em planos (R$)", f"=SUM({P('valor')})", BRL0, "Preço final dos planos (Aprovados ou suspensos depois de pagos)."),
        ("Valor vindo de recompras – 2ª compra em diante (R$)", f'=SUMIFS({P("valor")},{P("n")},">1")', BRL0, ""),
        ("% do valor vindo de recompras", "=IFERROR(B20/B19,0)", PCT, "Quanto do faturamento em planos depende de clientes que voltam."),
        ("Ticket médio da 1ª compra (R$)", f'=IFERROR(AVERAGEIFS({P("valor")},{P("n")},1),0)', BRL0, ""),
        ("Ticket médio das recompras (R$)", f'=IFERROR(AVERAGEIFS({P("valor")},{P("n")},">1"),0)', BRL0, ""),
        ("Gasto médio de quem comprou uma vez (R$)", f'=IFERROR(AVERAGEIFS({C("total")},{C("n")},1),0)', BRL0, ""),
        ("Gasto médio de quem voltou a comprar (R$)", f'=IFERROR(AVERAGEIFS({C("total")},{C("n")},">1"),0)', BRL0, "Soma de todas as compras do cliente."),
        ("Trocas de tratamento não contadas como recompra", trocas.get("troca", 0), NUM,
         "Plano vendido até 3 dias da suspensão de outro, sem custar mais: continua a compra original (contado pelo script)."),
        ("Trocas com upgrade contadas como nova compra", trocas.get("upgrade", 0), NUM, "O novo plano custou mais que o trocado (contado pelo script)."),
    ]
    for i, (t, f, fmt, como) in enumerate(inds, 8):
        linha(rs, i, [t, f, como], [None, fmt, None])
    r = 8 + len(inds) + 1
    secao(rs, r, "2. Número de compras por cliente")
    r += 1
    cabecalho(rs, r, ["Nº de compras", "Clientes", "% dos clientes", "Valor total (R$)", "% do valor"])
    h2 = r
    faixas_n = [("1 compra", "=1", None), ("2 compras", "=2", None), ("3 compras", "=3", None), ("4 compras", "=4", None),
                ("5 compras", "=5", None), ("6 a 9 compras", ">=6", "<=9"), ("10 compras ou mais", ">=10", None)]
    t2 = h2 + len(faixas_n) + 1
    for nome_f, c1, c2 in faixas_n:
        r += 1
        crit = f'{C("n")},"{c1}"' + (f',{C("n")},"{c2}"' if c2 else "")
        linha(rs, r, [nome_f, f"=COUNTIFS({crit})", f"=IFERROR(B{r}/$B${t2},0)", f"=SUMIFS({C('total')},{crit})",
                      f"=IFERROR(D{r}/$D${t2},0)"], [None, NUM, PCT, BRL0, PCT])
    r += 1
    linha(rs, r, ["TOTAL"] + [f"=SUM({c}{h2 + 1}:{c}{r - 1})" for c in "BCDE"], [None, NUM, PCT, BRL0, PCT], total=True)
    notas = [
        "COMO LER",
        "• Compra = dia em que o cliente fechou pelo menos um plano pago (preço final maior que zero), Aprovado ou suspenso depois de pago. Vários planos no mesmo dia contam como uma compra.",
        "• Recompra = nova compra em outro dia. Troca de tratamento (plano suspenso e outro vendido em até 3 dias) não conta, a menos que o novo plano custe mais (upgrade).",
        "• O histórico da API começa em janeiro de 2021: quem já era cliente antes aparece com a 1ª compra em 2021.",
        "• Abas: 'Tempo entre compras', 'Momento da recompra', 'Perfis', 'Complementos', 'Faturamento' e 'Campanhas' são calculadas a partir das bases 'Clientes' e 'Compras'.",
        "• 'Campanhas' traz a situação atual de cada cliente e a lista de quem abordar agora; 'Exemplos' mostra jornadas reais.",
    ]
    r += 2
    for i, n_ in enumerate(notas):
        rs.cell(row=r + i, column=1, value=n_)
    rs.cell(row=r, column=1).font = FB
    larguras(rs, [52, 16, 90, 16, 12])
    rs.freeze_panes = "A5"

    # ---------------- Tempo entre compras
    wt = wb.create_sheet("Tempo entre compras")
    titulo(wt, "Tempo entre uma compra e a próxima", fonte)
    secao(wt, 4, "A. Da 1ª para a 2ª compra (clientes que voltaram)")
    cabecalho(wt, 5, ["Intervalo", "Clientes", "% de quem voltou", "% acumulado"])
    r = 5
    tA = 5 + len(FAIXAS_INT) + 1
    for nome_f, lo, hi in FAIXAS_INT:
        r += 1
        crit = f'{C("dias2")},">={lo}"' + (f',{C("dias2")},"<={hi}"' if hi is not None else "")
        linha(wt, r, [nome_f, f"=COUNTIFS({crit})", f"=IFERROR(B{r}/$B${tA},0)", f"=SUM($C$6:C{r})"], [None, NUM, PCT, PCT])
    r += 1
    linha(wt, r, ["TOTAL", "=SUM(B6:B%d)" % (r - 1), "=SUM(C6:C%d)" % (r - 1), ""], [None, NUM, PCT, None], total=True)
    r += 2
    secao(wt, r, "B. Entre compras seguidas, pelo número da compra (quantas recompras caíram em cada faixa)")
    r += 1
    hB = r
    trans = [("1ª → 2ª", "=2"), ("2ª → 3ª", "=3"), ("3ª → 4ª", "=4"), ("4ª → 5ª em diante", ">=5"), ("Todas", ">=2")]
    cabecalho(wt, r, ["Intervalo"] + [t for t, _ in trans])
    for nome_f, _, _ in FAIXAS_INT:
        r += 1
        linha(wt, r, [nome_f] + [f'=COUNTIFS({P("n")},"{c}",{P("faixa")},$A{r})' for _, c in trans], [None] + [NUM] * len(trans))
    r += 1
    linha(wt, r, ["Total de recompras"] + [f"=SUM({get_column_letter(j)}{hB + 1}:{get_column_letter(j)}{r - 1})" for j in range(2, 2 + len(trans))],
          [None] + [NUM] * len(trans), total=True)
    r += 1
    linha(wt, r, ["Mediana (dias)"] + [""] * len(trans), [None] + [NUM] * len(trans), fill=SUB)
    for j, (_, c) in enumerate(trans, 2):
        cond = f'({P("n")}{"=" + c[1:] if c.startswith("=") else c})'
        cel = f"{get_column_letter(j)}{r}"
        wt[cel] = ArrayFormula(cel, f'=IFERROR(MEDIAN(IF({cond},{P("dias")})),"")')
    r += 1
    linha(wt, r, ["% em até 30 dias"] + [f"=IFERROR((COUNTIFS({P('n')},\"{c}\",{P('dias')},\"<=30\"))/{get_column_letter(j)}{r - 2},0)"
                                          for j, (_, c) in enumerate(trans, 2)], [None] + [PCT] * len(trans), fill=SUB)
    wt.cell(row=r + 2, column=1, value="Quanto mais compras o cliente já fez, mais curto o intervalo até a próxima: quem cria o hábito volta todo mês.").font = FI
    larguras(wt, [30, 14, 14, 14, 18, 14])
    wt.row_dimensions[hB].height = 30

    # ---------------- Momento da recompra
    wm = wb.create_sheet("Momento da recompra")
    titulo(wm, "Em que momento o cliente compra de novo", fonte)
    secao(wm, 4, "A. Momento de cada recompra (2ª compra em diante)")
    cabecalho(wm, 5, ["Momento", "Recompras", "% das recompras", "Só a 2ª compra", "% das 2ªs compras",
                      "Mediana de dias desde a compra anterior", "Valor médio (R$)", "% da mesma família da anterior"])
    r = 5
    tM = 5 + len(MOMENTOS) + 1
    for m in MOMENTOS:
        r += 1
        linha(wm, r, [m, f"=COUNTIFS({P('momento')},$A{r})", f"=IFERROR(B{r}/$B${tM},0)", f"=COUNTIFS({P('momento')},$A{r},{P('n')},2)",
                      f"=IFERROR(D{r}/$D${tM},0)", "", f"=IFERROR(AVERAGEIFS({P('valor')},{P('momento')},$A{r}),0)",
                      f'=IFERROR(COUNTIFS({P("momento")},$A{r},{P("mesma")},"SIM")/B{r},0)'], [None, NUM, PCT, NUM, PCT, NUM, BRL0, PCT])
        wm[f"F{r}"] = ArrayFormula(f"F{r}", f'=IFERROR(MEDIAN(IF({P("momento")}=$A{r},{P("dias")})),"")')
    r += 1
    linha(wm, r, ["TOTAL", f"=SUM(B6:B{r - 1})", f"=SUM(C6:C{r - 1})", f"=SUM(D6:D{r - 1})", f"=SUM(E6:E{r - 1})", "", "", ""],
          [None, NUM, PCT, NUM, PCT, None, None, None], total=True)
    leg = ["Complemento = comprou outro plano com o cliente vindo (até 30 dias sem vir) e menos da metade do plano anterior usado – em geral, outro tratamento.",
           "Renovação = comprou com o cliente vindo e metade ou mais do plano anterior usado, ou até 30 dias depois de concluir.",
           "Retorno = comprou depois de mais de 30 dias sem vir (sem sessão e sem compra)."]
    for i, t in enumerate(leg):
        wm.cell(row=r + 1 + i, column=1, value=t).font = FI
    r += len(leg) + 3
    secao(wm, r, "B. Chance de comprar de novo conforme o tempo sem vir (depois da última sessão)")
    r += 1
    hC = r
    cabecalho(wm, r, ["Dias depois da última sessão", "Ciclos que já tinham recomprado", "% acumulado",
                      "Se não recomprou até aqui: chance de recomprar até 1 ano", "Só clientes novos (1ª compra): % acumulado",
                      "Só clientes novos: chance de recomprar até 1 ano"])
    obs_all = f'{P("obs")},"SIM"'
    obs_1 = f'{P("obs")},"SIM",{P("n")},1'
    marcos = [7, 15, 30, 45, 60, 90, 180, 365]
    for d in marcos:
        r += 1
        linha(wm, r, [d, f'=COUNTIFS({obs_all},{P("diasProx")},"<={d}")', f"=IFERROR(B{r}/COUNTIFS({obs_all}),0)",
                      f'=IFERROR((COUNTIFS({obs_all},{P("diasProx")},"<=365")-B{r})/(COUNTIFS({obs_all})-B{r}),0)',
                      f'=IFERROR(COUNTIFS({obs_1},{P("diasProx")},"<={d}")/COUNTIFS({obs_1}),0)',
                      f'=IFERROR((COUNTIFS({obs_1},{P("diasProx")},"<=365")-COUNTIFS({obs_1},{P("diasProx")},"<={d}"))/'
                      f'(COUNTIFS({obs_1})-COUNTIFS({obs_1},{P("diasProx")},"<={d}")),0)'], [NUM, NUM, PCT, PCT, PCT, PCT])
    wm.cell(row=r + 1, column=1, value=(
        "Ciclo = período entre uma compra e a seguinte. Só entram ciclos cuja última sessão foi há 1 ano ou mais, para todos terem o mesmo prazo "
        "de observação. Ex.: na linha 60, a última coluna mostra a chance de um cliente que está há 60 dias sem vir (e não recomprou) ainda comprar até completar 1 ano.")).font = FI
    r += 3
    secao(wm, r, "C. Quanto do plano anterior já tinha sido usado na hora da recompra")
    r += 1
    hD = r
    cabecalho(wm, r, ["% do plano anterior usado", "Recompras", "% das recompras", "Mediana de dias desde a compra anterior"])
    bandas = [("Menos de 25%", 0, 0.25), ("25% a 49%", 0.25, 0.5), ("50% a 74%", 0.5, 0.75), ("75% a 99%", 0.75, 1), ("100% (concluído)", 1, None)]
    tD = hD + len(bandas) + 1
    for nome_f, lo, hi in bandas:
        r += 1
        crit = f'{P("pctAnt")},">={lo}"' + (f',{P("pctAnt")},"<{hi}"' if hi is not None else "")
        cond = f"({P('pctAnt')}>={lo})" + (f"*({P('pctAnt')}<{hi})" if hi is not None else "") + f'*({P("pctAnt")}<>"")'
        linha(wm, r, [nome_f, f"=COUNTIFS({crit})", f"=IFERROR(B{r}/$B${tD},0)", ""], [None, NUM, PCT, NUM])
        wm[f"D{r}"] = ArrayFormula(f"D{r}", f'=IFERROR(MEDIAN(IF({cond},{P("dias")})),"")')
    r += 1
    linha(wm, r, ["TOTAL", f"=SUM(B{hD + 1}:B{r - 1})", f"=SUM(C{hD + 1}:C{r - 1})", ""], [None, NUM, PCT, None], total=True)
    larguras(wm, [46, 16, 14, 22, 18, 20, 14, 16])
    for rr in (5, hC, hD):
        wm.row_dimensions[rr].height = 45

    # ---------------- Perfis
    wp = wb.create_sheet("Perfis")
    titulo(wp, "Quem volta mais: por unidade, tratamento, tipo e valor da 1ª compra e ano em que começou", fonte)
    cols_p = ["Grupo (pela 1ª compra)", "Clientes", "Voltaram a comprar", "% que voltou (até hoje)", "Clientes com 1ª compra há 1 ano ou mais",
              "% de recompra em até 12 meses", "Mediana de dias até a 2ª compra", "Ticket médio da 1ª compra (R$)", "Gasto médio por cliente (R$)"]
    fm_p = [None, NUM, NUM, PCT, NUM, PCT, NUM, BRL0, BRL0]
    r = 3

    def bloco(titulo_b, grupos, crit, cond):
        nonlocal r
        r += 1
        secao(wp, r, titulo_b)
        r += 1
        cabecalho(wp, r, cols_p)
        wp.row_dimensions[r].height = 45
        for g in grupos:
            r += 1
            k = crit(g, r)
            linha(wp, r, [g, f"=COUNTIFS({k})", f'=COUNTIFS({k},{C("n")},">1")', f"=IFERROR(C{r}/B{r},0)",
                          f'=COUNTIFS({k},{C("r12")},"SIM")+COUNTIFS({k},{C("r12")},"NÃO")',
                          f'=IFERROR(COUNTIFS({k},{C("r12")},"SIM")/E{r},"")', "",
                          f"=IFERROR(AVERAGEIFS({C('valor1')},{k}),0)", f"=IFERROR(AVERAGEIFS({C('total')},{k}),0)"], fm_p)
            wp[f"G{r}"] = ArrayFormula(f"G{r}", f'=IFERROR(MEDIAN(IF({cond(g, r)}*({C("data2")}<>""),{C("dias2")})),"")')
        r += 1

    bloco("A. Por unidade da 1ª compra", list(UNIDADES.values()), lambda g, r: f"{C('unid1')},$A{r}", lambda g, r: f"({C('unid1')}=$A{r})")
    bloco("B. Por tratamento principal da 1ª compra", FAMILIAS_TODAS, lambda g, r: f"{C('fam1')},$A{r}", lambda g, r: f"({C('fam1')}=$A{r})")
    bloco("C. Por tipo da 1ª compra", TIPOS, lambda g, r: f"{C('tipo1')},$A{r}", lambda g, r: f"({C('tipo1')}=$A{r})")
    lim = {f: (lo, hi) for f, lo, hi in FAIXAS_VALOR}
    bloco("D. Por valor da 1ª compra", [f for f, _, _ in FAIXAS_VALOR],
          lambda g, r: f'{C("valor1")},">={lim[g][0]}"' + (f',{C("valor1")},"<{lim[g][1]}"' if lim[g][1] else ""),
          lambda g, r: f"({C('valor1')}>={lim[g][0]})" + (f"*({C('valor1')}<{lim[g][1]})" if lim[g][1] else ""))
    anos = sorted({c["primeira"]["data"].year for c in clientes})
    bloco("E. Por ano da 1ª compra (safra)", anos, lambda g, r: f"{C('ano1')},$A{r}", lambda g, r: f"({C('ano1')}=$A{r})")
    wp.cell(row=r + 1, column=1, value=("A coluna '% de recompra em até 12 meses' compara grupos de forma justa: só entra quem já teve 1 ano para voltar. "
                                        "A safra de 2021 inclui clientes antigos (o histórico da API começa em 2021).")).font = FI
    larguras(wp, [44, 11, 11, 13, 16, 14, 14, 15, 15])

    # ---------------- Complementos (de → para)
    wc = wb.create_sheet("Complementos")
    titulo(wc, "O que o cliente compra na recompra, conforme o tratamento da compra anterior", fonte)
    secao(wc, 4, "% das recompras (linha = tratamento da compra anterior; coluna = tratamento da nova compra)")
    cabecalho(wc, 5, ["Compra anterior ↓ / Nova compra →"] + FAMILIAS_TODAS + ["Recompras"])
    wc.row_dimensions[5].height = 75
    for cell in wc[5]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = 5
    for f in FAMILIAS_TODAS:
        r += 1
        tot = get_column_letter(2 + len(FAMILIAS_TODAS))
        linha(wc, r, [f] + [f'=IFERROR(COUNTIFS({P("famAnt")},$A{r},{P("fam")},{get_column_letter(j)}$5)/${tot}{r},0)'
                            for j in range(2, 2 + len(FAMILIAS_TODAS))] + [f'=COUNTIFS({P("famAnt")},$A{r})'],
              [None] + [PCT] * len(FAMILIAS_TODAS) + [NUM])
    fim_col = get_column_letter(1 + len(FAMILIAS_TODAS))
    wc.conditional_formatting.add(f"B6:{fim_col}{r}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                     end_type="num", end_value=0.6, end_color="5B9BD5"))
    wc.cell(row=r + 2, column=1, value=("Leia por linha: de quem comprou Crioturbo, que % fez a compra seguinte em cada tratamento. "
                                        "A diagonal é recompra do mesmo tratamento; fora dela, complemento (oportunidade de venda cruzada).")).font = FI
    larguras(wc, [40] + [13] * len(FAMILIAS_TODAS) + [11])
    wc.freeze_panes = "B6"

    # ---------------- Faturamento
    wf = wb.create_sheet("Faturamento")
    titulo(wf, "Faturamento em planos: 1ª compra x recompras, por ano da compra", fonte)
    cabecalho(wf, 4, ["Ano", "1ªs compras", "Valor das 1ªs compras (R$)", "Recompras", "Valor das recompras (R$)", "Total (R$)",
                      "% do valor vindo de recompras", "Ticket médio da 1ª compra (R$)", "Ticket médio da recompra (R$)"])
    wf.row_dimensions[4].height = 45
    r = 4
    for a in sorted({x["data"].year for x in compras}):
        r += 1
        linha(wf, r, [a, f"=COUNTIFS({P('ano')},$A{r},{P('n')},1)", f"=SUMIFS({P('valor')},{P('ano')},$A{r},{P('n')},1)",
                      f'=COUNTIFS({P("ano")},$A{r},{P("n")},">1")', f'=SUMIFS({P("valor")},{P("ano")},$A{r},{P("n")},">1")',
                      f"=C{r}+E{r}", f"=IFERROR(E{r}/F{r},0)", f"=IFERROR(C{r}/B{r},0)", f"=IFERROR(E{r}/D{r},0)"],
              ["0", NUM, BRL0, NUM, BRL0, BRL0, PCT, BRL0, BRL0])
    r += 1
    linha(wf, r, ["TOTAL", f"=SUM(B5:B{r - 1})", f"=SUM(C5:C{r - 1})", f"=SUM(D5:D{r - 1})", f"=SUM(E5:E{r - 1})", f"=SUM(F5:F{r - 1})",
                  f"=IFERROR(E{r}/F{r},0)", f"=IFERROR(C{r}/B{r},0)", f"=IFERROR(E{r}/D{r},0)"],
          [None, NUM, BRL0, NUM, BRL0, BRL0, PCT, BRL0, BRL0], total=True)
    wf.cell(row=r + 1, column=1, value=f"{HOJE.year} vai até {HOJE:%d/%m}. Em 2021 quase todo cliente aparece como 1ª compra porque o histórico começa nesse ano.").font = FI
    larguras(wf, [10, 12, 18, 12, 18, 18, 14, 16, 16])

    # ---------------- Campanhas
    wk = wb.create_sheet("Campanhas")
    titulo(wk, "Situação atual de cada cliente e quem abordar agora", fonte + " • Situação pela última compra de cada cliente")
    secao(wk, 4, "A. Carteira por situação")
    cabecalho(wk, 5, ["Situação", "Clientes", "% dos clientes", "Já compraram 2 vezes ou mais", "Gasto médio por cliente (R$)",
                      "Ação sugerida", "Quando agir"])
    wk.row_dimensions[5].height = 30
    r = 5
    s0, s1 = 6, 5 + len(SITUACOES)
    for s, acao, quando in SITUACOES:
        r += 1
        linha(wk, r, [s, f"=COUNTIFS({C('situacao')},$A{r})", f"=IFERROR(B{r}/$B${s1 + 1},0)", f'=COUNTIFS({C("situacao")},$A{r},{C("n")},">1")',
                      f"=IFERROR(AVERAGEIFS({C('total')},{C('situacao')},$A{r}),0)", acao, quando], [None, NUM, PCT, NUM, BRL0, None, None])
    r += 1
    linha(wk, r, ["TOTAL", f"=SUM(B{s0}:B{s1})", f"=SUM(C{s0}:C{s1})", f"=SUM(D{s0}:D{s1})",
                  f"=IFERROR(AVERAGE({C('total')}),0)", "", ""], [None, NUM, PCT, NUM, BRL0, None, None], total=True)
    regras = [f"Em tratamento = tem sessões a agendar e veio nos últimos {ATIVO_DIAS} dias (ou tem agendamento marcado). Reta final = isso e {RETA_FINAL:.0%} ou mais do plano usado ou agendado.",
              f"Parado com sessões a usar = tem sessões a agendar, não vem há mais de {ATIVO_DIAS} dias e não tem agendamento. Concluiu = não tem mais sessões a agendar.",
              "Dias sem vir contam da última sessão realizada ou da última compra, o que for mais recente. Sessões a agendar = saldo do Belle na data da extração."]
    for i, t in enumerate(regras):
        wk.cell(row=r + 1 + i, column=1, value=t).font = FI
    r += len(regras) + 3
    lista = sorted((c for c in clientes if c["situacao"] in LISTA), key=lambda c: (LISTA.index(c["situacao"]), -c["total"]))
    secao(wk, r, f"B. Lista para abordar agora – {len(lista)} clientes ({', '.join(LISTA)}), em ordem de prioridade e de gasto")
    r += 1
    hL = r
    cols_l = [("Código", "cod", "0"), ("Cliente", "nome", None), ("Celular", "cel", None), ("Situação", "situacao", None),
              ("Última compra – planos", "planosU", None), ("Data da última compra", "dataU", DATA), ("Sessões a agendar", "restU", NUM),
              ("% usado ou agendado", "pctU", PCT), ("Última sessão", "ultS", DATA), ("Dias sem vir", "semVir", NUM),
              ("Nº de compras", "n", NUM), ("Total em planos (R$)", "total", BRL0), ("Próximo agendamento", "agenda", DATA),
              ("Validade do último plano", "validade", DATA), ("Ação sugerida", "acao", None)]
    cabecalho(wk, r, ["Prioridade"] + [t for t, _, _ in cols_l] + ["Linha na aba Clientes"])
    wk.row_dimensions[r].height = 30
    lin = get_column_letter(len(cols_l) + 2)
    for i, c in enumerate(lista, 1):
        r += 1
        vals = [i, c["cliente"]] + [f'=IF(INDEX({C(k)},${lin}{r})="","",INDEX({C(k)},${lin}{r}))' for _, k, _ in cols_l[1:]]
        vals.append(f"=MATCH(B{r},{C('cod')},0)")
        linha(wk, r, vals, [NUM] + [f for _, _, f in cols_l] + [NUM])
    larguras(wk, [26, 11, 30, 16, 24, 30, 11, 10, 10, 11, 9, 9, 13, 11, 11, 60, 10])
    wk.freeze_panes = wk.cell(row=hL + 1, column=4)

    # ---------------- Exemplos
    we = wb.create_sheet("Exemplos")
    titulo(we, "Exemplos reais de jornada de compra (um cliente por padrão)", fonte)
    r = 3
    cols_e = [("Nº da compra", "n", NUM), ("Data da compra", "data", DATA), ("Unidade", "unid", None), ("Planos", "planos", None),
              ("Família", "fam", None), ("Valor (R$)", "valor", BRL0), ("Dias desde a compra anterior", "dias", NUM),
              ("% da anterior usado", "pctAnt", PCT), ("Dias sem vir antes", "semVir", NUM), ("Momento da recompra", "momento", None)]
    for t, c in exemplos:
        r += 2
        we.cell(row=r, column=1, value=f"{t} – {c['nome']} (código {c['cliente']})").font = FB
        r += 1
        total = f"{c['total']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        we.cell(row=r, column=1, value=f"{c['n']} compra(s), total R$ {total} em planos • situação atual: {c['situacao']}").font = FI
        r += 1
        cabecalho(we, r, ["Chave"] + [h for h, _, _ in cols_e])
        for x in c["_todas"][:12]:
            r += 1
            linha(we, r, [x["chave"]] + [f'=IF(INDEX({P(k)},MATCH($A{r},{P("chave")},0))="","",INDEX({P(k)},MATCH($A{r},{P("chave")},0)))'
                                         for _, k, _ in cols_e], [None] + [f for _, _, f in cols_e])
        if c["n"] > 12:
            r += 1
            we.cell(row=r, column=1, value=f"… e mais {c['n'] - 12} compras (veja a aba Compras).").font = FI
    larguras(we, [12, 9, 11, 14, 40, 26, 12, 12, 11, 11, 40])

    # ---------------- Clientes (base)
    wcl = wb.create_sheet("Clientes")
    L = CL
    dados = []
    for i, c in enumerate(clientes, 2):
        p, u = c["primeira"], c["ultima"]
        cad = cadastro.get(c["cliente"]) or {}
        f = {k: f"{L[k]}{i}" for k in L}
        dados.append([
            c["cliente"], c["nome"], (cad.get("celular") or "").strip() or None, p["unidade"], p["data"], f"=YEAR({f['data1']})",
            p["planos"], p["valor"], p["servico"], p["familia"], p["tipo"], c["n"], f'=IF({f["n"]}>1,"SIM","NÃO")', c["segunda"],
            f'=IF({f["data2"]}="","",{f["data2"]}-{f["data1"]})',
            *[f'=IF(DataRef-{f["data1"]}<{w},"Em observação",IF(AND({f["data2"]}<>"",{f["data2"]}-{f["data1"]}<={w}),"SIM","NÃO"))'
              for w in (91, 183, 365)],
            u["data"], c["total"], f"=IFERROR({f['total']}/{f['n']},0)", f'=IF({f["n"]}>1,({f["dataU"]}-{f["data1"]})/({f["n"]}-1),"")',
            u["planos"], u["familia"], u["sessoes"], u["a_agendar"], f'=IF({f["sessU"]}>0,({f["sessU"]}-{f["restU"]})/{f["sessU"]},"")',
            c["ult_visita"], f"=DataRef-MAX({f['ultS']},{f['dataU']})", c["agendado"], c["validade"],
            (f'=IF(AND({f["restU"]}>0,OR({f["semVir"]}<={ATIVO_DIAS},{f["agenda"]}<>"")),IF({f["pctU"]}>={RETA_FINAL},"{SITUACOES[0][0]}","Em tratamento"),'
             f'IF(AND({f["restU"]}>0,{f["semVir"]}<=365),"Parado com sessões a usar",IF({f["semVir"]}<=90,"Concluiu há até 3 meses",'
             f'IF({f["semVir"]}<=365,"Concluiu há 3 a 12 meses","Sem vir há mais de 1 ano"))))'),
            f'=IFERROR(INDEX(Campanhas!$F${s0}:$F${s1},MATCH({f["situacao"]},Campanhas!$A${s0}:$A${s1},0)),"")'])
    base(wcl, [t for _, t, _, _ in CLIENTES], dados, [fm for _, _, fm, _ in CLIENTES], "TabClientes", [w for *_, w in CLIENTES])

    # ---------------- Compras (base)
    wcp = wb.create_sheet("Compras")
    L = CP
    dados = []
    for i, x in enumerate(compras, 2):
        f = {k: f"{L[k]}{i}" for k in L}
        faixa = '""'
        for nome_f, lo, hi in reversed(FAIXAS_INT):
            faixa = f'"{nome_f}"' if hi is None else f'IF({f["dias"]}<={hi},"{nome_f}",{faixa})'
        dados.append([
            x["chave"], x["cliente"], x["nome"], x["n"], x["data"], f"=YEAR({f['data']})", x["unidade"], x["planos"], x["qtd_planos"],
            x["valor"], x["servico"], x["familia"], x["tipo"], x["vendedor"] or None, x["sessoes"], x["anterior"],
            f'=IF({f["dataAnt"]}="","",{f["data"]}-{f["dataAnt"]})', f'=IF({f["dias"]}="","",{faixa})', x["fam_ant"] or None,
            f'=IF({f["famAnt"]}="","",IF({f["fam"]}={f["famAnt"]},"SIM","NÃO"))', x["pct_ant"], x["dias_sem_vir"],
            (f'=IF({f["n"]}=1,"",IF({f["semVir"]}<=30,IF(AND({f["pctAnt"]}<>"",{f["pctAnt"]}<0.5),"{MOMENTOS[0]}",'
             f'IF(AND({f["pctAnt"]}<>"",{f["pctAnt"]}<1),"{MOMENTOS[1]}","{MOMENTOS[2]}")),'
             f'IF({f["semVir"]}<=90,"{MOMENTOS[3]}",IF({f["semVir"]}<=365,"{MOMENTOS[4]}","{MOMENTOS[5]}"))))'),
            x["ultima_sessao"], x["proxima"], f'=IF({f["prox"]}="","",{f["prox"]}-{f["ultCiclo"]})',
            f'=IF(DataRef-{f["ultCiclo"]}>=365,"SIM","NÃO")', "SIM" if x["upgrade"] else None,
            ", ".join(str(o) for o in x["orcs"])])
    base(wcp, [t for _, t, _, _ in COMPRAS], dados, [fm for _, _, fm, _ in COMPRAS], "TabCompras", [w for *_, w in COMPRAS])

    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    if compactar:
        form_cl = {CL[k] for k in ("ano1", "voltou", "dias2", "r3", "r6", "r12", "ticket", "intervalo", "pctU", "semVir", "situacao", "acao")}
        form_cp = {CP[k] for k in ("ano", "dias", "faixa", "mesma", "momento", "diasProx", "obs")}
        form_k = {get_column_letter(j) for j in range(3, len(cols_l) + 3)}
        compactar_xlsx(destino, {"Clientes": form_cl, "Compras": form_cp, "Campanhas": form_k})


if __name__ == "__main__":
    planos, validade, sessoes, futuros, visitas, agendado, cadastro = carregar()
    compras, trocas = montar_compras(planos, sessoes, futuros)
    clientes = analisar(compras, visitas, agendado, validade)
    por_cliente = collections.defaultdict(list)
    for x in compras:
        por_cliente[x["cliente"]].append(x)
    for c in clientes:
        c["_todas"] = por_cliente[c["cliente"]]
    exemplos = escolher_exemplos(clientes)
    out = os.environ.get("BELLE_OUT", "Recompra_Clientes_Drenesse.xlsx")
    montar(compras, clientes, cadastro, trocas, exemplos, out)
    k = numeros(compras, clientes, trocas)
    print(f"{k['clientes']} clientes, {k['compras']} compras, {k['recompraram']} voltaram a comprar ({k['pct']:.1%}) -> {out}")
