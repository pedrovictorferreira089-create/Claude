"""Auditoria de preços: atendimentos x tabela, planos, serviços e faturamento mensal (API Belle Software).

Uso: python3 belle_auditoria.py
Busca todo o histórico disponível até ontem das unidades Petrópolis, Lagoa Nova, Capim Macio e Norte Shopping
e gera Auditoria_Precos_Drenesse.xlsx (ou BELLE_OUT). BELLE_DESDE e BELLE_ATE (dd/mm/aaaa) limitam o período.
Regras de valor iguais às de belle_sabados.py: sessão de plano Aprovado = preço final rateado pelas sessões,
descontada a parte do plano em parcelas vencidas e não pagas; avulso = venda registrada no dia.
"""
import collections
import datetime as dt
import html
import os
import re
import zipfile

import openpyxl.workbook.workbook as _wbmod
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from belle_sabados import (BOX, BRL, CENTER, F, FB, FI, FT, HOJE, TOT, api, br, cabecalho, carregar_planos, historico,
                           larguras, valorar)

# Fonte padrão do arquivo = Arial 10: as abas de dados (centenas de milhares de linhas) dispensam estilo célula a célula.
_wbmod.DEFAULT_FONT = Font(name="Arial", size=10)

UNIDS = {1: "Petrópolis", 2: "Lagoa Nova", 3: "Capim Macio", 6: "Norte Shopping"}
ATE = HOJE - dt.timedelta(days=1)
CORTE = dt.date(2027, 10, 31)  # planos com vencimento (validade) depois de outubro/2027
NUM = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
DATA = "dd/mm/yyyy"
MES = "mmm/yyyy"
CORTESIA = "Cortesia / plano 100% desconto"
MOTIVOS = ["Desconto no plano", CORTESIA, "Parcelas vencidas em aberto", "Plano não aprovado", "Plano não localizado",
           "Desconto na venda avulsa", "Avulso sem venda no dia", "Acima da tabela", "Sem diferença",
           "Sem preço de tabela", "Sem serviço vinculado"]
ORIGENS = ["Plano", "Plano c/ atraso", "Plano não aprovado", "Plano não localizado", "Avulso", "Avulso s/ venda", "Sem serviço"]
FAIXAS = [("Sem desconto (0% ou acima da tabela)", ["<=0"]), ("Até 10%", [">0", "<=0.1"]), ("De 10% a 30%", [">0.1", "<=0.3"]),
          ("De 30% a 50%", [">0.3", "<=0.5"]), ("De 50% a 70%", [">0.5", "<=0.7"]), ("De 70% a 90%", [">0.7", "<=0.9"]),
          ("Acima de 90%", [">0.9"]), ("Sem valor de tabela", [""])]


def data_br(s):
    return dt.datetime.strptime(s, "%d/%m/%Y").date() if s else None


def nome(s):
    # sem curingas de SOMASES (* ? ~) para os nomes poderem ser usados como critério nas fórmulas
    return " ".join((s or "").replace("*", "x").replace("?", "").replace("~", "-").split())


def canonizar(registros, campo):
    """Unifica grafias que só diferem em maiúsculas/minúsculas (SOMASES não diferencia) pela mais frequente."""
    variantes = collections.defaultdict(collections.Counter)
    for x in registros:
        variantes[x[campo].casefold()][x[campo]] += 1
    mapa = {k: c.most_common(1)[0][0] for k, c in variantes.items()}
    for x in registros:
        x[campo] = mapa[x[campo].casefold()]


def motivo(origem, tab, val):
    fixos = {"Sem serviço": "Sem serviço vinculado", "Plano não aprovado": "Plano não aprovado",
             "Plano não localizado": "Plano não localizado", "Avulso s/ venda": "Avulso sem venda no dia",
             "Plano c/ atraso": "Parcelas vencidas em aberto"}
    if origem in fixos:
        return fixos[origem]
    if origem == "Plano" and val == 0:
        return CORTESIA
    if tab is None:
        return "Sem preço de tabela"
    if val < tab - 0.005:
        return "Desconto no plano" if origem == "Plano" else "Desconto na venda avulsa"
    if val > tab + 0.005:
        return "Acima da tabela"
    return "Sem diferença"


def coletar(desde=None, ate=None):
    tabela = {str(s["codServico"]): br(s["valor"]) for s in api("servico/listar")}
    planos = carregar_planos(tabela)
    for e in UNIDS:  # vencimento (validade) de cada plano, do relatório Sessões de Planos
        for x in historico("relatorios/sessoes_planos", "codEstab", e, HOJE, p_ini="dtVndIni", p_fim="dtVndFim"):
            p, v = planos.get(x["idPlano"]), data_br(x.get("dtValidade"))
            if p is not None and v:
                p["validade"] = max(v, p.get("validade") or v)
    linhas, agenda = [], collections.Counter()
    for e, unidade in UNIDS.items():
        ags = historico("relatorios/relatorio_atendimentos", "codEstab", e, ATE)
        if not ags:
            continue
        ini = min(data_br(a["dataAgendamento"]) for a in ags)
        vendas = historico("vendas_detalhado", "estab", e, ATE, desde=dt.date(ini.year, 3 * ((ini.month - 1) // 3) + 1, 1))
        por_dia = collections.defaultdict(list)
        for v in vendas:
            if v.get("data_venda"):
                por_dia[dt.date.fromisoformat(v["data_venda"][:10])].extend(
                    (v["cod_cliente"], i) for i in v["itens_venda"] or [] if i["tipo"] == "Serviço")
        for a in sorted(ags, key=lambda x: (data_br(x["dataAgendamento"]), x["horarioAgendamento"] or "")):
            d = data_br(a["dataAgendamento"])
            if (desde and d < desde) or (ate and d > ate):
                continue
            agenda[(unidade, d.replace(day=1), a["statusAgendamento"])] += 1
            if a["statusAgendamento"] != "Atendido":
                continue
            servico = (a["nomeServico"] or "").strip() or a["tipoAgendamento"] or ""
            tab, val, origem, obs = valorar(a, servico, planos, tabela, por_dia[d])
            orc = a["idOrcamento"] or ""
            p = planos.get(orc) if orc else None
            linhas.append(dict(
                unidade=unidade, data=d, mes=d.replace(day=1), cliente=nome(a["nomeCliente"]), servico=nome(servico),
                prof=nome(a["nomeProfissional"]), origem=origem, orc=orc, plano=p["nome"] if p else "",
                venda=data_br(p["raw"]["dataVenda"]) if p else None, vendedor=nome(p["raw"]["vendedor"]) if p else "",
                unit=p["unit"].get(str(a["codigoServico"] or "")) if p else None,
                tabela=tab, valor=round(val, 2), motivo=motivo(origem, tab, val),
                fonte=("Preço do plano" if "preço unitário do plano" in obs else
                       "Venda avulsa (bruto)" if "valor bruto da venda" in obs else
                       "Tabela atual" if tab is not None else "")))
    canonizar(linhas, "servico")
    return tabela, planos, linhas, agenda


def auditar_planos(planos, tabela, inicio, fim=ATE):
    """Planos Aprovados vendidos pelas unidades auditadas no período, com desconto, parcelas e saldo de sessões."""
    ordem = {u: i for i, u in enumerate(UNIDS)}
    out = []
    for orc, p in planos.items():
        r = p["raw"]
        venda = data_br(r["dataVenda"])
        if p["estab"] not in UNIDS or p["status"] != "Aprovado" or not venda or not inicio <= venda <= fim:
            continue
        servs = r["servicos"]
        parc, cota, validade = p["parcelas"], p["cota"], p.get("validade")  # parcelas da venda inteira
        vencs = [data_br(x["dataVencimento"]) for x in parc]
        aberto = [x for x in parc if x["confirmado"] != "Sim"]
        saldo = collections.Counter()
        for s in r.get("saldoPlano") or []:
            saldo[str(s["codigoServico"])] += int(br(s["qtdSessaoRestante"]))
        out.append(dict(
            estab=p["estab"], unidade=UNIDS[p["estab"]], orc=orc, venda=venda,
            cliente=nome((r["cliente"] or "").split(" - ", 1)[-1]), plano=p["nome"], tipo=r["tipoPlano"] or "",
            vendedor=nome(r["vendedor"]) or "(sem vendedor)",
            sessoes=sum(s["qtdSessoes"] or 0 for s in servs),
            cortesia=sum(s["qtdSessoes"] or 0 for s in servs if s["cortesia"] == "Sim"),
            tabela=round(sum((s["qtdSessoes"] or 0) * (tabela.get(str(s["codigoServico"])) or br(s["valorServico"]))
                             for s in servs), 2),
            cheio=round(sum(br(s["valorTotalServico"]) for s in servs), 2), final=br(r["precoFinal"]),
            validade=validade, meses=(validade.year - venda.year) * 12 + validade.month - venda.month if validade else None,
            venda_id=int(r["idVenda"]) if str(r.get("idVenda") or "").isdigit() else None,
            parcelas=len(parc), formas=", ".join(sorted({(x["formaPagamento"] or "").strip() for x in parc} - {""})),
            primeiro=min(vencs) if vencs else None, ultimo=max(vencs) if vencs else None,
            pago=round(cota * sum(br(x["valorLiquido"]) for x in parc if x["confirmado"] == "Sim"), 2),
            vencido=round(cota * sum(br(x["valorLiquido"]) for x in aberto if data_br(x["dataVencimento"]) < HOJE), 2),
            a_vencer=round(cota * sum(br(x["valorLiquido"]) for x in aberto if data_br(x["dataVencimento"]) >= HOJE), 2),
            restantes=sum(saldo.values()), a_realizar=round(sum(q * p["svc"].get(c, 0.0) for c, q in saldo.items()), 2)))
    canonizar(out, "vendedor")
    return sorted(out, key=lambda x: (ordem[x["estab"]], x["venda"], x["orc"]))


# ---------------------------------------------------------------- planilha
def titulo(ws, texto, sub):
    ws["A1"], ws["A2"] = texto, sub
    ws["A1"].font, ws["A2"].font = FT, FI


def secao(ws, r, texto):
    ws.cell(row=r, column=1, value=texto).font = FB


def linha(ws, r, valores, fmts, total=False, fill=None):
    for j, v in enumerate(valores, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = (FB if total else F), BOX
        if fmts[j - 1]:
            c.number_format = fmts[j - 1]
        if total:
            c.fill = TOT
        elif fill:
            c.fill = fill


def base(ws, cols, dados, fmts, nome_tabela, larg):
    """Aba de dados: cabeçalho na linha 1, uma linha por registro, tabela do Excel com filtro."""
    cabecalho(ws, 1, cols)
    for vals in dados:
        ws.append(vals)
    for row in ws.iter_rows(min_row=2, max_row=len(dados) + 1):
        for c, fmt in zip(row, fmts):
            if fmt and c.value is not None:
                c.number_format = fmt
    ultima = get_column_letter(len(cols))
    ws.add_table(Table(displayName=nome_tabela, ref=f"A1:{ultima}{max(len(dados) + 1, 2)}",
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))
    larguras(ws, larg)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def montar(planos_aud, linhas, agenda, inicio, destino, compactar=True, fim=ATE):
    wb = Workbook()
    unids = list(UNIDS.values())
    N, P = max(len(linhas), 1) + 1, max(len(planos_aud), 1) + 1  # última linha de cada base (nunca antes da 2)
    meses = sorted({l["mes"] for l in linhas})
    ag = sorted(agenda.items(), key=lambda x: (unids.index(x[0][0]), x[0][1]))
    ag_linhas = sorted({(u, m) for (u, m, _), _ in ag}, key=lambda x: (unids.index(x[0]), x[1]))
    G = max(len(ag_linhas), 1) + 1

    def A(c):
        return f"Atendimentos!${c}$2:${c}${N}"

    def PL(c):
        return f"Planos!${c}$2:${c}${P}"

    def AG(c):
        return f"Agenda!${c}$2:${c}${G}"

    fonte = f"Período: {inicio:%d/%m/%Y} a {fim:%d/%m/%Y} • Unidades: {', '.join(unids)} • Fonte: API Belle Software, extraído em {HOJE:%d/%m/%Y}"

    # ---------------- Resumo por Unidade
    rs = wb.active
    rs.title = "Resumo por Unidade"
    titulo(rs, "Drenesse – Auditoria de preços: valor de tabela x valor faturado", fonte)
    secao(rs, 4, "1. Atendimentos realizados – quanto seria cobrado pela tabela e quanto foi faturado")
    cabecalho(rs, 5, ["Unidade", "Atendimentos", "Se cobrado sem desconto (tabela, R$)", "Ganhou com o desconto (faturado, R$)",
                      "Diferença (faturado − tabela, R$)", "Deixou de ganhar (R$)", "% de desconto médio", "Ticket médio faturado (R$)"])
    fm1 = [None, NUM, BRL, BRL, BRL, BRL, PCT, BRL]
    r = 5
    for u in unids:
        r += 1
        linha(rs, r, [u, f"=COUNTIFS({A('A')},$A{r})", f"=SUMIFS({A('M')},{A('A')},$A{r})", f"=SUMIFS({A('N')},{A('A')},$A{r})",
                      f"=SUMIFS({A('O')},{A('A')},$A{r})", f"=SUMIFS({A('Q')},{A('A')},$A{r})",
                      f"=IFERROR(-E{r}/C{r},0)", f"=IFERROR(D{r}/B{r},0)"], fm1)
    r += 1
    linha(rs, r, ["TOTAL"] + [f"=SUM({c}6:{c}{r - 1})" for c in "BCDEF"] + [f"=IFERROR(-E{r}/C{r},0)", f"=IFERROR(D{r}/B{r},0)"],
          fm1, total=True)

    r += 2
    secao(rs, r, "2. De onde vem a diferença – por motivo (todas as unidades e deixou de ganhar por unidade)")
    r += 1
    rs.cell(row=r, column=8, value="Deixou de ganhar por unidade (R$)").font = FB
    rs.merge_cells(start_row=r, start_column=8, end_row=r, end_column=11)
    rs.cell(row=r, column=8).alignment = CENTER
    r += 1
    h2 = r
    cabecalho(rs, r, ["Motivo", "Atendimentos", "Valor de tabela (R$)", "Valor faturado (R$)", "Diferença (R$)",
                      "Deixou de ganhar (R$)", "% do total deixado de ganhar"] + unids)
    fm2 = [None, NUM, BRL, BRL, BRL, BRL, PCT, BRL, BRL, BRL, BRL]
    tot2 = h2 + len(MOTIVOS) + 1
    for m in MOTIVOS:
        r += 1
        linha(rs, r, [m, f"=COUNTIFS({A('R')},$A{r})", f"=SUMIFS({A('M')},{A('R')},$A{r})", f"=SUMIFS({A('N')},{A('R')},$A{r})",
                      f"=SUMIFS({A('O')},{A('R')},$A{r})", f"=SUMIFS({A('Q')},{A('R')},$A{r})", f"=IFERROR(F{r}/$F${tot2},0)"]
              + [f"=SUMIFS({A('Q')},{A('R')},$A{r},{A('A')},{get_column_letter(8 + k)}${h2})" for k in range(len(unids))], fm2)
    r += 1
    linha(rs, r, ["TOTAL"] + [f"=SUM({c}{h2 + 1}:{c}{r - 1})" for c in "BCDEFGHIJK"], fm2, total=True)

    r += 2
    secao(rs, r, "3. Planos aprovados vendidos no período – desconto sobre a tabela e vencimento após outubro/2027")
    r += 1
    h3 = r
    cabecalho(rs, r, ["Unidade", "Planos vendidos", "Valor de tabela (R$)", "Valor vendido (R$)", "Desconto concedido (R$)",
                      "Desconto médio", "Planos com desconto > 50% sobre a tabela", "% dos planos",
                      "Planos com desconto > 50% sobre o preço cheio registrado no plano",
                      "Planos com vencimento (validade) após out/2027", "Valor vendido nesses planos (R$)"])
    fm3 = [None, NUM, BRL, BRL, BRL, PCT, NUM, PCT, NUM, NUM, BRL]
    for u in unids:
        r += 1
        linha(rs, r, [u, f"=COUNTIFS({PL('A')},$A{r})", f"=SUMIFS({PL('J')},{PL('A')},$A{r})", f"=SUMIFS({PL('L')},{PL('A')},$A{r})",
                      f"=C{r}-D{r}", f"=IFERROR(E{r}/C{r},0)", f'=COUNTIFS({PL("A")},$A{r},{PL("O")},"SIM")',
                      f"=IFERROR(G{r}/B{r},0)", f'=COUNTIFS({PL("A")},$A{r},{PL("P")},">0.5")',
                      f'=COUNTIFS({PL("A")},$A{r},{PL("R")},"SIM")', f'=SUMIFS({PL("L")},{PL("A")},$A{r},{PL("R")},"SIM")'], fm3)
    r += 1
    linha(rs, r, ["TOTAL", f"=SUM(B{h3 + 1}:B{r - 1})", f"=SUM(C{h3 + 1}:C{r - 1})", f"=SUM(D{h3 + 1}:D{r - 1})", f"=C{r}-D{r}",
                  f"=IFERROR(E{r}/C{r},0)", f"=SUM(G{h3 + 1}:G{r - 1})", f"=IFERROR(G{r}/B{r},0)", f"=SUM(I{h3 + 1}:I{r - 1})",
                  f"=SUM(J{h3 + 1}:J{r - 1})", f"=SUM(K{h3 + 1}:K{r - 1})"], fm3, total=True)

    notas = [
        "COMO LER",
        "• Valor de tabela = preço atual do serviço na tabela do Belle. Se o serviço saiu da tabela, usa o preço unitário registrado no plano ou o valor bruto da venda avulsa (ver 'Fonte do valor de tabela').",
        "• Valor faturado = sessão de plano Aprovado: preço final do plano rateado pelas sessões (cortesia = R$ 0), descontada a parte do plano em parcelas vencidas e não pagas. Avulso = venda registrada no mesmo dia.",
        "• Plano Suspenso/Pendente, plano não localizado na API e avulso sem venda no dia entram com faturado R$ 0 – veja o motivo na seção 2 antes de tratar como desconto.",
        "• Diferença = faturado − tabela. Deixou de ganhar = quanto a mais entraria se a sessão fosse cobrada pela tabela (só diferenças positivas). Ganhou com o desconto = valor faturado.",
        "• Planos: só status Aprovado, vendidos pelas 4 unidades no período. Desconto = 1 − preço final ÷ (sessões vendidas × preço de tabela, incluindo cortesias).",
        "• Parcelas: a API registra as parcelas por venda. Quando a venda tem mais de um plano (ou outros itens), cada plano recebe a parte proporcional ao seu preço final.",
        "• Vencimento do plano = data de validade (relatório Sessões de Planos do Belle). Nenhuma parcela de pagamento vence depois de outubro/2027.",
        "• A tabela atual pode ser diferente da tabela vigente na data da venda. Para planos antigos, compare também com o 'Preço cheio no plano' (aba Planos).",
    ]
    r += 2
    for i, n in enumerate(notas):
        rs.cell(row=r + i, column=1, value=n).font = FB if i == 0 else F
    larguras(rs, [34, 13, 20, 20, 20, 18, 16, 16, 18, 16, 18])
    rs.row_dimensions[5].height = rs.row_dimensions[h2].height = 45
    rs.row_dimensions[h3].height = 60

    # ---------------- Mensal
    wm = wb.create_sheet("Mensal")
    titulo(wm, "Valor faturado por mês e por unidade", fonte)
    cabecalho(wm, 4, ["Mês"] + unids + ["Total faturado (R$)", "Valor de tabela (R$)", "Deixou de ganhar (R$)", "% de desconto",
                                        "Atendimentos", "Planos vendidos", "Valor vendido em planos (R$)",
                                        "Desconto dos planos vendidos"])
    fmm = [MES, BRL, BRL, BRL, BRL, BRL, BRL, BRL, PCT, NUM, NUM, BRL, PCT]

    def mes_crit(rng, cel):
        return f'{rng},">="&{cel},{rng},"<"&DATE(YEAR({cel}),MONTH({cel})+1,1)'

    r = 4
    for m in meses:
        r += 1
        crit = mes_crit(PL("C"), f"$A{r}")
        linha(wm, r, [m] + [f"=SUMIFS({A('N')},{A('C')},$A{r},{A('A')},{get_column_letter(2 + k)}$4)" for k in range(len(unids))]
              + [f"=SUM(B{r}:E{r})", f"=SUMIFS({A('M')},{A('C')},$A{r})", f"=SUMIFS({A('Q')},{A('C')},$A{r})",
                 f"=IFERROR(-SUMIFS({A('O')},{A('C')},$A{r})/G{r},0)", f"=COUNTIFS({A('C')},$A{r})",
                 f"=COUNTIFS({crit})", f"=SUMIFS({PL('L')},{crit})", f"=IFERROR(1-L{r}/SUMIFS({PL('J')},{crit}),0)"], fmm)
    r += 1
    tm = r
    linha(wm, r, ["TOTAL"] + [f"=SUM({c}5:{c}{r - 1})" for c in "BCDEFGH"] + [f"=IFERROR(-SUM({A('O')})/G{r},0)"]
          + [f"=SUM({c}5:{c}{r - 1})" for c in "JKL"] + [f"=IFERROR(1-L{r}/SUM({PL('J')}),0)"], fmm, total=True)
    r += 2
    secao(wm, r, "Por ano")
    r += 1
    ha = r
    cabecalho(wm, r, ["Ano"] + unids + ["Total faturado (R$)", "Valor de tabela (R$)", "Deixou de ganhar (R$)", "% de desconto",
                                        "Atendimentos", "Planos vendidos", "Valor vendido em planos (R$)",
                                        "Desconto dos planos vendidos"])
    fma = ["0"] + fmm[1:]
    for ano in sorted({m.year for m in meses}):
        r += 1
        ca = f'{A("C")},">="&DATE($A{r},1,1),{A("C")},"<"&DATE($A{r}+1,1,1)'
        cp = f'{PL("C")},">="&DATE($A{r},1,1),{PL("C")},"<"&DATE($A{r}+1,1,1)'
        linha(wm, r, [ano] + [f"=SUMIFS({A('N')},{ca},{A('A')},{get_column_letter(2 + k)}${ha})" for k in range(len(unids))]
              + [f"=SUM(B{r}:E{r})", f"=SUMIFS({A('M')},{ca})", f"=SUMIFS({A('Q')},{ca})", f"=IFERROR(-SUMIFS({A('O')},{ca})/G{r},0)",
                 f"=COUNTIFS({ca})", f"=COUNTIFS({cp})", f"=SUMIFS({PL('L')},{cp})", f"=IFERROR(1-L{r}/SUMIFS({PL('J')},{cp}),0)"], fma)
    larguras(wm, [11, 15, 15, 15, 15, 17, 17, 17, 11, 13, 11, 17, 14])
    wm.row_dimensions[4].height = wm.row_dimensions[ha].height = 45
    wm.freeze_panes = "B5"
    wm.cell(row=tm + 1, column=1, value="Planos vendidos = planos Aprovados pela data da venda; o faturado reconhece o plano aos poucos, conforme as sessões acontecem.").font = FI

    # ---------------- Serviços
    wsv = wb.create_sheet("Serviços")
    titulo(wsv, "Serviços realizados – quantidade, faturamento, ticket médio e desconto", fonte)
    cabecalho(wsv, 4, ["#", "Serviço", "Qtd realizada", "Valor faturado (R$)", "Ticket médio (R$)", "% do faturado",
                       "Valor de tabela (R$)", "Deixou de ganhar (R$)", "% de desconto"])
    fat = collections.Counter()
    for l in linhas:
        fat[l["servico"]] += l["valor"]
    servs = sorted({l["servico"] for l in linhas}, key=lambda s: (-fat[s], s))
    ts = 5 + len(servs)
    fms = [NUM, None, NUM, BRL, BRL, PCT, BRL, BRL, PCT]
    for i, s in enumerate(servs, 5):
        linha(wsv, i, [i - 4, s, f"=COUNTIFS({A('E')},B{i})", f"=SUMIFS({A('N')},{A('E')},B{i})", f"=IFERROR(D{i}/C{i},0)",
                       f"=IFERROR(D{i}/$D${ts},0)", f"=SUMIFS({A('M')},{A('E')},B{i})", f"=SUMIFS({A('Q')},{A('E')},B{i})",
                       f"=IFERROR(-SUMIFS({A('O')},{A('E')},B{i})/G{i},0)"], fms)
    linha(wsv, ts, ["", "TOTAL", f"=SUM(C5:C{ts - 1})", f"=SUM(D5:D{ts - 1})", f"=IFERROR(D{ts}/C{ts},0)", f"=SUM(F5:F{ts - 1})",
                    f"=SUM(G5:G{ts - 1})", f"=SUM(H5:H{ts - 1})", f"=IFERROR(-SUM({A('O')})/G{ts},0)"], fms, total=True)
    larguras(wsv, [6, 50, 13, 18, 15, 12, 18, 18, 12])
    wsv.freeze_panes = "C5"
    wsv.row_dimensions[4].height = 30

    # ---------------- Planos (base)
    wp = wb.create_sheet("Planos")
    cols_p = ["Unidade", "Nº plano", "Data da venda", "Cliente", "Plano", "Tipo", "Vendedor", "Sessões vendidas",
              "Sessões cortesia", "Valor de tabela (R$)", "Preço cheio no plano (R$)", "Preço final (R$)",
              "Desconto s/ tabela (R$)", "% desconto s/ tabela", "Desconto > 50%?", "% desconto s/ preço cheio do plano",
              "Vencimento do plano (validade)", "Vence após out/2027?", "Nº da venda", "Parcelas da venda",
              "Forma(s) de pagamento", "1º vencimento de parcela", "Último vencimento de parcela", "Pago (R$)",
              "Vencido e não pago (R$)", "A vencer (R$)", "Sessões restantes", "Valor das sessões a realizar (R$)",
              "Meses de validade"]
    fmp = [None, "0", DATA, None, None, None, None, NUM, NUM, BRL, BRL, BRL, BRL, PCT, None, PCT, DATA, None, "0", NUM, None,
           DATA, DATA, BRL, BRL, BRL, NUM, BRL, NUM]
    dados_p = []
    for i, p in enumerate(planos_aud, 2):
        dados_p.append([p["unidade"], p["orc"], p["venda"], p["cliente"], p["plano"], p["tipo"], p["vendedor"], p["sessoes"],
                        p["cortesia"], p["tabela"], p["cheio"], p["final"], f"=J{i}-L{i}", f'=IF(J{i}>0,1-L{i}/J{i},"")',
                        f'=IF(N{i}="","",IF(N{i}>0.5,"SIM","NÃO"))', f'=IF(K{i}>0,1-L{i}/K{i},"")', p["validade"],
                        f'=IF(Q{i}="","",IF(Q{i}>DATE({CORTE.year},{CORTE.month},{CORTE.day}),"SIM","NÃO"))', p["venda_id"],
                        p["parcelas"], p["formas"] or None, p["primeiro"], p["ultimo"], p["pago"], p["vencido"], p["a_vencer"],
                        p["restantes"], p["a_realizar"], p["meses"]])
    base(wp, cols_p, dados_p, fmp, "TabPlanos",
         [15, 11, 11, 32, 30, 14, 30, 9, 9, 15, 15, 15, 15, 11, 10, 11, 12, 10, 11, 9, 26, 11, 11, 14, 14, 14, 10, 15, 10])

    # ---------------- Vencimentos após out/2027 (validade do plano)
    wv = wb.create_sheet("Venc após out-2027")
    lista = sorted((p for p in planos_aud if p["validade"] and p["validade"] > CORTE), key=lambda p: (p["validade"], p["orc"]), reverse=True)
    ult_parcela = max((p["ultimo"] for p in planos_aud if p["ultimo"]), default=None)
    titulo(wv, f"Planos aprovados com vencimento (validade) depois de outubro/2027 – {len(lista)} planos",
           fonte + (f" • Nenhuma parcela de pagamento vence depois de out/2027 (a mais distante vence em {ult_parcela:%d/%m/%Y})."
                    if ult_parcela else ""))
    cabecalho(wv, 4, ["Unidade", "Nº plano", "Cliente", "Data da venda", "Vencimento do plano", "Meses de validade", "Plano",
                      "Vendedor", "Preço final (R$)", "Sessões vendidas", "Sessões restantes", "Valor das sessões a realizar (R$)",
                      "Pago (R$)", "Vencido e não pago (R$)"])
    fmv = [None, "0", None, DATA, DATA, NUM, None, None, BRL, NUM, NUM, BRL, BRL, BRL]
    r = 4
    for p in lista:
        r += 1
        linha(wv, r, [p["unidade"], p["orc"], p["cliente"], p["venda"], p["validade"], p["meses"], p["plano"], p["vendedor"],
                      p["final"], p["sessoes"], p["restantes"], p["a_realizar"], p["pago"], p["vencido"]], fmv)
    if not lista:
        r += 1
        wv.cell(row=r, column=1, value="Nenhum plano aprovado vence depois de outubro/2027 neste período.").font = FI
    r += 1
    linha(wv, r, ["TOTAL", f"=COUNT(B5:B{r - 1})", "", "", "", "", "", ""] + [f"=SUM({c}5:{c}{r - 1})" for c in "IJKLMN"],
          fmv, total=True)
    larguras(wv, [15, 11, 32, 11, 12, 10, 30, 30, 14, 10, 10, 15, 14, 14])
    wv.freeze_panes = "A5"
    wv.row_dimensions[4].height = 30

    # ---------------- Indicadores
    wi = wb.create_sheet("Indicadores")
    titulo(wi, "Indicadores para análise", fonte)
    r = 4
    secao(wi, r, "A. Composição do faturado por origem do valor (R$)")
    r += 1
    hA = r
    cabecalho(wi, r, ["Origem do valor"] + unids + ["Total (R$)", "% do faturado", "Atendimentos"])
    fmA = [None, BRL, BRL, BRL, BRL, BRL, PCT, NUM]
    tA = hA + len(ORIGENS) + 1
    for o in ORIGENS:
        r += 1
        linha(wi, r, [o] + [f"=SUMIFS({A('N')},{A('G')},$A{r},{A('A')},{get_column_letter(2 + k)}${hA})" for k in range(len(unids))]
              + [f"=SUM(B{r}:E{r})", f"=IFERROR(F{r}/$F${tA},0)", f"=COUNTIFS({A('G')},$A{r})"], fmA)
    r += 1
    linha(wi, r, ["TOTAL"] + [f"=SUM({c}{hA + 1}:{c}{r - 1})" for c in "BCDEFGH"], fmA, total=True)

    r += 2
    secao(wi, r, "B. Planos aprovados por faixa de desconto sobre a tabela")
    r += 1
    hB = r
    cabecalho(wi, r, ["Faixa de desconto", "Planos", "% dos planos", "Valor de tabela (R$)", "Valor vendido (R$)", "Desconto (R$)"])
    fmB = [None, NUM, PCT, BRL, BRL, BRL]
    tB = hB + len(FAIXAS) + 1
    for faixa, crits in FAIXAS:
        r += 1
        c = ",".join(f'{PL("N")},"{x}"' for x in crits)
        linha(wi, r, [faixa, f"=COUNTIFS({c})", f"=IFERROR(B{r}/$B${tB},0)", f"=SUMIFS({PL('J')},{c})", f"=SUMIFS({PL('L')},{c})",
                      f"=D{r}-E{r}"], fmB)
    r += 1
    linha(wi, r, ["TOTAL"] + [f"=SUM({c}{hB + 1}:{c}{r - 1})" for c in "BCDEF"], fmB, total=True)

    r += 2
    secao(wi, r, "C. Planos aprovados com preço final zero (sessões entregues sem cobrança)")
    r += 1
    hZ = r
    cabecalho(wi, r, ["Unidade", "Planos com preço zero", "% dos planos da unidade", "Sessões vendidas", "Valor de tabela (R$)",
                      "Sessões restantes"])
    fmZ = [None, NUM, PCT, NUM, BRL, NUM]
    for u in unids:
        r += 1
        cz = f"{PL('A')},$A{r},{PL('L')},0"
        linha(wi, r, [u, f"=COUNTIFS({cz})", f"=IFERROR(B{r}/COUNTIFS({PL('A')},$A{r}),0)", f"=SUMIFS({PL('H')},{cz})",
                      f"=SUMIFS({PL('J')},{cz})", f"=SUMIFS({PL('AA')},{cz})"], fmZ)
    r += 1
    linha(wi, r, ["TOTAL", f"=SUM(B{hZ + 1}:B{r - 1})", f"=IFERROR(B{r}/COUNTA({PL('A')}),0)", f"=SUM(D{hZ + 1}:D{r - 1})",
                  f"=SUM(E{hZ + 1}:E{r - 1})", f"=SUM(F{hZ + 1}:F{r - 1})"], fmZ, total=True)

    r += 2
    secao(wi, r, "D. Recebimentos, inadimplência e sessões ainda a realizar (planos aprovados vendidos no período)")
    r += 1
    hC = r
    cabecalho(wi, r, ["Unidade", "Valor vendido (R$)", "Pago (R$)", "Vencido e não pago (R$)", "% inadimplência", "A vencer (R$)",
                      "Sessões restantes", "Valor das sessões a realizar (R$)"])
    fmC = [None, BRL, BRL, BRL, PCT, BRL, NUM, BRL]
    for u in unids:
        r += 1
        linha(wi, r, [u, f"=SUMIFS({PL('L')},{PL('A')},$A{r})", f"=SUMIFS({PL('X')},{PL('A')},$A{r})",
                      f"=SUMIFS({PL('Y')},{PL('A')},$A{r})", f"=IFERROR(D{r}/B{r},0)", f"=SUMIFS({PL('Z')},{PL('A')},$A{r})",
                      f"=SUMIFS({PL('AA')},{PL('A')},$A{r})", f"=SUMIFS({PL('AB')},{PL('A')},$A{r})"], fmC)
    r += 1
    linha(wi, r, ["TOTAL", f"=SUM(B{hC + 1}:B{r - 1})", f"=SUM(C{hC + 1}:C{r - 1})", f"=SUM(D{hC + 1}:D{r - 1})",
                  f"=IFERROR(D{r}/B{r},0)", f"=SUM(F{hC + 1}:F{r - 1})", f"=SUM(G{hC + 1}:G{r - 1})", f"=SUM(H{hC + 1}:H{r - 1})"],
          fmC, total=True)

    r += 2
    secao(wi, r, "E. Sessões realizadas sem receita")
    r += 1
    hD = r
    cabecalho(wi, r, ["Unidade", "Sessões cortesia / 100% desconto", "Valor de tabela das cortesias (R$)", "Avulsos sem venda no dia",
                      "Valor de tabela dos avulsos sem venda (R$)", "Sessões de planos não aprovados", "% dos atendimentos sem receita"])
    fmD = [None, NUM, BRL, NUM, BRL, NUM, PCT]
    for u in unids:
        r += 1
        cond = f"{A('A')},$A{r},{A('R')}"
        linha(wi, r, [u, f'=COUNTIFS({cond},"{CORTESIA}")', f'=SUMIFS({A("M")},{cond},"{CORTESIA}")',
                      f'=COUNTIFS({cond},"Avulso sem venda no dia")', f'=SUMIFS({A("M")},{cond},"Avulso sem venda no dia")',
                      f'=COUNTIFS({cond},"Plano não aprovado")',
                      f'=IFERROR(COUNTIFS({A("A")},$A{r},{A("N")},0)/COUNTIFS({A("A")},$A{r}),0)'], fmD)
    r += 1
    linha(wi, r, ["TOTAL"] + [f"=SUM({c}{hD + 1}:{c}{r - 1})" for c in "BCDEF"] + [f'=IFERROR(COUNTIFS({A("N")},0)/COUNTA({A("A")}),0)'],
          fmD, total=True)

    r += 2
    secao(wi, r, "F. Agenda – faltas e desmarcações (todos os agendamentos do período)")
    r += 1
    hE = r
    cabecalho(wi, r, ["Unidade", "Agendamentos", "Atendidos", "Faltas", "Desmarcados", "Cancelados", "Taxa de faltas",
                      "Taxa de desmarcação/cancelamento"])
    fmE = [None, NUM, NUM, NUM, NUM, NUM, PCT, PCT]
    for u in unids:
        r += 1
        linha(wi, r, [u] + [f"=SUMIFS({AG(c)},{AG('A')},$A{r})" for c in "HCDEF"]
              + [f"=IFERROR(D{r}/(C{r}+D{r}),0)", f"=IFERROR((E{r}+F{r})/B{r},0)"], fmE)
    r += 1
    linha(wi, r, ["TOTAL"] + [f"=SUM({c}{hE + 1}:{c}{r - 1})" for c in "BCDEF"]
          + [f"=IFERROR(D{r}/(C{r}+D{r}),0)", f"=IFERROR((E{r}+F{r})/B{r},0)"], fmE, total=True)
    wi.cell(row=r + 1, column=1, value="Taxa de faltas = faltas ÷ (atendidos + faltas).").font = FI

    r += 3
    secao(wi, r, "G. Validade dos planos (meses entre a venda e o vencimento)")
    r += 1
    hV = r
    cabecalho(wi, r, ["Validade", "Planos", "% dos planos", "Valor vendido (R$)", "Sessões restantes"])
    fmVal = [None, NUM, PCT, BRL, NUM]
    faixas_val = [("Até 6 meses", ["<=6"]), ("7 a 12 meses", [">6", "<=12"]), ("13 a 18 meses", [">12", "<=18"]),
                  ("19 a 24 meses", [">18", "<=24"]), ("Mais de 24 meses", [">24"]), ("Sem validade informada", [""])]
    tV = hV + len(faixas_val) + 1
    for faixa, crits in faixas_val:
        r += 1
        c = ",".join(f'{PL("AC")},"{x}"' for x in crits)
        linha(wi, r, [faixa, f"=COUNTIFS({c})", f"=IFERROR(B{r}/$B${tV},0)", f"=SUMIFS({PL('L')},{c})", f"=SUMIFS({PL('AA')},{c})"], fmVal)
    r += 1
    linha(wi, r, ["TOTAL"] + [f"=SUM({c}{hV + 1}:{c}{r - 1})" for c in "BCDE"], fmVal, total=True)

    r += 2
    secao(wi, r, "H. Desconto por vendedor (planos aprovados vendidos no período)")
    r += 1
    cabecalho(wi, r, ["Vendedor", "Planos", "Valor de tabela (R$)", "Valor vendido (R$)", "Desconto médio", "Planos com desconto > 50%",
                      "% dos planos > 50%", "Planos com preço zero"])
    fmF = [None, NUM, BRL, BRL, PCT, NUM, PCT, NUM]
    qtd_vend = collections.Counter(p["vendedor"] for p in planos_aud)
    for v, _ in sorted(qtd_vend.items(), key=lambda x: (-x[1], x[0])):
        r += 1
        linha(wi, r, [v, f"=COUNTIFS({PL('G')},$A{r})", f"=SUMIFS({PL('J')},{PL('G')},$A{r})", f"=SUMIFS({PL('L')},{PL('G')},$A{r})",
                      f"=IFERROR(1-D{r}/C{r},0)", f'=COUNTIFS({PL("G")},$A{r},{PL("O")},"SIM")', f"=IFERROR(F{r}/B{r},0)",
                      f"=COUNTIFS({PL('G')},$A{r},{PL('L')},0)"], fmF)

    r += 2
    secao(wi, r, "I. Os 20 planos com maior desconto em R$ (links para a aba Planos)")
    r += 1
    cabecalho(wi, r, ["Unidade", "Nº plano", "Data da venda", "Cliente", "Vendedor", "Valor de tabela (R$)", "Preço final (R$)",
                      "Desconto (R$)", "% desconto"])
    fmG = [None, "0", DATA, None, None, BRL, BRL, BRL, PCT]
    top = sorted(range(len(planos_aud)), key=lambda i: -(planos_aud[i]["tabela"] - planos_aud[i]["final"]))[:20]
    for i in top:
        r += 1
        pr = i + 2
        linha(wi, r, [f"=Planos!{c}{pr}" for c in "ABCDGJLMN"], fmG)

    r += 2
    secao(wi, r, "J. Os 15 serviços que mais deixaram de ganhar")
    r += 1
    cabecalho(wi, r, ["Serviço", "Qtd realizada", "Valor de tabela (R$)", "Valor faturado (R$)", "Deixou de ganhar (R$)", "% de desconto"])
    fmH = [None, NUM, BRL, BRL, BRL, PCT]
    perda = collections.Counter()
    for l in linhas:
        if l["tabela"] is not None:
            perda[l["servico"]] += max(0.0, l["tabela"] - l["valor"])
    for s, _ in perda.most_common(15):
        r += 1
        linha(wi, r, [s, f"=COUNTIFS({A('E')},$A{r})", f"=SUMIFS({A('M')},{A('E')},$A{r})", f"=SUMIFS({A('N')},{A('E')},$A{r})",
                      f"=SUMIFS({A('Q')},{A('E')},$A{r})", f"=IFERROR(-SUMIFS({A('O')},{A('E')},$A{r})/C{r},0)"], fmH)
    larguras(wi, [36, 15, 17, 17, 17, 17, 15, 15, 15])
    for rr in (hA, hB, hZ, hC, hD, hE, hV):
        wi.row_dimensions[rr].height = 45

    # ---------------- Atendimentos (base)
    wa = wb.create_sheet("Atendimentos")
    cols_a = ["Unidade", "Data", "Mês", "Cliente", "Serviço", "Profissional", "Origem do valor", "Nº plano", "Plano",
              "Venda do plano", "Vendedor do plano", "Preço unit. no plano (R$)", "Valor de tabela (R$)", "Valor faturado (R$)",
              "Diferença (R$)", "% desconto", "Deixou de ganhar (R$)", "Motivo da diferença", "Fonte do valor de tabela"]
    fma_ = [None, DATA, MES, None, None, None, None, "0", None, DATA, None, BRL, BRL, BRL, BRL, PCT, BRL, None, None]
    dados_a = []
    for i, l in enumerate(linhas, 2):
        dados_a.append([l["unidade"], l["data"], l["mes"], l["cliente"], l["servico"], l["prof"] or None, l["origem"],
                        l["orc"] or None, l["plano"] or None, l["venda"], l["vendedor"] or None, l["unit"], l["tabela"], l["valor"],
                        f'=IF(M{i}="","",N{i}-M{i})', f'=IF(OR(M{i}="",M{i}=0),"",1-N{i}/M{i})',
                        f'=IF(M{i}="","",MAX(0,M{i}-N{i}))', l["motivo"], l["fonte"] or None])
    base(wa, cols_a, dados_a, fma_, "TabAtendimentos",
         [15, 11, 10, 32, 36, 30, 17, 11, 26, 11, 26, 13, 13, 13, 13, 10, 13, 28, 18])

    # ---------------- Agenda (base)
    wg = wb.create_sheet("Agenda")
    cont = collections.defaultdict(collections.Counter)
    for (u, m, st), q in ag:
        cont[(u, m)][st] += q
    dados_g = []
    for i, (u, m) in enumerate(ag_linhas, 2):
        c = cont[(u, m)]
        outros = sum(q for st, q in c.items() if st not in ("Atendido", "Falhou", "Desmarcado", "Cancelado"))
        dados_g.append([u, m, c["Atendido"], c["Falhou"], c["Desmarcado"], c["Cancelado"], outros, f"=SUM(C{i}:G{i})"])
    base(wg, ["Unidade", "Mês", "Atendidos", "Faltas", "Desmarcados", "Cancelados", "Outros status", "Total de agendamentos"],
         dados_g, [None, MES, NUM, NUM, NUM, NUM, NUM, NUM], "TabAgenda", [15, 11, 11, 11, 12, 11, 12, 14])

    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    if compactar:
        compactar_xlsx(destino, {"Atendimentos": {"O", "P", "Q"}, "Planos": {"M", "N", "O", "P", "R"}, "Agenda": {"H"}})


def compactar_xlsx(caminho, formulas_por_aba):
    """Reescreve o .xlsx gerado pelo openpyxl de forma mais enxuta, sem mudar o conteúdo: os textos viram strings
    compartilhadas (o openpyxl grava cada texto por extenso em cada célula), as fórmulas que se repetem linha a
    linha viram fórmulas compartilhadas do Excel e o zip é recomprimido no nível máximo."""
    with zipfile.ZipFile(caminho) as z:
        ordem = [i.filename for i in z.infolist()]
        partes = {n: z.read(n) for n in ordem}
    rels = partes["xl/_rels/workbook.xml.rels"].decode("utf-8")
    alvo = {}
    for rel in re.findall(r"<Relationship\b[^>]*>", rels):
        destino_rel = re.search(r'\bTarget="([^"]+)"', rel).group(1)
        alvo[re.search(r'\bId="([^"]+)"', rel).group(1)] = destino_rel.lstrip("/") if destino_rel.startswith("/") else "xl/" + destino_rel
    abas = {}
    for sh in re.findall(r"<sheet\b[^>]*>", partes["xl/workbook.xml"].decode("utf-8")):
        abas[html.unescape(re.search(r'\bname="([^"]*)"', sh).group(1))] = alvo[re.search(r'\br:id="([^"]+)"', sh).group(1)]

    cel_f = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*)><f>([^<]*)</f>(<v\s*/>|<v></v>)?</c>')
    for aba, colunas in formulas_por_aba.items():
        xml = partes[abas[aba]].decode("utf-8")

        def modelo(m):  # fórmula com o número da própria linha trocado por #
            return re.sub(rf"(?<=[A-Z]){m.group(2)}(?!\d)", "#", m.group(4))
        mestre = {}  # coluna -> [linha mestre, última linha, modelo, si]
        for m in cel_f.finditer(xml):
            c = m.group(1)
            if c not in colunas:
                continue
            if c not in mestre:
                mestre[c] = [int(m.group(2)), int(m.group(2)), modelo(m), len(mestre)]
            elif modelo(m) == mestre[c][2]:
                mestre[c][1] = int(m.group(2))

        def troca_f(m):
            c, lin = m.group(1), int(m.group(2))
            if c not in mestre or modelo(m) != mestre[c][2]:
                return m.group(0)
            ini, fim, _, si = mestre[c]
            v = m.group(5) or ""
            if lin == ini:
                return f'<c r="{c}{lin}"{m.group(3)}><f t="shared" ref="{c}{ini}:{c}{fim}" si="{si}">{m.group(4)}</f>{v}</c>'
            return f'<c r="{c}{lin}"{m.group(3)}><f t="shared" si="{si}"/>{v}</c>'
        partes[abas[aba]] = cel_f.sub(troca_f, xml).encode("utf-8")

    textos, indice, usos = [], {}, 0
    cel_s = re.compile(r'<c r="([A-Z]+\d+)"([^>]*?) t="inlineStr"><is><t( xml:space="preserve")?>(.*?)</t></is></c>', re.S)

    def troca_s(m):
        nonlocal usos
        chave = (m.group(3) or "", m.group(4))
        if chave not in indice:
            indice[chave] = len(textos)
            textos.append(chave)
        usos += 1
        return f'<c r="{m.group(1)}"{m.group(2)} t="s"><v>{indice[chave]}</v></c>'
    for parte in abas.values():
        partes[parte] = cel_s.sub(troca_s, partes[parte].decode("utf-8")).encode("utf-8")
    partes["xl/sharedStrings.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{usos}" uniqueCount="{len(textos)}">'
        + "".join(f"<si><t{p}>{t}</t></si>" for p, t in textos) + "</sst>").encode("utf-8")
    tipos = partes["[Content_Types].xml"].decode("utf-8")
    if "/xl/sharedStrings.xml" not in tipos:
        tipos = tipos.replace("</Types>", '<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
                              'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    partes["[Content_Types].xml"] = tipos.encode("utf-8")
    if "sharedStrings" not in rels:
        novo_id = next(f"rId{i}" for i in range(len(alvo) + 1, len(alvo) + 100) if f"rId{i}" not in alvo)
        rels = rels.replace("</Relationships>", f'<Relationship Id="{novo_id}" Type="http://schemas.openxmlformats.org/'
                            'officeDocument/2006/relationships/sharedStrings" Target="/xl/sharedStrings.xml"/></Relationships>')
    partes["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
    temp = caminho + ".tmp"
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for n in ordem + (["xl/sharedStrings.xml"] if "xl/sharedStrings.xml" not in ordem else []):
            z.writestr(n, partes[n])
    os.replace(temp, caminho)


if __name__ == "__main__":
    desde = dt.datetime.strptime(os.environ["BELLE_DESDE"], "%d/%m/%Y").date() if os.environ.get("BELLE_DESDE") else None
    ate = dt.datetime.strptime(os.environ["BELLE_ATE"], "%d/%m/%Y").date() if os.environ.get("BELLE_ATE") else None
    tabela, planos, linhas, agenda = coletar(desde, ate)
    inicio = desde or min(l["data"] for l in linhas).replace(day=1)
    fim = ate or ATE
    planos_aud = auditar_planos(planos, tabela, inicio, fim)
    out = os.environ.get("BELLE_OUT", "Auditoria_Precos_Drenesse.xlsx")
    montar(planos_aud, linhas, agenda, inicio, out, fim=fim)
    print(f"{len(linhas)} atendimentos, {len(planos_aud)} planos aprovados -> {out}")
