"""Vendas de avaliação x Cabine SDR x Cabine – API Belle Software.

Uso: python3 belle_avaliacao.py [dd/mm/aaaa início] [dd/mm/aaaa fim] [codEstab ...]
     padrão: 01/01/2026 a 26/09/2026, unidades Petrópolis, Lagoa Nova, Capim Macio e Norte Shopping.
     Gera Avaliacao_x_Cabine_Drenesse.xlsx (ou BELLE_OUT).

Regras (definidas pela Drenesse):
- Sessão de avaliação = agendamento do tipo "Avaliação" (ou o serviço AVALIAÇÃO ESTÉTICA).
- Sessão experimental = agendamento avulso (sem plano) de um dos serviços de entrada de cliente novo (EXPERIMENTAL).
- Avaliação = venda feita no mesmo dia de uma sessão de avaliação com status Atendido (com ou sem sessão experimental).
- Cabine SDR = venda no dia de uma sessão experimental atendida quando a avaliação do dia teve falta ou foi desmarcada,
  ou quando só houve a sessão experimental e o cliente nunca teve plano.
- Cabine = qualquer outra venda (inclusive sessão experimental de quem já tinha plano).
Comparecimento = atendidos ÷ (atendidos + faltas). Fechamento da avaliação e da Cabine SDR = sessões atendidas
(cliente/dia) com venda de plano no dia ÷ sessões atendidas; da Cabine = clientes atendidos na cabine no mês que
compraram plano de cabine no mesmo mês ÷ clientes atendidos na cabine no mês.
"""
import collections
import datetime as dt
import os
import sys

from openpyxl import Workbook

from belle_auditoria import DATA, MES, NUM, PCT, base, compactar_xlsx, data_br, linha, nome, secao, titulo
from belle_sabados import BRL, FB, FI, HOJE, UNIDADES, br, cabecalho, carregar_planos, historico, larguras

# Serviços com que as sessões experimentais (agendamentos avulsos de clientes novos) são lançadas no Belle.
EXPERIMENTAL = {
    22: "DRENAGEM MÉTODO DRENESSE",
    56210744: "SESSÃO EXPERIMENTAL DRENAGEM - MÉTODO DRENESSE",
    56260425: "DRENAGEM MÉTODO DRENESSE 98,70",
    33353403: "LIMPEZA DE PELE PREMIUM",
    56210746: "SESSÃO EXPERIMENTAL - STIMULUS",
    56210745: "SESSÃO EXPERIMENTAL - DIÁSTASE",
}
SERV_AVALIACAO = {52: "AVALIAÇÃO ESTÉTICA"}  # avaliação lançada como serviço (mesmas avaliadoras do tipo "Avaliação")
ESTABS = [1, 2, 3, 6]  # unidade Laser (4) fora da análise
ANTES = 90  # dias antes do início buscados só para achar a avaliação de vendas do começo do período
AV, SDR, CAB, FORA = "Avaliação", "Cabine SDR", "Cabine", "Fora da análise"
T_AV, T_CAB, T_OUT = "Avaliação", "Sessão de cabine", "Retorno, consulta ou sem serviço"
T_EXP = "Sessão experimental"  # provisório: vira um dos três abaixo conforme o dia e o histórico do cliente
T_EXP_AV = "Sessão experimental no dia da avaliação atendida"
T_SDR = "Sessão experimental SDR"
T_EXP_REC = "Sessão experimental de quem já tinha plano"
GRUPO = {T_AV: AV, T_EXP_AV: AV, T_SDR: SDR, T_EXP_REC: CAB, T_CAB: CAB, T_OUT: FORA}
STATUS = {"Atendido": "Atendido", "Falhou": "Falta", "Desmarcado": "Desmarcado", "Cancelado": "Cancelado"}
NAO_TINHA = "Não tinha"
P_AV = "Comprou no dia da avaliação atendida"
P_AV_FALTA = "Avaliação do dia não atendida – comprou na sessão experimental"
P_EXP_NOVO = "Só sessão experimental no dia – nunca teve plano"
P_EXP_REC = "Só sessão experimental no dia – já tinha plano"
P_DEPOIS = "Fez avaliação em outro dia – primeira compra"
P_REC = "Cliente que já tinha plano"
P_SEM = "Sem avaliação nem sessão experimental no dia – primeira compra"
PERFIS = [(AV, P_AV), (SDR, P_AV_FALTA), (SDR, P_EXP_NOVO), (CAB, P_EXP_REC), (CAB, P_DEPOIS), (CAB, P_REC), (CAB, P_SEM)]


def codigo(x):
    s = str(x if x is not None else "").strip()
    return int(s) if s.isdigit() else s


def campo(d, *chaves):
    return next((d[k] for k in chaves if d.get(k) not in (None, "")), None)


def reais(v):
    return f"R$ {v:,.2f}".replace(",", "x").replace(".", ",").replace("x", ".")


def meses(ini, fim):
    m, out = ini.replace(day=1), []
    while m <= fim:
        out.append(m)
        m = (m + dt.timedelta(days=32)).replace(day=1)
    return out


def tipo_sessao(a, sc, orc):
    if (a.get("tipoAgendamento") or "") == "Avaliação" or sc in SERV_AVALIACAO:
        return T_AV
    if sc in ("", 0):
        return T_OUT
    return T_EXP if sc in EXPERIMENTAL and not orc else T_CAB


def coletar(ini, fim, estabs):
    desde = ini - dt.timedelta(days=ANTES)
    desde = dt.date(desde.year, 3 * ((desde.month - 1) // 3) + 1, 1)  # janelas trimestrais da API
    planos = carregar_planos()
    nomes, compras, vendas = {}, collections.defaultdict(list), []
    for orc, p in planos.items():
        r = p["raw"]
        cliente = r.get("cliente") or ""
        c = codigo(campo(r, "codigoCliente", "codCliente") or cliente.split(" - ", 1)[0])
        nomes.setdefault(c, nome(cliente.split(" - ", 1)[-1]))
        d = data_br(r.get("dataVenda"))
        if not d:
            continue
        if p["status"] == "Aprovado":  # histórico completo de compras do cliente, em todas as unidades
            compras[c].append((d, p["nome"]))
        if p["estab"] in estabs and ini <= d <= fim:
            vendas.append(dict(unidade=UNIDADES[p["estab"]], data=d, cliente_cod=c, cliente=nomes[c], tipo="Plano",
                               desc=nome(p["nome"]), orc=orc, status=p["status"], vendedor=nome(r.get("vendedor")),
                               valor=br(r.get("precoFinal")), conta=p["status"] == "Aprovado",
                               av_belle=codigo(r.get("avaliacao")) or None,
                               obs="" if p["status"] == "Aprovado" else f"Plano {p['status']} – fora do faturamento"))
    for lst in compras.values():
        lst.sort()

    agenda, por_id = [], {}
    for e in estabs:
        for a in historico("relatorios/relatorio_atendimentos", "codEstab", e, fim, desde=desde):
            d, c = data_br(a.get("dataAgendamento")), codigo(a.get("codigoCliente"))
            if not d or not desde <= d <= fim or c in ("", 0):
                continue
            sc, orc = codigo(a.get("codigoServico")), a.get("idOrcamento") or ""
            t = tipo_sessao(a, sc, orc)
            nomes.setdefault(c, nome(a.get("nomeCliente")))
            st = (a.get("statusAgendamento") or "").strip()
            x = dict(id=codigo(a.get("idAgendamento")), unidade=UNIDADES[e], data=d, mes=d.replace(day=1),
                     hora=a.get("horarioAgendamento") or "", cliente_cod=c, cliente=nomes[c], servico_cod=sc,
                     servico=nome(a.get("nomeServico")) or nome(a.get("tipoAgendamento")), orc=orc, status_belle=st,
                     status=STATUS.get(st, "Outros"), prof=nome(a.get("nomeProfissional")), tipo=t)
            agenda.append(x)
            por_id[x["id"]] = x
    agenda.sort(key=lambda a: (a["data"], a["hora"], a["unidade"], a["cliente"]))

    # A sessão experimental pertence ao grupo da venda daquele dia: Avaliação se houve avaliação atendida; Cabine SDR se a
    # avaliação do dia não foi atendida ou se o cliente nunca teve plano; Cabine se ele já tinha plano.
    # Avaliação realizada = cliente + dia com avaliação atendida; sessão SDR realizada = cliente + dia com experimental SDR
    # atendida.
    dia = collections.defaultdict(list)
    for a in agenda:
        dia[(a["cliente_cod"], a["data"])].append(a)
    av_vis, exp_vis, exp_rec = {}, {}, set()
    for (c, d), lst in dia.items():
        avs = [a for a in lst if a["tipo"] == T_AV]
        atend = [a for a in avs if a["status"] == "Atendido"]
        tinha = any(x < d for x, _ in compras.get(c, []))
        for a in lst:
            if a["tipo"] == T_EXP:
                a["tipo"] = T_EXP_AV if atend else T_SDR if (avs or not tinha) else T_EXP_REC
        for a in lst:
            a["grupo"] = GRUPO[a["tipo"]]
        exps = [a for a in lst if a["tipo"] in (T_EXP_AV, T_SDR) and a["status"] == "Atendido"]
        if atend:
            av_vis[(c, d)] = dict(atend[0], exp=", ".join(a["servico"] for a in exps), vendas=[])
        elif exps:
            st_av = ", ".join(sorted({a["status_belle"] for a in avs})) or NAO_TINHA
            exp_vis[(c, d)] = dict(exps[0], outros=", ".join(a["servico"] for a in exps[1:]), av_status=st_av, vendas=[])
        if any(a["tipo"] == T_EXP_REC and a["status"] == "Atendido" for a in lst):
            exp_rec.add((c, d))
    av_cliente = collections.defaultdict(list)
    for v in sorted(av_vis.values(), key=lambda v: v["data"]):
        av_cliente[v["cliente_cod"]].append(v)

    def classificar(c, d):
        """Grupo, perfil e sessão do dia (avaliação ou experimental) de uma venda ao cliente c no dia d."""
        antes = [x for x, _ in compras.get(c, []) if x < d]
        if (c, d) in av_vis:
            return AV, P_AV, av_vis[(c, d)], None
        ex = exp_vis.get((c, d))
        if ex:
            return SDR, (P_EXP_NOVO if ex["av_status"] == NAO_TINHA else P_AV_FALTA), ex, None
        if (c, d) in exp_rec:
            return CAB, P_EXP_REC, None, None
        av = next((v for v in reversed(av_cliente.get(c, [])) if v["data"] < d), None)
        if av and not antes:
            return CAB, P_DEPOIS, None, av
        return CAB, (P_REC if antes else P_SEM), None, None

    tipos = collections.Counter()
    ids_planos = {str(p["raw"]["idVenda"]) for p in planos.values() if p["raw"].get("idVenda")}
    for e in estabs:
        for v in historico("vendas_detalhado", "estab", e, fim, desde=ini):
            d = dt.date.fromisoformat(v["data_venda"][:10]) if v.get("data_venda") else None
            if not d or not ini <= d <= fim or "cancel" in str(campo(v, "status", "situacao") or "").casefold():
                continue
            c = codigo(v.get("cod_cliente"))
            idv = campo(v, "id_venda", "idVenda", "cod_venda", "codigo_venda", "codVenda")
            com_plano = idv is not None and str(idv) in ids_planos
            for i in v.get("itens_venda") or []:
                tipo = (i.get("tipo") or "").strip()
                tipos[tipo or "(sem tipo)"] += 1
                if "plano" in tipo.casefold() or "pacote" in tipo.casefold():
                    continue  # o plano já entra pelo relatório de planos, com status e preço final
                if "cancel" in str(i.get("status") or "").casefold():
                    continue
                repetido = com_plano and tipo == "Serviço"
                vendas.append(dict(unidade=UNIDADES[e], data=d, cliente_cod=c,
                                   cliente=nomes.get(c) or nome((v.get("dados_cliente") or {}).get("nome")),
                                   tipo="Serviço avulso" if tipo == "Serviço" else (tipo or "Item"), desc=nome(i.get("desc_item")),
                                   orc="", status="", vendedor=nome(campo(v, "vendedor", "nome_vendedor", "nomeVendedor")),
                                   valor=br(i.get("valor_liquido")), conta=not repetido, av_belle=None,
                                   obs="Serviço da mesma venda de um plano – já contado no preço do plano" if repetido else ""))
    vendas.sort(key=lambda x: (x["data"], x["unidade"], x["cliente"], x["tipo"] != "Plano"))
    for x in vendas:
        x["grupo"], x["perfil"], sessao, av_antes = classificar(x["cliente_cod"], x["data"])
        x["exp"] = any(a["tipo"] in (T_EXP_AV, T_SDR, T_EXP_REC) and a["status"] == "Atendido"
                       for a in dia.get((x["cliente_cod"], x["data"]), []))
        x["ja_cliente"] = any(d < x["data"] for d, _ in compras.get(x["cliente_cod"], []))
        x["av_data"] = av_antes["data"] if av_antes else None
        x["dias"] = (x["data"] - av_antes["data"]).days if av_antes else None
        if sessao is not None and x["tipo"] == "Plano" and x["conta"]:
            sessao["vendas"].append(x)

    for v in list(av_vis.values()) + list(exp_vis.values()):
        cmp = compras.get(v["cliente_cod"], [])
        v["ja_cliente"] = any(x < v["data"] for x, _ in cmp)
        prox = [x for x in cmp if v["data"] <= x[0] <= fim]
        v["dias_compra"] = (prox[0][0] - v["data"]).days if prox else None
        v["primeira"] = ", ".join(n for x, n in prox if x == prox[0][0]) if prox else ""

    periodo = [a for a in agenda if a["data"] >= ini]
    cab_mes = collections.Counter((a["unidade"], a["mes"], a["cliente_cod"]) for a in periodo
                                  if a["grupo"] == CAB and a["status"] == "Atendido")
    compra_cab = collections.defaultdict(list)
    for x in vendas:
        if x["grupo"] == CAB and x["tipo"] == "Plano" and x["conta"]:
            compra_cab[(x["cliente_cod"], x["data"].replace(day=1))].append(x["valor"])
    resumo = collections.Counter((a["unidade"], a["mes"], a["tipo"], a["status"]) for a in periodo)

    # Conferência com o vínculo plano → atendimento que o próprio Belle guarda no campo "avaliação" do plano (preenchido só
    # em parte dos planos, e às vezes apontando para sessões que não são de avaliação, como acompanhamento estético).
    ligados = [x for x in vendas if x["tipo"] == "Plano" and x["conta"] and x["av_belle"] in por_id]
    lig_av = [x for x in ligados if por_id[x["av_belle"]]["tipo"] == T_AV and por_id[x["av_belle"]]["data"] == x["data"]
              and por_id[x["av_belle"]]["status"] == "Atendido"]
    diag = dict(
        tipos=tipos, repetidos=sum(1 for x in vendas if not x["conta"] and x["tipo"] != "Plano"),
        fora=collections.Counter(a["servico"] or "(sem serviço)" for a in periodo if a["tipo"] == T_OUT),
        exp_plano=sum(1 for a in periodo if a["servico_cod"] in EXPERIMENTAL and a["tipo"] == T_CAB),
        serv52=sum(1 for a in periodo if a["servico_cod"] in SERV_AVALIACAO),
        ligados=len(ligados), lig_av=len(lig_av), lig_av_ok=sum(1 for x in lig_av if x["grupo"] == AV),
        lig_outros=collections.Counter(por_id[x["av_belle"]]["servico"] or por_id[x["av_belle"]]["tipo"]
                                       for x in ligados if x not in lig_av))
    return dict(vendas=vendas,
                av_vis=sorted((v for v in av_vis.values() if v["data"] >= ini), key=lambda v: (v["data"], v["hora"], v["unidade"])),
                exp_vis=sorted((v for v in exp_vis.values() if v["data"] >= ini), key=lambda v: (v["data"], v["hora"], v["unidade"])),
                agenda_det=[a for a in periodo if a["tipo"] in (T_AV, T_EXP_AV, T_SDR, T_EXP_REC)], cab_mes=cab_mes, compra_cab=compra_cab,
                resumo=resumo, diag=diag, nomes=nomes)


# ---------------------------------------------------------------- planilha
ABAS = {"V": "Vendas", "AV": "Avaliações", "EX": "Sessões SDR", "AA": "Agenda avaliação-experimental",
        "AR": "Agenda resumo", "CC": "Clientes cabine"}
DIM = {"unidade": {"AR": "A", "AV": "A", "EX": "A", "CC": "A", "V": "A"},
       "mes": {"AR": "B", "AV": "C", "EX": "C", "CC": "B", "V": "C"}}
# Como cada grupo é lido nas abas de base: agenda (coluna, valor), base e "fechou" do fechamento, filtros das vendas.
# O comparecimento da Avaliação conta só as sessões de avaliação (não as experimentais feitas no mesmo dia).
GR = {AV: dict(ar=("D", T_AV), base=("AV", "A"), fech=("AV", "J"), vendas=[("L", AV)]),
      SDR: dict(ar=("C", SDR), base=("EX", "A"), fech=("EX", "L"), vendas=[("L", SDR)]),
      CAB: dict(ar=("C", CAB), base=("CC", "A"), fech=("CC", "F"), vendas=[("L", CAB)])}


def exemplos(d, fim):
    """Casos reais de cada situação, para mostrar como cada atendimento e cada venda foi classificado."""
    def rec(lst, n):  # os mais recentes
        return list(reversed(lst))[:n]
    av, ex, ag, vd = d["av_vis"], d["exp_vis"], d["agenda_det"], d["vendas"]
    out = []

    def add(caso, unidade, data, cliente, sessoes, situacao, grupo, porque):
        out.append([caso, unidade, data, cliente, sessoes, situacao, grupo, porque])

    def planos(v):
        return f"{', '.join(x['desc'] for x in v['vendas'])} ({reais(sum(x['valor'] for x in v['vendas']))})"

    for v in rec([v for v in av if v["vendas"] and v["exp"]], 2):
        add("Avaliação + sessão experimental, comprou", v["unidade"], v["data"], v["cliente"],
            f"Avaliação ({v['prof']}) + {v['exp']}", f"Avaliação atendida • {planos(v)}", AV,
            "Teve sessão de avaliação com status Atendido no dia da compra.")
    for v in rec([v for v in av if v["vendas"] and not v["exp"]], 2):
        add("Só avaliação, comprou", v["unidade"], v["data"], v["cliente"], f"Avaliação ({v['prof']})",
            f"Avaliação atendida • {planos(v)}", AV, "Sessão de avaliação atendida e compra no mesmo dia.")
    for v in rec([v for v in ex if v["vendas"] and v["av_status"] != NAO_TINHA], 2):
        add("Avaliação não atendida + experimental, comprou", v["unidade"], v["data"], v["cliente"],
            f"Avaliação ({v['av_status']}) + {v['servico']}", f"Experimental atendida • {planos(v)}", SDR,
            f"A sessão de avaliação do dia ficou como '{v['av_status']}' e a experimental foi atendida.")
    for v in rec([v for v in ex if v["vendas"] and v["av_status"] == NAO_TINHA], 2):
        add("Só sessão experimental, nunca teve plano, comprou", v["unidade"], v["data"], v["cliente"], v["servico"],
            f"Experimental atendida • {planos(v)}", SDR, "Sem avaliação no dia e nenhum plano aprovado antes.")
    for x in rec([x for x in vd if x["perfil"] == P_EXP_REC and x["tipo"] == "Plano" and x["conta"]], 2):
        add("Só sessão experimental, já tinha plano, comprou", x["unidade"], x["data"], x["cliente"], x["desc"],
            f"Plano {reais(x['valor'])}", CAB, "Sessão experimental sem avaliação no dia, mas a cliente já tinha plano aprovado.")
    for v in rec([v for v in av if not v["vendas"] and v["dias_compra"] is None and v["data"] <= fim - dt.timedelta(days=30)], 2):
        add("Avaliação atendida sem compra", v["unidade"], v["data"], v["cliente"], f"Avaliação ({v['prof']})",
            "Avaliação atendida • não comprou", AV, f"Entra na base do fechamento da avaliação; não comprou plano até {fim:%d/%m/%Y}.")
    for a in rec([a for a in ag if a["tipo"] == T_AV and a["status"] == "Falta"], 2):
        add("Falta na avaliação", a["unidade"], a["data"], a["cliente"], f"Avaliação ({a['prof']})", "Falhou (não compareceu)",
            AV, "Reduz a taxa de comparecimento da avaliação.")
    for x in rec([x for x in vd if x["perfil"] == P_DEPOIS and x["tipo"] == "Plano" and x["conta"]], 2):
        add("Fez avaliação e comprou em outro dia", x["unidade"], x["data"], x["cliente"], x["desc"], f"Plano {reais(x['valor'])}",
            CAB, f"Avaliação atendida em {x['av_data']:%d/%m/%Y}, compra {x['dias']} dias depois, sem avaliação no dia da compra.")
    for x in rec([x for x in vd if x["perfil"] == P_REC and x["tipo"] == "Plano" and x["conta"]], 2):
        add("Cliente que já tinha plano, comprou na cabine", x["unidade"], x["data"], x["cliente"], x["desc"],
            f"Plano {reais(x['valor'])}", CAB, "Renovação ou novo plano de quem já tinha plano aprovado antes.")
    for x in rec([x for x in vd if x["grupo"] == AV and x["tipo"] != "Plano" and x["conta"]], 2):
        add("Pagamento avulso no dia da avaliação", x["unidade"], x["data"], x["cliente"], x["desc"],
            f"{x['tipo']} {reais(x['valor'])}", AV, "Valor pago à parte no dia da avaliação atendida: faturamento avulso da avaliação.")
    return out


def montar(d, ini, fim, estabs, destino):
    wb = Workbook()
    unids = [UNIDADES[e] for e in estabs]
    lista_meses = meses(ini, fim)
    vendas, av_vis, exp_vis, agenda_det = d["vendas"], d["av_vis"], d["exp_vis"], d["agenda_det"]
    res = d["resumo"]
    tipos_ordem = [T_AV, T_EXP_AV, T_SDR, T_EXP_REC, T_CAB, T_OUT]
    resumo_ag = sorted({(u, m, t) for (u, m, t, _) in res}, key=lambda x: (unids.index(x[0]), x[1], tipos_ordem.index(x[2])))
    cab = sorted(d["cab_mes"].items(), key=lambda x: (unids.index(x[0][0]), x[0][1], x[0][2]))
    n = {"V": len(vendas), "AV": len(av_vis), "EX": len(exp_vis), "AA": len(agenda_det), "AR": len(resumo_ag), "CC": len(cab)}

    def rg(b, col):
        return f"'{ABAS[b]}'!${col}$2:${col}${max(n[b], 1) + 1}"

    fonte = (f"Período: {ini:%d/%m/%Y} a {fim:%d/%m/%Y} • Unidades: {', '.join(unids)} • Fonte: API Belle Software, extraído em "
             f"{HOJE:%d/%m/%Y} • Venda de avaliação = compra no dia de uma sessão de avaliação com status Atendido")

    def formulas(g, crit=None):
        """Fórmulas de cada indicador de um grupo; `crit` acrescenta, por aba, um critério (unidade ou mês) às contagens."""
        s, crit = GR[g], crit or {}

        def x(b):
            return f",{crit[b]}" if b in crit else ""
        ar = f"{rg('AR', s['ar'][0])},\"{s['ar'][1]}\"" + x("AR")
        bs = f"{rg(*s['base'])},\"<>\"" + x(s["base"][0])
        vf = ",".join(f"{rg('V', c)},\"{v}\"" for c, v in s["vendas"]) + x("V") + f",{rg('V', 'F')}"
        vv, vn = f"{vf},\"Plano\",{rg('V', 'S')},\"Sim\"", f"{vf},\"Plano\",{rg('V', 'S')},\"Não\""
        va = f"{vf},\"<>Plano\",{rg('V', 'S')},\"Sim\""
        return {"agend": f"SUMIFS({rg('AR', 'J')},{ar})", "atend": f"SUMIFS({rg('AR', 'E')},{ar})",
                "faltas": f"SUMIFS({rg('AR', 'F')},{ar})",
                "desm": f"SUMIFS({rg('AR', 'G')},{ar})+SUMIFS({rg('AR', 'H')},{ar})",
                "base": f"COUNTIFS({bs})", "fech": f"COUNTIFS({bs},{rg(s['base'][0], s['fech'][1])},\"Sim\")",
                "planos": f"COUNTIFS({vv})", "fatpl": f"SUMIFS({rg('V', 'K')},{vv})", "avulso": f"SUMIFS({rg('V', 'K')},{va})",
                "naoapr": f"COUNTIFS({vn})"}

    # ---------------- Resumo
    rs = wb.active
    rs.title = "Resumo"
    titulo(rs, "Drenesse – Vendas de Avaliação x Cabine SDR x Cabine", fonte)
    secao(rs, 4, "1. Comparativo geral")
    h = 5
    cabecalho(rs, h, ["Indicador", AV, SDR, CAB, "Total", "Como é calculado"])
    fs = [formulas(g) for g in (AV, SDR, CAB)]
    r1 = h + 1
    ind = [
        ("Agendamentos", "agend", NUM, "soma", "Avaliação = sessões de avaliação (tipo 'Avaliação' e serviço AVALIAÇÃO ESTÉTICA). Cabine SDR = sessões experimentais de quem nunca teve plano ou com avaliação do dia não atendida. Cabine = demais sessões de serviço."),
        ("Atendidos", "atend", NUM, "soma", "Status 'Atendido' no Belle."),
        ("Faltas (não compareceram)", "faltas", NUM, "soma", "Status 'Falhou' no Belle."),
        ("Desmarcados / cancelados", "desm", NUM, "soma", "Status 'Desmarcado' ou 'Cancelado' (não entram na taxa de comparecimento)."),
        ("Taxa de comparecimento", f"=IFERROR({{c}}{r1 + 1}/({{c}}{r1 + 1}+{{c}}{r1 + 2}),0)", PCT, "taxa", "Atendidos ÷ (atendidos + faltas)."),
        ("Comparecimento sobre todos os agendamentos", f"=IFERROR({{c}}{r1 + 1}/{{c}}{r1},0)", PCT, "taxa", "Atendidos ÷ agendamentos (inclui desmarcados, cancelados e outros status)."),
        ("Base do fechamento", "base", NUM, "-", "Avaliação: avaliações atendidas (cliente/dia). Cabine SDR: sessões experimentais SDR atendidas (cliente/dia). Cabine: clientes atendidos na cabine em cada mês (cliente/mês)."),
        ("Fecharam (compraram plano)", "fech", NUM, "-", "Avaliação e Cabine SDR: compraram plano no mesmo dia. Cabine: compraram plano de cabine no mesmo mês em que foram atendidos."),
        ("Taxa de fechamento", f"=IFERROR({{c}}{r1 + 7}/{{c}}{r1 + 6},0)", PCT, "-", "Fecharam ÷ base do fechamento."),
        ("Planos vendidos (aprovados)", "planos", NUM, "soma", "Planos com status Aprovado, pela data da venda."),
        ("Faturamento em planos (R$)", "fatpl", BRL, "soma", "Preço final dos planos aprovados."),
        ("Ticket médio do plano (R$)", f"=IFERROR({{c}}{r1 + 10}/{{c}}{r1 + 9},0)", BRL, "taxa", "Faturamento em planos ÷ planos vendidos."),
        ("Faturamento avulso (R$)", "avulso", BRL, "soma", "Serviços vendidos à parte (Vendas Detalhado, valor líquido)."),
        ("Faturamento total (R$)", f"={{c}}{r1 + 10}+{{c}}{r1 + 12}", BRL, "soma", "Planos + avulso."),
        ("% do faturamento total", f"=IFERROR({{c}}{r1 + 13}/$E${r1 + 13},0)", PCT, "taxa", "Participação no faturamento total."),
        ("Planos não aprovados (fora do faturamento)", "naoapr", NUM, "soma", "Planos Pendentes/Suspensos/Cancelados vendidos no período – só informativo."),
    ]
    for i, (lbl, chave, fmt, tot, expl) in enumerate(ind):
        r = r1 + i
        vals = ["=" + f[chave] if chave in f else chave.format(c=col) for f, col in zip(fs, "BCD")]
        ve = f"=B{r}+C{r}+D{r}" if tot == "soma" else chave.format(c="E") if tot == "taxa" else "—"
        linha(rs, r, [lbl] + vals + [ve, expl], [None, fmt, fmt, fmt, fmt, None])
        rs.cell(row=r, column=6).font = FI
    r = r1 + len(ind) - 1

    r += 2
    secao(rs, r, "2. Avaliação – em quanto tempo quem fez a avaliação comprou plano")
    r += 1
    cabecalho(rs, r, ["Prazo", "Avaliações atendidas", "Compraram plano", "Taxa", "Observação"])
    tb = f"COUNTIFS({rg('AV', 'A')},\"<>\")"
    for lbl, col, obs in [("No dia da avaliação (venda de avaliação)", "J", "É a taxa de fechamento da avaliação no comparativo."),
                          ("Em até 7 dias", "O", "Inclui o dia. Compras em outro dia contam como Cabine SDR ou Cabine, conforme a regra."),
                          ("Em até 30 dias", "P", f"Avaliações do fim do período têm menos de 30 dias até {fim:%d/%m/%Y}.")]:
        r += 1
        linha(rs, r, [lbl, f"={tb}", f"=COUNTIFS({rg('AV', col)},\"Sim\")", f"=IFERROR(C{r}/B{r},0)", obs], [None, NUM, NUM, PCT, None])
    r += 1
    linha(rs, r, ["Avaliações de clientes que já tinham plano", f"={tb}", f"=COUNTIFS({rg('AV', 'I')},\"Sim\")", f"=IFERROR(C{r}/B{r},0)",
                  "Fazem parte da avaliação pela regra (avaliação atendida no dia); marcadas na aba Avaliações."], [None, NUM, NUM, PCT, None])

    r += 2
    secao(rs, r, "3. De onde vêm as vendas – situação do cliente no dia da compra")
    r += 1
    hp = r
    cabecalho(rs, r, ["Situação no dia da compra", "Grupo", "Planos vendidos", "Faturamento em planos (R$)", "Faturamento avulso (R$)",
                      "Faturamento total (R$)", "% do faturamento total"])
    tp = hp + len(PERFIS) + 1
    for g, perfil in PERFIS:
        r += 1
        c = f"{rg('V', 'M')},$A{r},{rg('V', 'S')},\"Sim\""
        linha(rs, r, [perfil, g, f"=COUNTIFS({c},{rg('V', 'F')},\"Plano\")", f"=SUMIFS({rg('V', 'K')},{c},{rg('V', 'F')},\"Plano\")",
                      f"=SUMIFS({rg('V', 'K')},{c},{rg('V', 'F')},\"<>Plano\")", f"=D{r}+E{r}", f"=IFERROR(F{r}/$F${tp},0)"],
              [None, None, NUM, BRL, BRL, BRL, PCT])
    r += 1
    linha(rs, r, ["TOTAL", ""] + [f"=SUM({c}{hp + 1}:{c}{r - 1})" for c in "CDEFG"], [None, None, NUM, BRL, BRL, BRL, PCT], total=True)

    def bloco(ws, r, grupo, dim, chaves, fmt_chave):
        """Tabela de um grupo por unidade ou por mês. Devolve a linha do total."""
        rot = {AV: ("Avaliações atendidas (cliente/dia)", "Avaliações que fecharam"),
               SDR: ("Sessões SDR atendidas (cliente/dia)", "Sessões SDR que fecharam"),
               CAB: ("Clientes atendidos (cliente/mês)", "Clientes que compraram plano")}[grupo]
        cabecalho(ws, r, ["Unidade" if dim == "unidade" else "Mês", "Agendamentos", "Atendidos", "Faltas", "Desmarcados / cancelados",
                          "Taxa de comparecimento", rot[0], rot[1], "Taxa de fechamento", "Planos vendidos",
                          "Faturamento em planos (R$)", "Ticket médio do plano (R$)", "Faturamento avulso (R$)", "Faturamento total (R$)"])
        ws.row_dimensions[r].height = 45
        h0 = r
        fm = [fmt_chave, NUM, NUM, NUM, NUM, PCT, NUM, NUM, PCT, NUM, BRL, BRL, BRL, BRL]
        for k in chaves:
            r += 1
            f = formulas(grupo, {b: f"{rg(b, DIM[dim][b])},$A{r}" for b in DIM[dim]})
            linha(ws, r, [k, "=" + f["agend"], "=" + f["atend"], "=" + f["faltas"], "=" + f["desm"],
                          f"=IFERROR(C{r}/(C{r}+D{r}),0)", "=" + f["base"], "=" + f["fech"], f"=IFERROR(H{r}/G{r},0)",
                          "=" + f["planos"], "=" + f["fatpl"], f"=IFERROR(K{r}/J{r},0)", "=" + f["avulso"], f"=K{r}+M{r}"], fm)
        r += 1
        s = {c: f"=SUM({c}{h0 + 1}:{c}{r - 1})" for c in "BCDEGHJKMN"}
        linha(ws, r, ["TOTAL", s["B"], s["C"], s["D"], s["E"], f"=IFERROR(C{r}/(C{r}+D{r}),0)", s["G"], s["H"],
                      f"=IFERROR(H{r}/G{r},0)", s["J"], s["K"], f"=IFERROR(K{r}/J{r},0)", s["M"], s["N"]],
              [None] + fm[1:], total=True)
        return r

    r += 2
    secao(rs, r, "4a. Avaliação por unidade")
    r = bloco(rs, r + 1, AV, "unidade", unids, None)
    r += 2
    secao(rs, r, "4b. Cabine SDR por unidade")
    r = bloco(rs, r + 1, SDR, "unidade", unids, None)
    r += 2
    secao(rs, r, "4c. Cabine por unidade")
    r = bloco(rs, r + 1, CAB, "unidade", unids, None)
    r += 2
    notas = [
        "COMO LER",
        "• Avaliação = compra no mesmo dia de uma sessão de avaliação com status Atendido (com ou sem sessão experimental no dia).",
        "• Cabine SDR = compra no dia da sessão experimental quando a avaliação do dia teve falta ou foi desmarcada, ou quando só houve a experimental e a cliente nunca teve plano.",
        "• Cabine = qualquer outra compra (inclusive sessão experimental de quem já tinha plano). Fechamento medido por cliente/mês.",
        "• Os números são fórmulas sobre as abas de base (Vendas, Avaliações, Sessões SDR, Agenda resumo, Clientes cabine).",
        "• Passo a passo completo na aba Metodologia; casos reais na aba Exemplos.",
    ]
    for i, t in enumerate(notas):
        c = rs.cell(row=r + i, column=1, value=t)
        if i == 0:
            c.font = FB
    larguras(rs, [46, 15, 15, 15, 18, 15, 18, 16, 13, 12, 17, 15, 16, 17])
    rs.row_dimensions[h].height = rs.row_dimensions[hp].height = 45

    # ---------------- Mensal
    wm = wb.create_sheet("Mensal")
    titulo(wm, "Avaliação x Cabine SDR x Cabine – mês a mês", fonte)
    r = 4
    for g, lbl in [(AV, "A. Avaliação"), (SDR, "B. Cabine SDR"), (CAB, "C. Cabine")]:
        secao(wm, r, lbl)
        r = bloco(wm, r + 1, g, "mes", lista_meses, MES) + 2
    larguras(wm, [11, 14, 12, 10, 14, 14, 18, 14, 13, 11, 17, 15, 16, 17])
    wm.freeze_panes = "B6"

    # ---------------- Serviços experimentais
    wsv = wb.create_sheet("Serviços SDR")
    titulo(wsv, "Cabine SDR por serviço experimental – comparecimento e fechamento no dia", fonte)
    cabecalho(wsv, 4, ["ID do serviço", "Serviço", "Agendamentos", "Atendidos", "Faltas", "Desmarcados / cancelados",
                       "Taxa de comparecimento", "Sessões SDR atendidas (cliente/dia)", "Fecharam no dia",
                       "Taxa de fechamento", "Planos vendidos", "Valor vendido (R$)", "Ticket médio (R$)"])
    fms = ["0", None, NUM, NUM, NUM, NUM, PCT, NUM, NUM, PCT, NUM, BRL, BRL]
    nomes_srv = {}
    for a in agenda_det:
        nomes_srv.setdefault(a["servico_cod"], a["servico"])
    r = 4
    for cod_s, nm in EXPERIMENTAL.items():
        r += 1
        aa = f"{rg('AA', 'H')},$A{r},{rg('AA', 'G')},\"{T_SDR}\""
        ex_ = f"{rg('EX', 'F')},$A{r}"
        linha(wsv, r, [cod_s, nomes_srv.get(cod_s, nm), f"=COUNTIFS({aa})", f"=COUNTIFS({aa},{rg('AA', 'K')},\"Atendido\")",
                       f"=COUNTIFS({aa},{rg('AA', 'K')},\"Falta\")",
                       f"=COUNTIFS({aa},{rg('AA', 'K')},\"Desmarcado\")+COUNTIFS({aa},{rg('AA', 'K')},\"Cancelado\")",
                       f"=IFERROR(D{r}/(D{r}+E{r}),0)", f"=COUNTIFS({ex_})", f"=COUNTIFS({ex_},{rg('EX', 'L')},\"Sim\")",
                       f"=IFERROR(I{r}/H{r},0)", f"=SUMIFS({rg('EX', 'M')},{ex_})", f"=SUMIFS({rg('EX', 'N')},{ex_})",
                       f"=IFERROR(L{r}/K{r},0)"], fms)
    r += 1
    s = {c: f"=SUM({c}5:{c}{r - 1})" for c in "CDEFHIKL"}
    linha(wsv, r, ["", "TOTAL", s["C"], s["D"], s["E"], s["F"], f"=IFERROR(D{r}/(D{r}+E{r}),0)", s["H"], s["I"],
                   f"=IFERROR(I{r}/H{r},0)", s["K"], s["L"], f"=IFERROR(L{r}/K{r},0)"], fms, total=True)
    wsv.cell(row=r + 2, column=1, value="Ficam fora daqui as sessões experimentais no dia de uma avaliação atendida (Avaliação) e as de "
             "quem já tinha plano (Cabine) – ver 'Tipo de sessão' na aba Agenda avaliação-experimental.").font = FI
    wsv.cell(row=r + 3, column=1, value="Sessões com esses IDs vinculadas a um plano são sessões de cabine comuns (não experimentais).").font = FI
    larguras(wsv, [11, 46, 13, 11, 9, 14, 14, 18, 10, 13, 10, 16, 14])
    wsv.row_dimensions[4].height = 45

    # ---------------- Exemplos
    we = wb.create_sheet("Exemplos")
    titulo(we, "Exemplos reais – como cada atendimento e cada venda foi classificado", fonte)
    cabecalho(we, 4, ["Caso", "Unidade", "Data", "Cliente", "Sessões no dia / plano", "Situação no Belle", "Contado como", "Por quê"])
    r = 4
    for vals in exemplos(d, fim):
        r += 1
        linha(we, r, vals, [None, None, DATA, None, None, None, None, None])
    larguras(we, [40, 15, 11, 32, 52, 46, 13, 80])
    we.freeze_panes = "A5"

    # ---------------- Metodologia
    wt = wb.create_sheet("Metodologia")
    titulo(wt, "Passo a passo – como os números foram obtidos", fonte)
    passos = [
        "1. Dados extraídos da API do Belle, por unidade, em janelas de até 3 meses (limite da API):",
        "   • Relatório de Atendimentos (agenda): data, cliente, tipo de agendamento, serviço, status e plano vinculado.",
        "   • Venda de Planos: data da venda, cliente, plano, status, vendedor e preço final (todo o histórico, todas as unidades).",
        "   • Vendas Detalhado: serviços vendidos à parte, com valor líquido.",
        f"   • A agenda também foi lida nos {ANTES} dias anteriores ao início, só para saber quem já tinha feito avaliação.",
        "2. Cada agendamento foi classificado:",
        "   • Sessão de avaliação = agendamento do tipo 'Avaliação' no Belle, ou o serviço 52 – AVALIAÇÃO ESTÉTICA (feito pelas mesmas avaliadoras).",
        "   • Sessão experimental = um dos 6 IDs abaixo, sem plano vinculado (agendamento avulso). Ela entra no grupo da venda do dia:",
        "     Avaliação (se houve avaliação atendida no dia), Cabine SDR (avaliação do dia não atendida, ou cliente que nunca teve plano)",
        "     ou Cabine (cliente que já tinha plano).",
        "   • Sessão de cabine = qualquer outro agendamento de serviço, inclusive sessões com os 6 IDs que estão dentro de um plano.",
        "   • Retorno, consulta e agendamentos sem serviço ficam fora das contas (ver diagnóstico).",
        "3. Comparecimento = atendidos ÷ (atendidos + faltas), pelo status do Belle (Atendido, Falhou, Desmarcado, Cancelado).",
        "4. Cada venda (plano ou serviço avulso) foi olhada pelo cliente e pelo dia da compra:",
        "   • Teve sessão de avaliação com status Atendido no dia → venda de AVALIAÇÃO (com ou sem sessão experimental no dia).",
        "   • Avaliação do dia com falta ou desmarcada, mas sessão experimental atendida → venda de CABINE SDR.",
        "   • Só sessão experimental no dia, para quem nunca teve plano aprovado → venda de CABINE SDR.",
        "   • Qualquer outra compra (experimental de quem já tinha plano, renovações, avaliação em outro dia, sem sessão no dia) → venda de CABINE.",
        "5. Fechamento:",
        "   • Avaliação: avaliações atendidas (cliente/dia) com plano comprado no dia ÷ avaliações atendidas. Também em até 7 e 30 dias.",
        "   • Cabine SDR: sessões experimentais SDR atendidas (cliente/dia) com plano comprado no dia ÷ essas sessões.",
        "   • Cabine: para cada mês, clientes atendidos na cabine que compraram plano de cabine no mesmo mês ÷ clientes atendidos na cabine.",
        "6. Faturamento = valor vendido na data da venda: preço final dos planos Aprovados + valor líquido dos serviços avulsos.",
        "   Planos Pendentes/Suspensos/Cancelados aparecem na aba Vendas, mas ficam fora do faturamento e do fechamento.",
        "7. Todas as contas do Resumo e do Mensal são fórmulas sobre as abas de base: dá para conferir filtrando essas abas.",
    ]
    r = 4
    for t in passos:
        c = wt.cell(row=r, column=1, value=t)
        if t[0].isdigit():
            c.font = FB
        r += 1
    r += 1
    secao(wt, r, "IDs das sessões experimentais (informados pela Drenesse)")
    r += 1
    cabecalho(wt, r, ["ID", "Serviço"])
    for cod_s, nm in EXPERIMENTAL.items():
        r += 1
        linha(wt, r, [cod_s, nm], ["0", None])
    dg = d["diag"]
    r += 2
    secao(wt, r, "Conferência e diagnóstico dos dados")
    linhas_diag = [
        ("Planos aprovados que o Belle liga a um atendimento (campo 'avaliação' do plano, preenchido só em parte dos planos)", dg["ligados"]),
        ("   … ligados a uma sessão de avaliação atendida no dia da venda", dg["lig_av"]),
        ("   … desses, classificados como venda de avaliação pela regra do dia (conferência)", dg["lig_av_ok"]),
        ("   … ligados a outro atendimento, ou a avaliação de outro dia ou não atendida (seguem a regra do dia; principais abaixo)",
         sum(dg["lig_outros"].values())),
    ] + [(f"         – {k}", q) for k, q in dg["lig_outros"].most_common(6)] + [
        ("Agendamentos do serviço 52 – AVALIAÇÃO ESTÉTICA (contados como sessão de avaliação)", dg["serv52"]),
        ("Sessões com os IDs experimentais vinculadas a plano (contadas como sessão de cabine)", dg["exp_plano"]),
        ("Serviços avulsos da mesma venda de um plano (fora do faturamento para não contar 2 vezes)", dg["repetidos"]),
    ] + [(f"Fora da análise – '{k}'", q) for k, q in dg["fora"].most_common()] \
      + [(f"Itens no Vendas Detalhado – tipo '{k}'" + (" (os planos entram pelo relatório de planos)" if k == "Plano" else ""), q)
         for k, q in dg["tipos"].most_common()]
    for lbl, q in linhas_diag:
        r += 1
        linha(wt, r, [lbl, q], [None, NUM])
    larguras(wt, [130, 14])

    # ---------------- Bases
    base(wb.create_sheet(ABAS["V"]),
         ["Unidade", "Data da venda", "Mês", "Cliente", "Cód. cliente", "Tipo", "Descrição", "Nº plano", "Status do plano", "Vendedor",
          "Valor (R$)", "Grupo", "Situação no dia da compra", "Sessão experimental no dia?", "Já tinha plano antes?",
          "Avaliação anterior (data)", "Dias após a avaliação", "Avaliação ligada ao plano no Belle (ID)", "Conta no faturamento?",
          "Observação"],
         [[x["unidade"], x["data"], x["data"].replace(day=1), x["cliente"], x["cliente_cod"], x["tipo"], x["desc"], x["orc"] or None,
           x["status"] or None, x["vendedor"] or None, x["valor"], x["grupo"], x["perfil"], "Sim" if x["exp"] else "Não",
           "Sim" if x["ja_cliente"] else "Não", x["av_data"], x["dias"], x["av_belle"], "Sim" if x["conta"] else "Não",
           x["obs"] or None] for x in vendas],
         [None, DATA, MES, None, "0", None, None, "0", None, None, BRL, None, None, None, None, DATA, NUM, "0", None, None],
         "TabVendas", [15, 11, 10, 32, 10, 14, 34, 11, 11, 26, 13, 11, 44, 12, 10, 11, 9, 12, 11, 40])
    base(wb.create_sheet(ABAS["AV"]),
         ["Unidade", "Data", "Mês", "Cliente", "Cód. cliente", "Avaliadora", "Hora", "Sessão experimental no dia",
          "Já tinha plano antes?", "Fechou plano no dia?", "Planos vendidos no dia", "Valor vendido no dia (R$)", "Vendedor",
          "Dias até a 1ª compra de plano", "Comprou em até 7 dias?", "Comprou em até 30 dias?", "1ª compra de plano após a avaliação"],
         [[v["unidade"], v["data"], v["mes"], v["cliente"], v["cliente_cod"], v["prof"] or None, v["hora"] or None, v["exp"] or "Não",
           "Sim" if v["ja_cliente"] else "Não", "Sim" if v["vendas"] else "Não", len(v["vendas"]),
           round(sum(x["valor"] for x in v["vendas"]), 2), ", ".join(sorted({x["vendedor"] for x in v["vendas"] if x["vendedor"]})) or None,
           v["dias_compra"], "Sim" if v["dias_compra"] is not None and v["dias_compra"] <= 7 else "Não",
           "Sim" if v["dias_compra"] is not None and v["dias_compra"] <= 30 else "Não", v["primeira"] or None] for v in av_vis],
         [None, DATA, MES, None, "0", None, None, None, None, None, NUM, BRL, None, NUM, None, None, None], "TabAvaliacoes",
         [15, 11, 10, 32, 10, 28, 7, 34, 10, 11, 10, 13, 26, 11, 10, 10, 34])
    base(wb.create_sheet(ABAS["EX"]),
         ["Unidade", "Data", "Mês", "Cliente", "Cód. cliente", "ID do serviço", "Serviço experimental", "Outros serviços no dia",
          "Profissional", "Avaliação no dia", "Já tinha plano antes?", "Fechou plano no dia?", "Planos vendidos no dia",
          "Valor vendido no dia (R$)", "Vendedor"],
         [[v["unidade"], v["data"], v["mes"], v["cliente"], v["cliente_cod"], v["servico_cod"], v["servico"], v["outros"] or None,
           v["prof"] or None, v["av_status"], "Sim" if v["ja_cliente"] else "Não", "Sim" if v["vendas"] else "Não", len(v["vendas"]),
           round(sum(x["valor"] for x in v["vendas"]), 2),
           ", ".join(sorted({x["vendedor"] for x in v["vendas"] if x["vendedor"]})) or None] for v in exp_vis],
         [None, DATA, MES, None, "0", "0", None, None, None, None, None, None, NUM, BRL, None], "TabExperimentais",
         [15, 11, 10, 32, 10, 11, 36, 26, 28, 14, 10, 11, 10, 13, 26])
    base(wb.create_sheet(ABAS["AA"]),
         ["Unidade", "Data", "Mês", "Hora", "Cliente", "Cód. cliente", "Tipo de sessão", "ID do serviço", "Serviço",
          "Status no Belle", "Status", "Profissional"],
         [[a["unidade"], a["data"], a["mes"], a["hora"], a["cliente"], a["cliente_cod"], a["tipo"], a["servico_cod"] or None,
           a["servico"], a["status_belle"], a["status"], a["prof"] or None] for a in agenda_det],
         [None, DATA, MES, None, None, "0", None, "0", None, None, None, None], "TabAgendaAvExp",
         [15, 11, 10, 7, 32, 10, 18, 11, 40, 12, 11, 28])
    base(wb.create_sheet(ABAS["AR"]),
         ["Unidade", "Mês", "Grupo", "Tipo de sessão", "Atendidos", "Faltas", "Desmarcados", "Cancelados", "Outros status",
          "Total de agendamentos"],
         [[u, m, GRUPO[t], t, res[(u, m, t, "Atendido")], res[(u, m, t, "Falta")], res[(u, m, t, "Desmarcado")],
           res[(u, m, t, "Cancelado")], res[(u, m, t, "Outros")], f"=SUM(E{i}:I{i})"] for i, (u, m, t) in enumerate(resumo_ag, 2)],
         [None, MES, None, None, NUM, NUM, NUM, NUM, NUM, NUM], "TabAgendaResumo", [15, 11, 15, 30, 11, 10, 12, 11, 11, 13])
    base(wb.create_sheet(ABAS["CC"]),
         ["Unidade", "Mês", "Cliente", "Cód. cliente", "Sessões atendidas na cabine", "Comprou plano na cabine no mês?",
          "Planos comprados no mês", "Valor comprado no mês (R$)"],
         [[u, m, d["nomes"].get(c, ""), c, q, "Sim" if d["compra_cab"].get((c, m)) else "Não", len(d["compra_cab"].get((c, m), [])),
           round(sum(d["compra_cab"].get((c, m), [])), 2)] for (u, m, c), q in cab],
         [None, MES, None, "0", NUM, None, NUM, BRL], "TabClientesCabine", [15, 11, 32, 10, 12, 13, 11, 14])

    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    compactar_xlsx(destino, {ABAS["AR"]: {"J"}})


if __name__ == "__main__":
    datas = [dt.datetime.strptime(x, "%d/%m/%Y").date() for x in sys.argv[1:] if "/" in x]
    ini = datas[0] if datas else dt.date(2026, 1, 1)
    fim = datas[1] if len(datas) > 1 else dt.date(2026, 9, 26)
    estabs = [int(x) for x in sys.argv[1:] if "/" not in x] or ESTABS
    dados = coletar(ini, fim, estabs)
    out = os.environ.get("BELLE_OUT", "Avaliacao_x_Cabine_Drenesse.xlsx")
    montar(dados, ini, fim, estabs, out)
    dg = dados["diag"]
    print(f"{len(dados['vendas'])} vendas, {len(dados['av_vis'])} avaliações atendidas, {len(dados['exp_vis'])} sessões SDR "
          f"atendidas -> {out}")
    print(f"Planos ligados a atendimento no Belle: {dg['ligados']}; a sessão de avaliação atendida no dia: {dg['lig_av']} "
          f"(classificados como avaliação: {dg['lig_av_ok']})")
