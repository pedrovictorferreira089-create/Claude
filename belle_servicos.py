"""Ranking de serviços realizados por período, todas as unidades (API Belle Software).

Uso: python3 belle_servicos.py [dd/mm/aaaa início] [dd/mm/aaaa fim]   (padrão: 01/09/2026 a 28/09/2026)
Reaproveita a regra de valor faturado de belle_sabados.py (plano rateado / venda avulsa do dia).
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

from belle_sabados import (BOX, BRL, F, FB, FH, FI, FT, HEAD, INT, SUB, TOT, UNIDADES, api, br, cabecalho, carregar_planos,
                           larguras, valorar)
from openpyxl.utils import get_column_letter


def coletar(ini, fim):
    tabela = {str(s["codServico"]): br(s["valor"]) for s in api("servico/listar")}
    planos = carregar_planos()
    linhas = []
    for e, unidade in UNIDADES.items():
        ags = api("relatorios/relatorio_atendimentos", dtInicio=ini, dtFim=fim, codEstab=str(e))
        vendas = api("vendas_detalhado", estab=str(e), dtInicio=ini, dtFim=fim)
        por_dia = {}
        for v in vendas:
            d = dt.date.fromisoformat(v["data_venda"]).strftime("%d/%m/%Y")
            por_dia.setdefault(d, []).extend((v["cod_cliente"], i) for i in v["itens_venda"] if i["tipo"] == "Serviço")
        for a in sorted(ags, key=lambda x: (x["dataAgendamento"][6:] + x["dataAgendamento"][3:5] + x["dataAgendamento"][:2], x["horarioAgendamento"])):
            if a["statusAgendamento"] != "Atendido":
                continue
            servico = (a["nomeServico"] or "").strip() or a["tipoAgendamento"]
            tab, val, origem, obs = valorar(a, servico, planos, tabela, por_dia.setdefault(a["dataAgendamento"], []))
            linhas.append(dict(unidade=unidade, data=dt.datetime.strptime(a["dataAgendamento"], "%d/%m/%Y").date(),
                               cliente=" ".join(a["nomeCliente"].split()), servico=" ".join(servico.split()),
                               prof=" ".join((a["nomeProfissional"] or "").split()), origem=origem,
                               tabela=tab, valor=round(val, 2), obs=obs))
    return linhas


def montar(linhas, ini, fim, destino):
    wb = Workbook()
    N = len(linhas) + 1
    U, E, V = f"Atendimentos!$A$2:$A${N}", f"Atendimentos!$D$2:$D${N}", f"Atendimentos!$I$2:$I${N}"
    servs = sorted({l["servico"] for l in linhas}, key=lambda s: -sum(l["valor"] for l in linhas if l["servico"] == s))
    unids = list(UNIDADES.values())

    # ---- Serviços (consolidado)
    ws = wb.active
    ws.title = "Serviços"
    ws["A1"] = "Drenesse – Serviços realizados, todas as unidades"
    ws["A1"].font = FT
    ws["A2"] = (f"Período: {ini} a {fim} • Somente agendamentos com status 'Atendido' • Fonte: API Belle Software. "
                "Valor faturado = sessão de plano rateada pelo preço final do plano, ou venda avulsa registrada no dia.")
    ws["A2"].font = FI
    cabecalho(ws, 4, ["#", "Serviço", "Qtd realizada", "Valor faturado (R$)", "Ticket médio (R$)", "% do faturado"])
    last = 4 + len(servs)
    for i, s in enumerate(servs, 5):
        vals = [i - 4, s, f"=COUNTIFS({E},B{i})", f"=SUMIFS({V},{E},B{i})", f"=IFERROR(D{i}/C{i},0)", f"=IFERROR(D{i}/$D${last + 1},0)"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = F, BOX
    t = last + 1
    for j, v in enumerate(["", "TOTAL", f"=SUM(C5:C{last})", f"=SUM(D5:D{last})", f"=IFERROR(D{t}/C{t},0)", f"=SUM(F5:F{last})"], 1):
        c = ws.cell(row=t, column=j, value=v)
        c.font, c.fill, c.border = FB, TOT, BOX
    for r in range(5, t + 1):
        ws.cell(row=r, column=3).number_format = INT
        ws.cell(row=r, column=4).number_format = BRL
        ws.cell(row=r, column=5).number_format = BRL
        ws.cell(row=r, column=6).number_format = "0.0%"
    larguras(ws, [5, 48, 13, 18, 16, 13])
    ws.freeze_panes = "C5"

    # ---- Serviços x Unidade
    wu = wb.create_sheet("Por Unidade")
    wu["A1"] = "Serviços por unidade – quantidade e valor faturado"
    wu["A1"].font = FT
    wu["A2"] = f"Período: {ini} a {fim}"
    wu["A2"].font = FI
    wu.cell(row=4, column=1, value="Serviço")
    col = 2
    for u in unids + ["TOTAL"]:
        wu.cell(row=4, column=col, value=u)
        wu.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        wu.cell(row=5, column=col, value="Qtd")
        wu.cell(row=5, column=col + 1, value="Faturado (R$)")
        col += 2
    wu.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    for rr in (4, 5):
        for c in range(1, col):
            cell = wu.cell(row=rr, column=c)
            cell.font, cell.fill, cell.border = FH, HEAD, BOX
            cell.alignment = Alignment(horizontal="center", vertical="center")
    last_u = 5 + len(servs)
    for i, s in enumerate(servs, 6):
        wu.cell(row=i, column=1, value=s)
        for k, u in enumerate(unids):
            cq = 2 + 2 * k
            hdr = f"{get_column_letter(cq)}$4"
            wu.cell(row=i, column=cq, value=f"=COUNTIFS({E},$A{i},{U},{hdr})")
            wu.cell(row=i, column=cq + 1, value=f"=SUMIFS({V},{E},$A{i},{U},{hdr})")
        tq = 2 + 2 * len(unids)
        wu.cell(row=i, column=tq, value="=" + "+".join(f"{get_column_letter(2 + 2 * k)}{i}" for k in range(len(unids))))
        wu.cell(row=i, column=tq + 1, value="=" + "+".join(f"{get_column_letter(3 + 2 * k)}{i}" for k in range(len(unids))))
    tr = last_u + 1
    wu.cell(row=tr, column=1, value="TOTAL")
    for c in range(2, col):
        L = get_column_letter(c)
        wu.cell(row=tr, column=c, value=f"=SUM({L}6:{L}{last_u})")
    for r in range(6, tr + 1):
        for c in range(1, col):
            cell = wu.cell(row=r, column=c)
            cell.font, cell.border = (FB if r == tr else F), BOX
            if c > 1:
                cell.number_format = INT if c % 2 == 0 else BRL
            if r == tr:
                cell.fill = TOT
            elif c >= col - 2:
                cell.fill = SUB
    larguras(wu, [44] + [8, 15] * (len(unids) + 1))
    wu.freeze_panes = "B6"

    # ---- Atendimentos (base)
    wa = wb.create_sheet("Atendimentos")
    cabecalho(wa, 1, ["Unidade", "Data", "Cliente", "Serviço", "Profissional", "Origem do valor", "Valor tabela (R$)", "Obs.", "Valor faturado (R$)"])
    for i, l in enumerate(linhas, 2):
        for j, v in enumerate([l["unidade"], l["data"], l["cliente"], l["servico"], l["prof"], l["origem"], l["tabela"], l["obs"], l["valor"]], 1):
            c = wa.cell(row=i, column=j, value=v)
            c.font = F
        wa.cell(row=i, column=2).number_format = "dd/mm/yyyy"
        wa.cell(row=i, column=7).number_format = BRL
        wa.cell(row=i, column=9).number_format = BRL
    wa.add_table(Table(displayName="Base", ref=f"A1:I{N}", tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))
    larguras(wa, [14, 11, 36, 40, 32, 15, 14, 55, 15])
    wa.freeze_panes = "A2"

    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)


if __name__ == "__main__":
    ini = sys.argv[1] if len(sys.argv) > 1 else "01/09/2026"
    fim = sys.argv[2] if len(sys.argv) > 2 else "28/09/2026"
    linhas = coletar(ini, fim)
    out = os.environ.get("BELLE_OUT", "Servicos_Drenesse.xlsx")
    montar(linhas, ini, fim, out)
    print(f"{len(linhas)} atendimentos -> {out}")
