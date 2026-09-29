"""Gera planilha de faturamento dos sábados (últimos 90 dias) a partir da API Belle Software.

Uso: python3 belle_sabados.py [codEstab ...]   (padrão: 2 = Lagoa Nova)
Os dados brutos da API ficam em cache em ./data para não estourar o rate limit (40 req/min).
"""
import datetime as dt
import glob
import json
import os
import sys
import time
import urllib.parse
import urllib.request

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = "https://app.bellesoftware.com.br/api/release/controller/IntegracaoExterna/v1.0/"
TOKEN = os.environ.get("BELLE_TOKEN", "409746c0fb619acbe444d0834766505c")
CACHE = os.environ.get("BELLE_CACHE", os.path.join(os.path.dirname(os.path.abspath(__file__)), "data"))
HOJE = dt.date(2026, 9, 29)
UNIDADES = {1: "Petrópolis", 2: "Lagoa Nova", 3: "Capim Macio", 6: "Norte Shopping", 4: "Laser"}
MAX_ROWS = 5000  # alcance das fórmulas do resumo


def api(path, **q):
    os.makedirs(CACHE, exist_ok=True)
    fn = os.path.join(CACHE, path.replace("/", "_") + "__" + "_".join(f"{k}-{v}".replace("/", "") for k, v in q.items()) + ".json")
    if os.path.exists(fn):
        return json.load(open(fn))
    req = urllib.request.Request(BASE + path + "?" + urllib.parse.urlencode(q), headers={"Authorization": TOKEN})
    for tentativa in range(4):
        try:
            data = json.loads(urllib.request.urlopen(req, timeout=180).read())
            break
        except Exception:
            if tentativa == 3:
                raise
            time.sleep(20)
    time.sleep(1.7)
    json.dump(data, open(fn, "w"), ensure_ascii=False)
    return data


def br(v):
    if v in (None, ""):
        return 0.0
    s = str(v)
    return float(s.replace(".", "").replace(",", ".")) if "," in s else float(s)


def sabados():
    d = HOJE - dt.timedelta(days=(HOJE.weekday() - 5) % 7)
    out = []
    while d >= HOJE - dt.timedelta(days=90):
        out.append(d)
        d -= dt.timedelta(weeks=1)
    return sorted(out)


def carregar_planos():
    """Valor por sessão de cada serviço em cada plano vendido (rateio do preço final)."""
    for e in UNIDADES:
        fim = HOJE
        while fim > dt.date(2023, 10, 1):
            ini = fim - dt.timedelta(days=89)
            api("venda_planos", dtInicio=ini.strftime("%d/%m/%Y"), dtFim=fim.strftime("%d/%m/%Y"), codEstab=str(e))
            fim = ini - dt.timedelta(days=1)
    valores = {}
    for f in glob.glob(os.path.join(CACHE, "venda_planos__*")):
        for p in json.load(open(f)):
            liq = {}
            for s in p["servicos"]:
                if s["cortesia"] == "Sim":
                    v = 0.0
                elif s["tipoDesconto"] == "%":
                    v = br(s["valorTotalServico"]) * (1 - br(s["desconto"]) / 100)
                else:
                    v = br(s["valorTotalServico"]) - br(s["desconto"])
                liq.setdefault(str(s["codigoServico"]), []).append((v, s["qtdSessoes"] or 1))
            soma = sum(v for lst in liq.values() for v, _ in lst)
            fator = br(p["precoFinal"]) / soma if soma else 0.0
            valores[p["codOrcamento"]] = {
                "nome": p["nomePlano"].strip(),
                "svc": {c: sum(v for v, _ in lst) * fator / sum(q for _, q in lst) for c, lst in liq.items()},
            }
    return valores


def coletar(estabs):
    tabela = {str(s["codServico"]): br(s["valor"]) for s in api("servico/listar")}
    planos = carregar_planos()
    atend, faltas = [], []
    for e in estabs:
        for d in sabados():
            D = d.strftime("%d/%m/%Y")
            ags = api("relatorios/relatorio_atendimentos", dtInicio=D, dtFim=D, codEstab=str(e))
            vendas = api("vendas_detalhado", estab=str(e), dtInicio=D, dtFim=D)
            avulsos = [(v["cod_cliente"], i) for v in vendas for i in v["itens_venda"] if i["tipo"] == "Serviço"]
            for a in sorted(ags, key=lambda x: (x["horarioAgendamento"], x["nomeCliente"])):
                base = dict(unidade=UNIDADES.get(e, str(e)), data=d, hora=a["horarioAgendamento"],
                            cliente=" ".join(a["nomeCliente"].split()), servico=(a["nomeServico"] or "").strip() or a["tipoAgendamento"],
                            prof=" ".join((a["nomeProfissional"] or "").split()))
                if a["statusAgendamento"] == "Falhou":
                    faltas.append(base)
                if a["statusAgendamento"] != "Atendido":
                    continue
                cod, orc = str(a["codigoServico"] or ""), a["idOrcamento"]
                tab = tabela.get(cod)
                if not cod:
                    val, origem, obs = 0.0, "Sem serviço", "Agendamento sem serviço vinculado"
                elif orc and orc in planos and cod in planos[orc]["svc"]:
                    val, origem = planos[orc]["svc"][cod], "Plano"
                    obs = f"Sessão do plano {orc} ({planos[orc]['nome']})"
                    if val == 0:
                        obs += " – cortesia/100% desconto"
                elif orc:
                    val, origem, obs = tab or 0.0, "Tabela", f"Plano {orc} não localizado – usado valor de tabela"
                else:
                    m = next((x for x in avulsos if x[0] == a["codigoCliente"] and x[1]["desc_item"].strip() == base["servico"]), None)
                    if m:
                        avulsos.remove(m)
                        val, origem, obs = br(m[1]["valor_liquido"]), "Avulso", "Venda avulsa registrada no dia"
                    else:
                        val, origem, obs = 0.0, "Avulso s/ venda", "Sem venda registrada no dia (experimental, cortesia ou pago em outra data)"
                atend.append(dict(base, origem=origem, orc=orc or "", tabela=tab, valor=round(val, 2), obs=obs))
    return atend, faltas


# ---------------------------------------------------------------- planilha
ARIAL = "Arial"
F = Font(name=ARIAL, size=10)
FB = Font(name=ARIAL, size=10, bold=True)
FH = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
FT = Font(name=ARIAL, size=14, bold=True)
FI = Font(name=ARIAL, size=9, italic=True, color="555555")
HEAD = PatternFill("solid", fgColor="1F4E78")
SUB = PatternFill("solid", fgColor="D9E1F2")
TOT = PatternFill("solid", fgColor="FFF2CC")
FILL_IN = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BRL = 'R$ #,##0.00;-R$ #,##0.00;"-"'
INT = '0;-0;"-"'
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def cabecalho(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.alignment, cell.border = FH, HEAD, CENTER, BOX


def larguras(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def montar(atend, faltas, destino):
    wb = Workbook()
    datas = sabados()

    # ---- Resumo
    rs = wb.active
    rs.title = "Resumo"
    rs["A1"] = "Drenesse – Faturamento dos sábados (últimos 90 dias)"
    rs["A1"].font = FT
    rs["A2"] = (f"Período: {datas[0]:%d/%m/%Y} a {datas[-1]:%d/%m/%Y} • Fonte: API Belle Software, extraído em {HOJE:%d/%m/%Y}. "
                "Valores calculados automaticamente a partir das abas 'Atendimentos' e 'Faltas'.")
    rs["A2"].font = FI
    unids = list(UNIDADES.values())
    r0 = 4
    rs.cell(row=r0, column=1, value="Sábado")
    col = 2
    for u in unids + ["TOTAL GERAL"]:
        rs.cell(row=r0, column=col, value=u)
        rs.merge_cells(start_row=r0, start_column=col, end_row=r0, end_column=col + 2)
        for j, t in enumerate(["Atendidos", "Faturado (R$)", "Faltas"]):
            rs.cell(row=r0 + 1, column=col + j, value=t)
        col += 3
    rs.cell(row=r0, column=col, value="Observação")
    rs.merge_cells(start_row=r0, start_column=1, end_row=r0 + 1, end_column=1)
    rs.merge_cells(start_row=r0, start_column=col, end_row=r0 + 1, end_column=col)
    last_col = col
    for rr in (r0, r0 + 1):
        for c in range(1, last_col + 1):
            cell = rs.cell(row=rr, column=c)
            cell.font, cell.alignment, cell.border = FH, CENTER, BOX
            cell.fill = HEAD
    A = f"Atendimentos!$A$2:$A${MAX_ROWS}"
    B = f"Atendimentos!$B$2:$B${MAX_ROWS}"
    K = f"Atendimentos!$K$2:$K${MAX_ROWS}"
    FA = f"Faltas!$A$2:$A${MAX_ROWS}"
    FBd = f"Faltas!$B$2:$B${MAX_ROWS}"
    for i, d in enumerate(datas):
        r = r0 + 2 + i
        c = rs.cell(row=r, column=1, value=d)
        c.number_format = "dd/mm/yyyy"
        tot_cols = {0: [], 1: [], 2: []}
        for k, u in enumerate(unids):
            cc = 2 + 3 * k
            L = [get_column_letter(cc + j) for j in range(3)]
            rs.cell(row=r, column=cc, value=f'=COUNTIFS({A},{L[0]}${r0},{B},$A{r})')
            rs.cell(row=r, column=cc + 1, value=f'=SUMIFS({K},{A},{L[0]}${r0},{B},$A{r})')
            rs.cell(row=r, column=cc + 2, value=f'=COUNTIFS({FA},{L[0]}${r0},{FBd},$A{r})')
            for j in range(3):
                tot_cols[j].append(f"{L[j]}{r}")
        tc = 2 + 3 * len(unids)
        for j in range(3):
            rs.cell(row=r, column=tc + j, value="=" + "+".join(tot_cols[j]))
        fl = get_column_letter(tc + 2)
        rs.cell(row=r, column=last_col, value=f'=IF({fl}{r}>0,{fl}{r}&" falta(s) no dia","Sem faltas")')
        for c in range(1, last_col + 1):
            cell = rs.cell(row=r, column=c)
            cell.font, cell.border = F, BOX
            if c > 1 and c < last_col:
                cell.number_format = BRL if (c - 2) % 3 == 1 else INT
                cell.alignment = Alignment(horizontal="right")
            if tc <= c < last_col:
                cell.fill = SUB
    rt = r0 + 2 + len(datas)
    rs.cell(row=rt, column=1, value="TOTAL")
    for c in range(2, last_col):
        L = get_column_letter(c)
        rs.cell(row=rt, column=c, value=f"=SUM({L}{r0 + 2}:{L}{rt - 1})").number_format = BRL if (c - 2) % 3 == 1 else INT
    fl = get_column_letter(last_col - 1)
    rs.cell(row=rt, column=last_col, value=f'="Total de faltas no período: "&{fl}{rt}')
    for c in range(1, last_col + 1):
        cell = rs.cell(row=rt, column=c)
        cell.font, cell.fill, cell.border = FB, TOT, BOX
    rs.cell(row=rt + 1, column=1, value="Média por sábado").font = FI
    for c in range(2, last_col):
        L = get_column_letter(c)
        cell = rs.cell(row=rt + 1, column=c, value=f"=IFERROR({L}{rt}/COUNT($A${r0 + 2}:$A${rt - 1}),0)")
        cell.number_format = BRL if (c - 2) % 3 == 1 else '0.0;-0.0;"-"'
        cell.font = FI
    notas = [
        "COMO USAR / LEGENDA",
        "• Para incluir outra unidade: cole as linhas dela nas abas 'Atendimentos' e 'Faltas', escrevendo na coluna 'Unidade' exatamente o nome do cabeçalho acima (ex.: Petrópolis). O resumo se atualiza sozinho.",
        "• Para uma unidade nova (fora da lista), troque o nome de um dos cabeçalhos em amarelo da linha 4.",
        "• Faturado = valor efetivamente apropriado a cada atendimento: sessão de plano = preço final do plano rateado pelas sessões; avulso = valor da venda do dia.",
        "• Faltas = agendamentos com status 'Falhou' no Belle. Desmarcações não entram na conta.",
    ]
    for i, n in enumerate(notas):
        c = rs.cell(row=rt + 3 + i, column=1, value=n)
        c.font = FB if i == 0 else F
    for k in range(len(unids)):
        rs.cell(row=r0, column=2 + 3 * k).fill = FILL_IN
        rs.cell(row=r0, column=2 + 3 * k).font = FB
    rs.cell(row=r0, column=2 + 3 * unids.index("Lagoa Nova")).comment = Comment("Única unidade já preenchida nesta versão.", "Claude")
    larguras(rs, [12] + [11, 14, 8] * (len(unids) + 1) + [26])
    rs.row_dimensions[r0 + 1].height = 28
    rs.freeze_panes = rs.cell(row=r0 + 2, column=2)

    # ---- Atendimentos
    wa = wb.create_sheet("Atendimentos")
    cols = ["Unidade", "Data", "Hora", "Cliente", "Serviço", "Profissional", "Origem do valor", "Nº plano", "Valor tabela (R$)", "Dif. p/ tabela (R$)", "Valor faturado (R$)", "Observação"]
    cabecalho(wa, 1, cols)
    for i, a in enumerate(atend, 2):
        vals = [a["unidade"], a["data"], a["hora"], a["cliente"], a["servico"], a["prof"], a["origem"], a["orc"], a["tabela"], f'=IF(I{i}="","",K{i}-I{i})', a["valor"], a["obs"]]
        for j, v in enumerate(vals, 1):
            c = wa.cell(row=i, column=j, value=v)
            c.font, c.border = F, BOX
        wa.cell(row=i, column=2).number_format = "dd/mm/yyyy"
        for j in (9, 10, 11):
            wa.cell(row=i, column=j).number_format = BRL
    n = len(atend) + 1
    wa.add_table(Table(displayName="Atendimentos", ref=f"A1:L{n}", tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))
    larguras(wa, [13, 11, 7, 38, 34, 32, 15, 11, 14, 14, 15, 60])
    wa.freeze_panes = "A2"

    # ---- Faltas
    wf = wb.create_sheet("Faltas")
    cabecalho(wf, 1, ["Unidade", "Data", "Hora", "Cliente", "Serviço", "Profissional"])
    for i, f in enumerate(faltas, 2):
        for j, v in enumerate([f["unidade"], f["data"], f["hora"], f["cliente"], f["servico"], f["prof"]], 1):
            c = wf.cell(row=i, column=j, value=v)
            c.font, c.border = F, BOX
        wf.cell(row=i, column=2).number_format = "dd/mm/yyyy"
    wf.add_table(Table(displayName="Faltas", ref=f"A1:F{max(len(faltas) + 1, 2)}", tableStyleInfo=TableStyleInfo(name="TableStyleLight10", showRowStripes=True)))
    larguras(wf, [13, 11, 7, 38, 34, 32])
    wf.freeze_panes = "A2"

    # ---- Por serviço
    wsv = wb.create_sheet("Por Serviço")
    wsv["A1"] = "Serviços atendidos nos sábados – todas as unidades lançadas"
    wsv["A1"].font = FT
    cabecalho(wsv, 3, ["Serviço", "Qtd atendida", "Valor faturado (R$)", "Ticket médio (R$)", "% do faturado"])
    servs = sorted({a["servico"] for a in atend}, key=lambda s: -sum(x["valor"] for x in atend if x["servico"] == s))
    E = f"Atendimentos!$E$2:$E${MAX_ROWS}"
    last = 3 + len(servs)
    for i, s in enumerate(servs, 4):
        vals = [s, f"=COUNTIFS({E},A{i})", f"=SUMIFS({K},{E},A{i})", f"=IFERROR(C{i}/B{i},0)", f"=IFERROR(C{i}/$C${last + 1},0)"]
        for j, v in enumerate(vals, 1):
            c = wsv.cell(row=i, column=j, value=v)
            c.font, c.border = F, BOX
        wsv.cell(row=i, column=3).number_format = BRL
        wsv.cell(row=i, column=4).number_format = BRL
        wsv.cell(row=i, column=5).number_format = "0.0%"
    for j, v in enumerate(["TOTAL", f"=SUM(B4:B{last})", f"=SUM(C4:C{last})", f"=IFERROR(C{last + 1}/B{last + 1},0)", f"=SUM(E4:E{last})"], 1):
        c = wsv.cell(row=last + 1, column=j, value=v)
        c.font, c.fill, c.border = FB, TOT, BOX
    wsv.cell(row=last + 1, column=3).number_format = BRL
    wsv.cell(row=last + 1, column=4).number_format = BRL
    wsv.cell(row=last + 1, column=5).number_format = "0.0%"
    larguras(wsv, [42, 13, 18, 16, 13])

    wb.save(destino)


if __name__ == "__main__":
    estabs = [int(x) for x in sys.argv[1:]] or [2]
    atend, faltas = coletar(estabs)
    out = os.environ.get("BELLE_OUT", "Faturamento_Sabados_Drenesse.xlsx")
    montar(atend, faltas, out)
    print(f"{len(atend)} atendimentos, {len(faltas)} faltas -> {out}")
