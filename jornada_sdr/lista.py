"""Lista de clientes para o SDR: inativas por faixa (30/60/90+ dias sem vir), valor gasto, saldo de sessões e próximo passo."""
import collections
import datetime as dt
import glob
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from base import D, HOJE, compras_por_cliente, data, planos

ps = planos()
C = compras_por_cliente(ps)
INI24 = HOJE - dt.timedelta(days=730)

# cross-sell mais frequente depois de cada categoria (matriz calculada na análise de 12 meses)
CROSS = {"Crioturbo": "Stimulus / Lipo Gessada / Crioturbo (nova área)", "Drenagem": "Crioturbo / Lipo Gessada / Stimulus / Enzimas / facial de tratamento",
         "Stimulus/Glúteo": "Crioturbo / Lipo Gessada", "Lipo Gessada/Modelagem": "Crioturbo / Stimulus",
         "Diástase": "Stimulus / Crioturbo / Enzimas", "Facial": "Facial de tratamento (Mellan, microagulhamento) / Crioturbo / Lipo Gessada", "Laser": "Facial de tratamento / Crioturbo / Lipo Gessada",
         "Enzimas/Injetáveis": "Crioturbo / Lipo Gessada", "Pós-Crioturbo": "Crioturbo (nova área) / Stimulus"}

cad, nomes = {}, {}
for f in glob.glob(D + "vendas_detalhado__*"):
    for v in json.load(open(f)):
        if v.get("dados_cliente") and v.get("cod_cliente"):
            cad[int(v["cod_cliente"])] = v["dados_cliente"]
for f in glob.glob(D + "venda_planos__*"):
    for p in json.load(open(f)):
        if not p.get("cliente") or " - " not in p["cliente"]:
            continue
        c, n = p["cliente"].split(" - ", 1)
        nomes[int(c)] = " ".join(n.split())

ult, unid, prox = {}, {}, {}
for f in glob.glob(D + "relatorios_relatorio_atendimentos__*"):
    if "codEstab-4" in f:
        continue
    e = {"1": "Petrópolis", "2": "Lagoa Nova", "3": "Capim Macio", "6": "Norte Shopping"}[re.search(r"codEstab-(\d+)", f).group(1)]
    for a in json.load(open(f)):
        d, c = data(a.get("dataAgendamento")), a.get("codigoCliente")
        if not d or not c:
            continue
        c = int(c)
        nomes.setdefault(c, " ".join((a.get("nomeCliente") or "").split()))
        if a.get("statusAgendamento") == "Atendido" and d <= HOJE and d > ult.get(c, dt.date.min):
            ult[c], unid[c] = d, e
        if d > HOJE and a.get("statusAgendamento") not in ("Desmarcado", "Cancelado"):
            prox[c] = min(prox.get(c, d), d)

linhas = []
for c, l in C.items():
    g24 = sum(x["v"] for x in l if x["d"] >= INI24)
    u = ult.get(c)
    if not u or c in prox:  # sem sessão na agenda desde 2023, ou já tem horário marcado
        continue
    dias = (HOJE - u).days
    if dias < 30:
        continue
    faixa = "30-59 dias" if dias < 60 else "60-89 dias" if dias < 90 else "90-179 dias" if dias < 180 else "180-365 dias" if dias <= 365 else None
    if not faixa:
        continue
    ativos = [p for x in l if x["d"] >= HOJE - dt.timedelta(days=365) for p in x["planos"]]
    saldo = sum(p["restante"] for p in ativos)
    nivel = "A (≥ R$ 6 mil)" if g24 >= 6060 else "B (R$ 2,5–6 mil)" if g24 >= 2553 else "C (< R$ 2,5 mil)"
    cat = l[-1]["cat"]
    dc = cad.get(c) or {}
    tel = dc.get("celular") or dc.get("telefone") or ""
    if saldo > 0:
        passo = "Retomar sessões do plano (saldo) + acompanhamento estético"
    elif dias < 60:
        passo = "Reavaliação com proposta de upgrade (procedimento de maior valor)"
    elif dias < 90:
        passo = "Presente (Spa Face) + reavaliação; oferta de upgrade"
    else:
        passo = "Reativação: vídeo/convite + reavaliação; condição de aquisição só se ≥ 6 meses"
    linhas.append([nivel, faixa, dias, nomes.get(c, ""), tel, unid.get(c, ""), u, l[-1]["d"], cat, round(g24, 2),
                   round(sum(x["v"] for x in l), 2), len(l), saldo, CROSS.get(cat, "Crioturbo / Lipo Gessada / Stimulus / Enzimas"), passo, c])

ordem_f = {"30-59 dias": 0, "60-89 dias": 1, "90-179 dias": 2, "180-365 dias": 3}
linhas.sort(key=lambda r: (r[0][0], ordem_f[r[1]], -r[9]))
print(len(linhas), collections.Counter((r[0][0], r[1]) for r in linhas))

wb = Workbook()
ws = wb.active
ws.title = "Resumo"
ws["A1"] = "Base de reativação – clientes de plano sem vir há 30 a 365 dias e sem horário marcado"
ws["A1"].font = Font(bold=True, size=12)
ws["A2"] = (f"Fonte: Belle (planos desde 2021, agenda desde 2023), extraído em {HOJE:%d/%m/%Y}. Nível pelo gasto em planos nos últimos 24 meses. "
            "Saldo = sessões restantes de planos comprados nos últimos 12 meses.")
ws["A2"].font = Font(italic=True, size=9)
cnt = collections.Counter((r[0], r[1]) for r in linhas)
val = collections.Counter()
for r in linhas:
    val[(r[0], r[1])] += r[9]
ws.append([])
ws.append(["Nível", "30-59 dias", "60-89 dias", "90-179 dias", "180-365 dias", "Total"])
for c_ in ws[4]:
    c_.font = Font(bold=True)
for nv in ("A (≥ R$ 6 mil)", "B (R$ 2,5–6 mil)", "C (< R$ 2,5 mil)"):
    ws.append([nv] + [cnt[(nv, f)] for f in ordem_f] + [sum(cnt[(nv, f)] for f in ordem_f)])
ws.append([])
ws.append(["Prioridade sugerida: A e B em 30-89 dias primeiro (maior chance de voltar: 24% e 14% voltam em 60 dias), depois A em 90+."])
ws.column_dimensions["A"].width = 22

wl = wb.create_sheet("Clientes")
cols = ["Nível (gasto 24m)", "Faixa sem vir", "Dias sem vir", "Cliente", "Celular", "Unidade (última sessão)", "Última sessão",
        "Última compra de plano", "Categoria da última compra", "Gasto 24m (R$)", "Gasto total (R$)", "Nº de compras",
        "Saldo de sessões", "Upgrade sugerido", "Próximo passo", "Cód. Belle"]
wl.append(cols)
for r in linhas:
    wl.append(r)
for i, w in enumerate([16, 13, 10, 34, 16, 16, 12, 12, 22, 13, 13, 9, 9, 36, 52, 11], 1):
    wl.column_dimensions[get_column_letter(i)].width = w
for row in wl.iter_rows(min_row=2):
    row[6].number_format = row[7].number_format = "dd/mm/yyyy"
    row[9].number_format = row[10].number_format = "#,##0.00"
for c_ in wl[1]:
    c_.font = Font(bold=True, color="FFFFFF")
    c_.fill = PatternFill("solid", fgColor="C8003C")
wl.add_table(Table(displayName="Reativacao", ref=f"A1:{get_column_letter(len(cols))}{len(linhas) + 1}",
                   tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))
wl.freeze_panes = "E2"
wb.save("Base_Reativacao_SDR.xlsx")
