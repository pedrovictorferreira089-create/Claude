"""Planilha de planos com desconto acima de 50% (clientes ativos) e planos de 2026 com validade acima de 12 meses:
principais clientes, quantidade, colaboradora que mais vende e unidade com maior incidência.

Uso: python belle_desc_validade.py [saida.xlsx]
Com BELLE_PKL apontando para um pickle (tabela, planos, linhas, agenda) já coletado, não chama a API."""
import collections
import datetime as dt
import os
import pickle
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment

import belle_auditoria as au
from belle_sabados import BRL, F, FB, FI, FT, INT, SUB, cabecalho, larguras

PCT = "0.0%"
DATA = "DD/MM/YYYY"
UN = list(au.UNIDS.values())
ANO = 2026
WRAP = Alignment(wrap_text=True, vertical="top")


def dados():
    if os.environ.get("BELLE_PKL"):
        tabela, planos, _, _ = pickle.load(open(os.environ["BELLE_PKL"], "rb"))
    else:
        tabela, planos, _, _ = au.coletar()
    return tabela, planos, au.auditar_planos(planos, tabela, dt.date(2020, 12, 1))


def ativo(p):
    """Plano ativo: ainda tem sessões a realizar e não venceu."""
    return p["restantes"] > 0 and (p["validade"] is None or p["validade"] >= au.HOJE)


def desc50(p):
    return p["tabela"] > 0 and 1 - p["final"] / p["tabela"] > 0.5


def servicos(planos, p):
    return " · ".join(f'{(s["nomeServico"] or "").strip()} ({s["qtdSessoes"] or 0})' for s in planos[p["orc"]]["raw"]["servicos"])


def celula(ws, r, c, v, fmt=None, bold=False, fill=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = FB if bold else F
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    return x


def auxiliares(P):
    """Colunas R e S da lista: 1 na primeira linha de cada cliente (geral e dentro da unidade), para contar clientes diferentes."""
    vistos, vistos_u = set(), set()
    out = []
    for p in P:
        out.append([int(p["cliente"] not in vistos), int((p["cliente"], p["unidade"]) not in vistos_u)])
        vistos.add(p["cliente"])
        vistos_u.add((p["cliente"], p["unidade"]))
    return out


AUX = ["Conta cliente (1ª linha)", "Conta cliente na unidade (1ª linha)"]


def lista(ws, cols, rows, fmts, larg):
    cabecalho(ws, 1, cols)
    for i, vals in enumerate(rows, 2):
        for j, (v, f) in enumerate(zip(vals, fmts), 1):
            celula(ws, i, j, v, f)
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cols)).column_letter}{max(2, len(rows) + 1)}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    larguras(ws, larg)


def resumo(ws, titulo, sub, notas, aba, n, base_un, base_rot):
    """Aba-resumo com fórmulas sobre a aba-lista `aba` (n linhas). Colunas da lista: A cliente, B unidade,
    C colaboradora, H tabela, I final, J desconto R$, M sessões restantes, N a realizar."""
    rng = lambda c: f"'{aba}'!${c}$2:${c}${n + 1}"  # noqa: E731
    ws["A1"], ws["A2"] = titulo, sub
    ws["A1"].font, ws["A2"].font = FT, FI
    r = 3
    for t in notas:
        celula(ws, r, 1, t).font = FI
        r += 1
    r += 1
    celula(ws, r, 1, "A. Visão geral", bold=True)
    r += 1
    ger = r
    for rot, f, fmt in [("Planos encontrados", f"=COUNTA({rng('A')})", INT),
                        ("Clientes diferentes", f"=SUM({rng('R')})", INT),
                        ("Valor de tabela", f"=SUM({rng('H')})", BRL), ("Valor final vendido", f"=SUM({rng('I')})", BRL),
                        ("Desconto (R$)", f"=SUM({rng('J')})", BRL), ("Desconto médio", f"=B{r + 4}/B{r + 2}", PCT),
                        ("Sessões ainda a realizar", f"=SUM({rng('M')})", INT), ("Valor das sessões a realizar", f"=SUM({rng('N')})", BRL)]:
        celula(ws, r, 1, rot)
        celula(ws, r, 2, f, fmt, bold=True)
        r += 1
    r += 1

    celula(ws, r, 1, "B. Por unidade (incidência)", bold=True)
    r += 1
    cabecalho(ws, r, ["Unidade", "Planos encontrados", base_rot, "Incidência", "Clientes", "% dos planos encontrados", "Desconto (R$)", "Valor final"])
    r += 1
    ini = r
    for u in UN:
        vals = [u, f"=COUNTIF({rng('B')},A{r})", base_un[u], f"=IF(C{r}=0,0,B{r}/C{r})",
                f"=SUMIF({rng('B')},A{r},{rng('S')})",
                f"=IF(B{r + len(UN) - UN.index(u)}=0,0,B{r}/B{r + len(UN) - UN.index(u)})",
                f"=SUMIF({rng('B')},A{r},{rng('J')})", f"=SUMIF({rng('B')},A{r},{rng('I')})"]
        for j, (v, fm) in enumerate(zip(vals, [None, INT, INT, PCT, INT, PCT, BRL, BRL]), 1):
            celula(ws, r, j, v, fm)
        r += 1
    tot = ["Total", f"=SUM(B{ini}:B{r - 1})", f"=SUM(C{ini}:C{r - 1})", f"=IF(C{r}=0,0,B{r}/C{r})", f"=SUM(E{ini}:E{r - 1})", f"=IF(B{r}=0,0,1)",
           f"=SUM(G{ini}:G{r - 1})", f"=SUM(H{ini}:H{r - 1})"]
    for j, (v, fm) in enumerate(zip(tot, [None, INT, INT, PCT, INT, PCT, BRL, BRL]), 1):
        celula(ws, r, j, v, fm, bold=True, fill=SUB)
    r += 2
    return r, rng, ger


def ranking_colab(ws, r, rng, colabs, principal):
    celula(ws, r, 1, "C. Colaboradoras que mais vendem (todas, da que mais vendeu para a que menos vendeu)", bold=True)
    r += 1
    cabecalho(ws, r, ["Colaboradora", "Planos", "% do total", "Unidade principal", "Planos nessa unidade", "Desconto (R$)", "Valor final"])
    r += 1
    ini, fim = r, r + len(colabs) - 1
    for v in colabs:
        u = principal[v]
        vals = [v, f"=COUNTIF({rng('C')},A{r})", f"=B{r}/SUM($B${ini}:$B${fim})", u, f"=COUNTIFS({rng('C')},A{r},{rng('B')},D{r})",
                f"=SUMIF({rng('C')},A{r},{rng('J')})", f"=SUMIF({rng('C')},A{r},{rng('I')})"]
        for j, (x, fm) in enumerate(zip(vals, [None, INT, PCT, None, INT, BRL, BRL]), 1):
            celula(ws, r, j, x, fm)
        r += 1
    return r + 1


def ranking_clientes(ws, r, rng, clientes, titulo, extra):
    celula(ws, r, 1, titulo, bold=True)
    r += 1
    cabecalho(ws, r, ["Cliente", "Unidade", "Planos", "Valor de tabela", "Valor final", "Desconto (R$)", "Desconto %",
                      "Sessões a realizar", "Valor a realizar", "Colaboradora(s)"] + [e[0] for e in extra])
    r += 1
    for c in clientes:
        vals = [c["cliente"], c["unidade"], f"=COUNTIF({rng('A')},A{r})", f"=SUMIF({rng('A')},A{r},{rng('H')})",
                f"=SUMIF({rng('A')},A{r},{rng('I')})", f"=D{r}-E{r}", f"=IF(D{r}=0,0,F{r}/D{r})",
                f"=SUMIF({rng('A')},A{r},{rng('M')})", f"=SUMIF({rng('A')},A{r},{rng('N')})", c["colabs"]] + [c[e[1]] for e in extra]
        for j, (x, fm) in enumerate(zip(vals, [None, None, INT, BRL, BRL, BRL, PCT, INT, BRL, None] + [None] * len(extra)), 1):
            celula(ws, r, j, x, fm)
        r += 1
    return r + 1


def por_cliente(P):
    g = collections.defaultdict(list)
    for p in P:
        g[p["cliente"]].append(p)
    out = []
    for c, ps in g.items():
        out.append(dict(cliente=c, unidade=", ".join(sorted({p["unidade"] for p in ps})), planos=len(ps),
                        desconto=sum(p["tabela"] - p["final"] for p in ps), final=sum(p["final"] for p in ps),
                        colabs=", ".join(sorted({p["vendedor"] for p in ps})),
                        ativo="Sim" if any(ativo(p) for p in ps) else "Não",
                        validade=max(p["validade"] for p in ps).strftime("%d/%m/%Y") if all(p["validade"] for p in ps) else ""))
    return out


def colab_principal(P):
    cont = collections.Counter(p["vendedor"] for p in P)
    pu = {v: collections.Counter(p["unidade"] for p in P if p["vendedor"] == v).most_common(1)[0][0] for v in cont}
    return [v for v, _ in sorted(cont.items(), key=lambda x: (-x[1], x[0]))], pu


def montar(destino):
    tabela, planos, pa = dados()
    A = [p for p in pa if ativo(p)]
    D = sorted([p for p in A if desc50(p)], key=lambda p: (-(p["tabela"] - p["final"]), p["cliente"]))
    P26 = [p for p in pa if p["venda"].year == ANO]
    V = sorted([p for p in P26 if p["meses"] is not None and p["meses"] > 12], key=lambda p: (p["venda"], p["cliente"]))

    wb = Workbook()
    ws, wl = wb.active, wb.create_sheet("Planos com desconto > 50")
    ws.title = "Desconto > 50"
    wv, wvl = wb.create_sheet("Validade > 1 ano 2026"), wb.create_sheet("Planos validade 2026")

    # ---------------- lista: desconto > 50%, ativos
    cols = ["Cliente", "Unidade", "Colaboradora (venda)", "Plano", "Data da venda", "Ano da venda", "Validade", "Valor de tabela", "Valor final",
            "Desconto (R$)", "Desconto %", "Sessões contratadas", "Sessões a realizar", "Valor a realizar", "Parcelas vencidas não pagas", "Serviços (sessões)", "Orçamento"]
    rows = [[p["cliente"], p["unidade"], p["vendedor"], p["plano"], p["venda"], p["venda"].year, p["validade"], p["tabela"], p["final"],
             f"=H{i}-I{i}", f"=IF(H{i}=0,0,J{i}/H{i})", p["sessoes"], p["restantes"], p["a_realizar"], p["vencido"], servicos(planos, p), str(p["orc"])]
            for i, p in enumerate(D, 2)]
    rows = [r + a for r, a in zip(rows, auxiliares(D))]
    lista(wl, cols + AUX, rows, [None, None, None, None, DATA, "0", DATA, BRL, BRL, BRL, PCT, INT, INT, BRL, BRL, None, None, "0", "0"],
          [34, 15, 30, 30, 12, 8, 12, 14, 14, 14, 10, 11, 11, 14, 14, 70, 12, 12, 12])

    notas = [f"Planos aprovados ATIVOS em {au.HOJE:%d/%m/%Y} (com sessões a realizar e validade não vencida), de qualquer ano de venda, "
             "com desconto acima de 50% sobre o valor de tabela (sessões × preço atual da tabela de serviços).",
             "Colaboradora = quem fez a venda do plano. Incidência = planos com desconto > 50% ÷ planos ativos da unidade.",
             "Lista completa na aba 'Planos com desconto > 50'. Os números desta aba são fórmulas sobre aquela lista."]
    r, rng, _ = resumo(ws, "Planos com desconto acima de 50% – clientes ativos", "Drenesse · planos aprovados vendidos de 2021 a 29/09/2026",
                       notas, wl.title, len(D), {u: sum(1 for p in A if p["unidade"] == u) for u in UN}, "Planos ativos da unidade")
    # ano da venda
    celula(ws, r, 1, "D. Ano da venda desses planos", bold=True)
    r += 1
    cabecalho(ws, r, ["Ano da venda", "Planos", "% do total"])
    r += 1
    ini = r
    anos = sorted({p["venda"].year for p in D})
    for y in anos:
        celula(ws, r, 1, y, "0")
        celula(ws, r, 2, f"=COUNTIF({rng('F')},A{r})", INT)
        celula(ws, r, 3, f"=B{r}/SUM($B${ini}:$B${ini + len(anos) - 1})", PCT)
        r += 1
    r += 1
    colabs, pu = colab_principal(D)
    r = ranking_colab(ws, r, rng, colabs, pu)
    cli = sorted(por_cliente(D), key=lambda c: (-c["desconto"], c["cliente"]))[:50]
    ranking_clientes(ws, r, rng, cli, "E. Principais clientes ativos (50 com maior desconto em R$ somando os planos com desconto > 50%)", [])
    larguras(ws, [40, 16, 16, 18, 14, 16, 16, 14, 14, 40])

    # ---------------- lista: validade > 12 meses, vendas 2026
    cols = ["Cliente", "Unidade", "Colaboradora (venda)", "Plano", "Data da venda", "Validade", "Meses de validade", "Valor de tabela", "Valor final",
            "Desconto (R$)", "Desconto %", "Sessões contratadas", "Sessões a realizar", "Valor a realizar", "Ativo", "Serviços (sessões)", "Orçamento"]
    rows = [[p["cliente"], p["unidade"], p["vendedor"], p["plano"], p["venda"], p["validade"], p["meses"], p["tabela"], p["final"],
             f"=H{i}-I{i}", f"=IF(H{i}=0,0,J{i}/H{i})", p["sessoes"], p["restantes"], p["a_realizar"], "Sim" if ativo(p) else "Não",
             servicos(planos, p), str(p["orc"])] for i, p in enumerate(V, 2)]
    rows = [r + a for r, a in zip(rows, auxiliares(V))]
    lista(wvl, cols + AUX, rows, [None, None, None, None, DATA, DATA, "0", BRL, BRL, BRL, PCT, INT, INT, BRL, None, None, None, "0", "0"],
          [34, 15, 30, 30, 12, 12, 10, 14, 14, 14, 10, 11, 11, 14, 8, 80, 12, 12, 12])
    notas = [f"Planos aprovados vendidos em {ANO} (01/01 a 29/09) com validade superior a 12 meses (da data da venda até a validade do plano).",
             f"Incidência = planos com validade > 12 meses ÷ planos aprovados vendidos em {ANO} pela unidade. Ativo = ainda tem sessões a realizar e não venceu.",
             "Lista completa na aba 'Planos validade 2026'. Os números desta aba são fórmulas sobre aquela lista."]
    r, rng, ger = resumo(wv, f"Planos com validade acima de 1 ano – vendas de {ANO}", "Drenesse · validade conforme relatório Sessões de Planos",
                         notas, wvl.title, len(V), {u: sum(1 for p in P26 if p["unidade"] == u) for u in UN}, f"Planos vendidos em {ANO}")
    celula(wv, ger + 8, 1, "Clientes com plano ativo")
    celula(wv, ger + 8, 2, f"=SUMPRODUCT(({rng('O')}=\"Sim\")/COUNTIFS({rng('A')},{rng('A')},{rng('O')},{rng('O')}))", INT, bold=True)
    colabs, pu = colab_principal(V)
    r = ranking_colab(wv, r + 1, rng, colabs, pu)
    cli = sorted(por_cliente(V), key=lambda c: (-c["final"], c["cliente"]))
    ranking_clientes(wv, r, rng, cli, f"D. Clientes (todos os {len(cli)}, do maior para o menor valor vendido)",
                     [("Ativo", "ativo"), ("Validade (mais longa)", "validade")])
    larguras(wv, [40, 16, 16, 18, 14, 16, 16, 14, 14, 40, 8, 14])

    for w in (ws, wv):
        for row in w.iter_rows():
            for c in row:
                if c.column == 10 or (c.column == 1 and c.row > 6):
                    c.alignment = WRAP
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    return D, V, A, P26


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Descontos_e_Validade_Drenesse.xlsx"
    D, V, A, P26 = montar(out)
    print(f"{len(D)} planos ativos com desconto > 50% ({len({p['cliente'] for p in D})} clientes) | "
          f"{len(V)} planos de {ANO} com validade > 12 meses -> {out}")
