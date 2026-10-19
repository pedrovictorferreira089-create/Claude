"""Clientes com plano ativo que só agendam aos sábados (API Belle Software).

Uso: python3 belle_sabado_clientes.py [saida.xlsx]
Plano ativo = plano Aprovado, validade não vencida e com sessões ainda a fazer.
Só sábado = todos os agendamentos da cliente (atendidos, faltas e marcados) dos últimos 90 dias e da agenda futura
caem em sábado; desmarcações e cancelamentos não contam. Uma linha por procedimento com sessões a fazer.
"""
import collections
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill

from belle_sabados import (BOX, F, FB, FI, FT, HOJE, INT, SUB, UNIDADES, api, cabecalho, historico, larguras)

JANELA = HOJE - dt.timedelta(days=90)
FIM_AGENDA = dt.date(2027, 12, 31)
CONTAM = {"Atendido", "Falhou", "Marcado", "Confirmado", "Aguardando", "Em Andamento"}
FUTURO = {"Marcado", "Confirmado", "Aguardando", "Em Andamento"}
DATA = "dd/mm/yyyy"
FAIXA = PatternFill("solid", fgColor="F2F2F2")


def data_br(s):
    return dt.datetime.strptime(s, "%d/%m/%Y").date() if s else None


def agenda_futura(e):
    out, ini = [], HOJE + dt.timedelta(days=1)
    while ini <= FIM_AGENDA:
        fim = min(dt.date(ini.year + (ini.month + 2) // 12, (ini.month + 2) % 12 + 1, 1) - dt.timedelta(days=1), FIM_AGENDA)
        out += api("relatorios/relatorio_atendimentos", dtInicio=ini.strftime("%d/%m/%Y"), dtFim=fim.strftime("%d/%m/%Y"), codEstab=str(e))
        ini = fim + dt.timedelta(days=1)
    return out


def telefone(c):
    nums = []
    for k in ("celular", "celular2", "telefone"):
        n = (c.get(k) or "").replace("(+55) ", "").strip()
        if n and n not in nums:
            nums.append(n)
    return " / ".join(nums)


def coletar():
    ags, sessoes = {}, {}
    for e in UNIDADES:
        for a in historico("relatorios/relatorio_atendimentos", "codEstab", e, HOJE) + agenda_futura(e):
            ags[a["idAgendamento"]] = dict(a, estab=e)
        for x in historico("relatorios/sessoes_planos", "codEstab", e, HOJE, p_ini="dtVndIni", p_fim="dtVndFim"):
            sessoes[(x["idPlano"], str(x["idServico"]))] = x
    futuros = collections.Counter((a["idOrcamento"], str(a["codigoServico"])) for a in ags.values()
                                  if a["statusAgendamento"] in FUTURO and a["idOrcamento"] and data_br(a["dataAgendamento"]) >= HOJE)
    por_cliente = collections.defaultdict(list)
    for a in ags.values():
        if a["statusAgendamento"] in CONTAM and data_br(a["dataAgendamento"]) >= JANELA:
            por_cliente[a["codigoCliente"]].append(a)

    linhas = []
    for (orc, cod), x in sessoes.items():
        validade = data_br(x["dtValidade"])
        if x["statusPlano"] != "Aprovado" or (validade and validade < HOJE):
            continue
        vendidas, livres, marcadas = int(x["sessoesVendidas"]), int(x["sessoesRestante"]), futuros[(orc, cod)]
        if livres + marcadas <= 0:
            continue
        lst = por_cliente[x["idCliente"]]
        if not lst or any(data_br(a["dataAgendamento"]).weekday() != 5 for a in lst):
            continue
        linhas.append(dict(cod=x["idCliente"], plano=orc, procedimento=" ".join(x["nomeServico"].split()),
                           venda=data_br(x["dtVenda"]), validade=validade, vendidas=vendidas,
                           feitas=max(vendidas - livres - marcadas, 0), marcadas=marcadas))

    clientes = {}
    for c in {l["cod"] for l in linhas}:
        cad = api("clientes", codEstab="1", id=str(c))
        cad = cad[0] if isinstance(cad, list) and cad else {}
        lst = sorted(por_cliente[c], key=lambda a: (data_br(a["dataAgendamento"]), a["horarioAgendamento"] or ""))
        feitos = [a for a in lst if a["statusAgendamento"] == "Atendido"]
        prox = [a for a in lst if a["statusAgendamento"] in FUTURO and data_br(a["dataAgendamento"]) >= HOJE]
        clientes[c] = dict(
            nome=" ".join((cad.get("nome") or lst[0]["nomeCliente"]).split()), fone=telefone(cad),
            unidade=", ".join(sorted({UNIDADES[a["estab"]] for a in lst})),
            sabados=len({a["dataAgendamento"] for a in lst if data_br(a["dataAgendamento"]) <= HOJE}),
            ultimo=data_br(feitos[-1]["dataAgendamento"]) if feitos else None,
            proximo=data_br(prox[0]["dataAgendamento"]) if prox else None,
            horario=collections.Counter(a["horarioAgendamento"] for a in lst).most_common(1)[0][0])
    for l in linhas:
        l.update(clientes[l["cod"]])
    return sorted(linhas, key=lambda l: (l["nome"].casefold(), l["venda"] or HOJE, l["procedimento"]))


def montar(linhas, destino):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes só sábado"
    ws["A1"] = "Drenesse – Clientes com plano ativo que só agendam aos sábados"
    ws["A1"].font = FT
    ws["A2"] = (f"Fonte: API Belle Software, extraído em {HOJE:%d/%m/%Y}. Plano ativo = aprovado, dentro da validade e com sessões a fazer. "
                f"Só sábado = todos os agendamentos de {JANELA:%d/%m/%Y} em diante (atendidos, faltas e já marcados) caem em sábado; "
                "desmarcações não contam. Uma linha por procedimento do plano.")
    ws["A2"].font = FI
    ws["A3"] = "Restantes = contratadas − feitas (inclui as sessões que já estão marcadas na agenda)."
    ws["A3"].font = FI
    cols = ["Cliente", "Telefone", "Procedimento", "Unidade", "Sessões contratadas", "Sessões feitas", "Sessões restantes",
            "Restantes já marcadas", "Validade do plano", "Sábados com agendamento (últ. 90 dias)", "Último atendimento",
            "Próximo agendamento", "Horário habitual", "Nº plano"]
    h = 5
    cabecalho(ws, h, cols)
    fmts = [None, None, None, None, INT, INT, INT, INT, DATA, INT, DATA, DATA, None, "0"]
    faixa, anterior = False, None
    for i, l in enumerate(linhas, h + 1):
        if l["cod"] != anterior:  # faixa de cor alternada por cliente
            faixa, anterior = not faixa, l["cod"]
        vals = [l["nome"], l["fone"], l["procedimento"], l["unidade"], l["vendidas"], l["feitas"], f"=E{i}-F{i}", l["marcadas"],
                l["validade"], l["sabados"], l["ultimo"], l["proximo"], l["horario"], l["plano"]]
        for j, (v, fm) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = F, BOX
            if fm:
                c.number_format = fm
            if faixa:
                c.fill = FAIXA
    t = h + len(linhas) + 1
    ult = t - 1
    tot = ["TOTAL", f'=SUMPRODUCT((A{h + 1}:A{ult}<>"")/COUNTIF(A{h + 1}:A{ult},A{h + 1}:A{ult}&""))&" clientes"',
           f'=COUNTA(C{h + 1}:C{ult})&" procedimentos"', None] + [f"=SUM({c}{h + 1}:{c}{ult})" for c in "EFGH"]
    for j, v in enumerate(tot, 1):
        c = ws.cell(row=t, column=j, value=v)
        c.font, c.fill, c.border = FB, SUB, BOX
        if j >= 5:
            c.number_format = INT
    ws.auto_filter.ref = f"A{h}:{ws.cell(row=h, column=len(cols)).column_letter}{ult}"
    ws.freeze_panes = ws.cell(row=h + 1, column=2)
    ws.row_dimensions[h].height = 42
    for r in ws.iter_rows(min_row=h + 1, max_row=ult, min_col=3, max_col=3):
        r[0].alignment = Alignment(wrap_text=True, vertical="top")
    larguras(ws, [38, 30, 46, 16, 12, 10, 11, 12, 12, 14, 13, 13, 10, 12])
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.environ.get("BELLE_OUT", "Clientes_So_Sabado_Drenesse.xlsx")
    linhas = coletar()
    montar(linhas, out)
    print(f"{len({l['cod'] for l in linhas})} clientes, {len(linhas)} procedimentos -> {out}")
